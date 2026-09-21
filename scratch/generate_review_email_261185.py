import os
import win32clipboard
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_out = os.path.join(desktop_dir, "OOS-261185 Review Email.docx")
html_out = os.path.join(desktop_dir, "OOS-261185_Review_Email_Robin_Format.html")

# --- 1. GENERATE WORD DOCX (FULL SENTENCES, NO BULLETS, ZERO INDENT, LEFT ALIGNED) ---
doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Greeting
p1 = doc.add_paragraph()
p1.paragraph_format.left_indent = Inches(0)
p1.paragraph_format.space_before = Pt(0)
p1.paragraph_format.space_after = Pt(6)
r1 = p1.add_run("Good morning @Simin Mohammad,")
r1.font.name = "Calibri"
r1.font.size = Pt(11)

p2 = doc.add_paragraph()
p2.paragraph_format.left_indent = Inches(0)
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after = Pt(12)
r2 = p2.add_run(
    "I have reviewed this OOS. Please find the summary of edits in the table below. "
    "All required updates, including the amended EM table and all form field revisions, have already been completed "
    "on your behalf and assembled into the attached finalized 7-page PDF package."
)
r2.font.name = "Calibri"
r2.font.size = Pt(11)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_borders(cell):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcBorders'):
            tcPr.remove(child)
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

table_data = [
    ["Section", "Unreviewed", "Reviewed", "Impact"],
    [
        "EM Table Attachment & Formatting",
        "The EM Table was omitted from the PDF submission. In the draft Word document, Table 1 mistakenly listed ETX-260526-0461 and Talaromyces purpurogenus instead of the actual isolates, Table 2 contained formatting inconsistencies such as unclosed parentheses in timing descriptions, and multi-organism entries ran together without line breaks.",
        "Amended both tables to reflect the correct ETX submission ID ETX-260518-0250, updated the microbial identification to Penicillium decumbens, Cladosporium tenuissimum, Cladosporium langeronii, and Cladosporium halotolerans, corrected all timing descriptions, properly formatted multi-organism line breaks, and attached the amended 1-page table as Page 7 to complete the 7-page package.",
        "Document Assembly & Technical Uniformity Correction. Completed on analyst's behalf."
    ],
    [
        "Initiator Name",
        "Listed as \"Simin Mohammad (written by Simin Mohammad)\".",
        "Revised to \"Simin Mohammad\" to remove redundant parenthetical phrasing.",
        "Clerical Correction."
    ],
    [
        "Name of Analyst who Performed the Test",
        "Listed setup and reader analysts running together with lowercase \"weekly\" and plural \"Readers\": \"Simin Mohammad (weekly Active air Sampling Plate Setup) Maraya Chukwumerije & Sophia Santamaria(Weekly active air Sampling Plate Readers)\".",
        "Formatted into distinct, properly capitalized entries separating Simin Mohammad as the setup analyst from Maraya Chukwumerije and Sophia Santamaria as individual plate readers:\nSimin Mohammad (Weekly Active Air Sampling Plate Setup)\nMaraya Chukwumerije (Weekly Active Air Sampling Plate Reader)\nSophia Santamaria (Weekly Active Air Sampling Plate Reader)",
        "Formatting & Administrative Correction."
    ],
    [
        "SOP Reference & Effective Date",
        "Listed as \"20600.002\" Revision 15 with Effective Date \"05AUG2025\".",
        "Revised to \"MICRO-SOP-2\" Revision 16 with Effective Date \"23-Jul-2026\".",
        "Compliance & Document Control Correction."
    ],
    [
        "Limits / Specification",
        "Listed as \"Action level >10\".",
        "Revised to \"Action level: >= 10 CFU/Plate\".",
        "Technical Specification Correction."
    ],
    [
        "Section B SOP Citations",
        "SOP reference in comments contained a typographical error listed as \"SOP 2.600.00\" missing the trailing digit.",
        "Revised all instances to \"MICRO-SOP-2\".",
        "SOP Citation Correction."
    ],
    [
        "Analyst Interview Comment",
        "\"Yes, analysts Simin Mohammad, Maraya Chukwumerije, & Sophia Santamaria were comprehensively interviewed.\"",
        "Revised ampersand to \"and\": \"Yes, analysts Simin Mohammad, Maraya Chukwumerije, and Sophia Santamaria were comprehensively interviewed.\"",
        "Grammatical Correction."
    ],
    [
        "Equipment & Calibration Details",
        "Equipment IDs and calibration dates were running together without proper spacing: \"Incubator E001034(Sensor E001501)Incubator E001031 (Sensor E001505)\" and \"Aug 2026 Feb 2027  Aug 2026 Feb 2027\".",
        "Revised with clean line breaks to clearly delineate Incubator E001034 with Sensor E001501 and Incubator E001031 with Sensor E001505 alongside their respective calibration dates of August 2026 and February 2027.",
        "Equipment Traceability & Readability Correction."
    ],
    [
        "Phase I Summary – Setup & Incubation Narrative",
        "Contained a spelling error in reader Sophia Santamaria's surname (\"Sophia Sanatamaria\"), grammatical disagreement regarding incubator functionality (\"incubators were verified\"), and an inaccurate reference to biosafety cabinet locations instead of cleanroom suites.",
        "Corrected the analyst's surname to Sophia Santamaria, corrected grammatical phrasing to indicate incubator functionality was verified, updated cleanroom locations to Suite 115 (ISO 8), Suite 115A (ISO 7), and Suite 115B (ISO 7), and cited MICRO-SOP-2 and MICRO-SOP-9.",
        "Grammatical & Scope Correction."
    ],
    [
        "Phase I Summary – Monthly Cleaning & Disinfection Narrative",
        "Contained contradictory monthly cleaning dates of 22 Feb 2026 and 26APR2026 within the same sentence, alongside a missing period and legacy SOP citations.",
        "Reconciled the cleaning date to 26-Apr-2026 performed by Tamiru Kotisso and Cuong Du, inserted the missing period before the H2O2 indicators statement, and updated the procedure reference to MICRO-SOP-9.",
        "Chronological Accuracy & Compliance Correction."
    ],
    [
        "Phase I Summary – Root Cause Statement",
        "Stated that the OOS result observed for the Environmental Monitoring (EM) Settling Sampling plate may be attributed to a potential analyst error.",
        "Corrected the sampling plate type from Settling Sampling plate to Active Air Sampling plate to accurately reflect the test performed, while maintaining the determination that the result may be attributed to a potential analyst error, noting that the contamination was transient in nature and no corrective actions are necessary.",
        "Sampling Type & Technical Accuracy Correction."
    ]
]

col_widths = [1.3, 2.0, 2.1, 1.1]
table = doc.add_table(rows=len(table_data), cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = False

tblPr = table._element.xpath('w:tblPr')
if tblPr:
    for ind in tblPr[0].xpath('w:tblInd'):
        tblPr[0].remove(ind)
    tblPr[0].append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>'))

for r_idx, row_content in enumerate(table_data):
    row = table.rows[r_idx]
    is_header = (r_idx == 0)
    
    for c_idx, cell_text in enumerate(row_content):
        cell = row.cells[c_idx]
        cell.width = Inches(col_widths[c_idx])
        set_cell_borders(cell)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        
        if is_header:
            set_cell_background(cell, "FFFF00")
        else:
            set_cell_background(cell, "FFFFFF")
            
        p = cell.paragraphs[0]
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        if is_header:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(cell_text)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            lines = cell_text.split('\n')
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                
                r = p.add_run(line)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                if c_idx == 0:
                    r.bold = True

p3 = doc.add_paragraph()
p3.paragraph_format.left_indent = Inches(0)
p3.paragraph_format.space_before = Pt(14)
p3.paragraph_format.space_after = Pt(6)
r3 = p3.add_run("@Simin Mohammad please review and sign the attached finalized 7-page PDF package.")
r3.font.name = "Calibri"
r3.font.size = Pt(11)

p4 = doc.add_paragraph()
p4.paragraph_format.left_indent = Inches(0)
p4.paragraph_format.space_before = Pt(0)
p4.paragraph_format.space_after = Pt(0)
r4 = p4.add_run("Thanks,")
r4.font.name = "Calibri"
r4.font.size = Pt(11)

doc.save(docx_out)
print(f"Saved Word review email: {docx_out}")

# --- 2. GENERATE HTML FILE (FULL SENTENCES, NO BULLETS, ZERO INDENT, LEFT ALIGNED) ---
html_rows = ""
for r_idx, row in enumerate(table_data[1:]):
    sec, unrev, rev, imp = row
    unrev_html = unrev.replace('\n', '<br>')
    rev_html = rev.replace('\n', '<br>')
    imp_html = imp.replace('\n', '<br>')
    
    html_rows += f"""
    <tr>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; font-weight: bold; vertical-align: top;">{sec}</td>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top;">{unrev_html}</td>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top;">{rev_html}</td>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top;">{imp_html}</td>
    </tr>
    """

full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #000; line-height: 1.4; margin: 0; padding: 0; }}
  table {{ border-collapse: collapse; width: 100%; font-family: Calibri, Arial, sans-serif; font-size: 10pt; margin: 10px 0 15px 0; }}
  th {{ background-color: #FFFF00 !important; color: #000; font-weight: bold; text-align: center; border: 1px solid #B0B0B0; padding: 6px 8px; }}
  td {{ border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top; }}
</style>
</head>
<body style="margin: 0; padding: 0;">
<p style="margin: 0 0 8px 0;">Good morning @Simin Mohammad,</p>
<p style="margin: 0 0 12px 0;">I have reviewed this OOS. Please find the summary of edits in the table below. All required updates, including the amended EM table and all form field revisions, have already been completed on your behalf and assembled into the attached finalized 7-page PDF package.</p>
<table style="border-collapse: collapse; width: 100%; margin: 10px 0 15px 0;">
  <thead>
    <tr>
      <th style="background-color: #FFFF00; width: 20%;">Section</th>
      <th style="background-color: #FFFF00; width: 30%;">Unreviewed</th>
      <th style="background-color: #FFFF00; width: 32%;">Reviewed</th>
      <th style="background-color: #FFFF00; width: 18%;">Impact</th>
    </tr>
  </thead>
  <tbody>
    {html_rows}
  </tbody>
</table>
<p style="margin: 12px 0 6px 0;">@Simin Mohammad please review and sign the attached finalized 7-page PDF package.</p>
<p style="margin: 0;">Thanks,</p>
</body>
</html>
"""

with open(html_out, "w", encoding="utf-8") as f:
    f.write(full_html)
print(f"Saved clean HTML email: {html_out}")

# --- 3. COPY TO CLIPBOARD WITH ZERO BULLETS AND FULL SENTENCES ---
HEADER = """Version:0.9
StartHTML:{start_html:08d}
EndHTML:{end_html:08d}
StartFragment:{start_fragment:08d}
EndFragment:{end_fragment:08d}
"""
fragment_wrapper = f"<html><body><!--StartFragment-->{full_html}<!--EndFragment--></body></html>"
dummy_header = HEADER.format(start_html=0, end_html=0, start_fragment=0, end_fragment=0)
start_html = len(dummy_header)
start_fragment = start_html + fragment_wrapper.find("<!--StartFragment-->") + len("<!--StartFragment-->")
end_fragment = start_html + fragment_wrapper.find("<!--EndFragment-->")
end_html = start_html + len(fragment_wrapper)

final_header = HEADER.format(
    start_html=start_html,
    end_html=end_html,
    start_fragment=start_fragment,
    end_fragment=end_fragment
)
payload = (final_header + fragment_wrapper).encode('utf-8')

win32clipboard.OpenClipboard()
try:
    win32clipboard.EmptyClipboard()
    CF_HTML = win32clipboard.RegisterClipboardFormat("HTML Format")
    win32clipboard.SetClipboardData(CF_HTML, payload)
    win32clipboard.SetClipboardText("OOS-261185 Review Email Table (Full Sentences - Review and Sign)")
    print("SUCCESS: Copied full sentence HTML table to Clipboard!")
finally:
    win32clipboard.CloseClipboard()
