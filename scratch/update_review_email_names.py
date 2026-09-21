import os, docx, win32clipboard
from docx.shared import Pt

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_path = os.path.join(desktop_dir, "OOS-261186 Review Email.docx")
html_path = os.path.join(desktop_dir, "OOS-261186_Review_Email_Robin_Format.html")

# 1. Update HTML
with open(html_path, "r", encoding="utf-8") as f:
    html_text = f.read()

old_close_pattern = "please review and sign the updated OOS investigation form. You may directly use the newly compiled 8-page PDF package attached here."
new_close_text = "please review and sign the updated OOS investigation form: <b>OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf</b>."
if old_close_pattern in html_text:
    html_text = html_text.replace(old_close_pattern, new_close_text)

with open(html_path, "w", encoding="utf-8") as f:
    f.write(html_text)
print("Updated HTML file.")

# 2. Update DOCX
doc = docx.Document(docx_path)
for p in doc.paragraphs:
    if "please review and sign" in p.text:
        p.text = ""
        r1 = p.add_run("@Simin Mohammad please review and sign the updated OOS investigation form: ")
        r1.font.name = "Calibri"
        r1.font.size = Pt(11)
        r2 = p.add_run("OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf")
        r2.font.name = "Calibri"
        r2.font.size = Pt(11)
        r2.bold = True
        r3 = p.add_run(".")
        r3.font.name = "Calibri"
        r3.font.size = Pt(11)

doc.save(docx_path)
print("Updated DOCX file.")

# 3. Update Windows Clipboard
def set_html_clipboard(html_fragment):
    MARKER_BLOCK = (
        "Version:0.9\r\n"
        "StartHTML:{start_html:08d}\r\n"
        "EndHTML:{end_html:08d}\r\n"
        "StartFragment:{start_frag:08d}\r\n"
        "EndFragment:{end_frag:08d}\r\n"
    )
    prefix = "<html><body><!--StartFragment-->"
    suffix = "<!--EndFragment--></body></html>"
    dummy = MARKER_BLOCK.format(start_html=0, end_html=0, start_frag=0, end_frag=0)
    start_html = len(dummy)
    start_frag = start_html + len(prefix)
    end_frag = start_frag + len(html_fragment.encode("utf-8"))
    end_html = end_frag + len(suffix.encode("utf-8"))
    header = MARKER_BLOCK.format(start_html=start_html, end_html=end_html, start_frag=start_frag, end_frag=end_frag)
    full_data = (header + prefix + html_fragment + suffix).encode("utf-8")

    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        cf_html = win32clipboard.RegisterClipboardFormat("HTML Format")
        win32clipboard.SetClipboardData(cf_html, full_data)
        plain = "Good morning @Simin Mohammad,\n\nI have reviewed this OOS. Please find the edits in the summary table.\n\n@Simin Mohammad please review and sign the updated OOS investigation form: OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf.\n\nThanks,"
        win32clipboard.SetClipboardText(plain)
    finally:
        win32clipboard.CloseClipboard()

body_start = html_text.find('<div class="email-container">')
body_end = html_text.find('</div>', body_start) + 6
fragment = html_text[body_start:body_end]
set_html_clipboard(fragment)
print("Clipboard refreshed with CF_HTML.")
