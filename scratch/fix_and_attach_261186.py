import os
import shutil
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader, PdfWriter
import win32com.client
import fitz

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docs_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
history_dir = os.path.join(docs_dir, ".history")
os.makedirs(history_dir, exist_ok=True)

src_docx = os.path.join(desktop_dir, "EM table OOS-261186 07MAY2026 (2).docx")
src_pdf = os.path.join(desktop_dir, "OOS-261186.pdf")

# 1. Backups
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
shutil.copy2(src_docx, os.path.join(history_dir, f"EM_table_261186_backup_{timestamp}.docx"))
shutil.copy2(src_pdf, os.path.join(history_dir, f"OOS_261186_backup_{timestamp}.pdf"))
print("Backed up original DOCX and PDF to .history")

# 2. Modify DOCX
doc = Document(src_docx)

# Update Table 2 title
for p in doc.paragraphs:
    if "Table 2:" in p.text:
        p.text = ""
        r = p.add_run("Table 2: Environmental Monitoring Plates for Analyst and Cleanroom Bracketing")
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.font.underline = True
        r.bold = False

# Table 1 Edits
t1 = doc.tables[0]
for r_idx, row in enumerate(t1.rows):
    for c_idx, cell in enumerate(row.cells):
        if "07MAY   2026" in cell.text:
            cell.text = cell.text.replace("07MAY   2026", "07MAY 2026")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(6.5)
        if cell.text.strip() == "Kocuria palustris":
            cell.text = "Kocuria palustris, Gram (+) cocci"
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(6.5)

# Table 2 Edits
t2 = doc.tables[1]
# Row 3: 07MAY 2026 Active Air
cell_r3_c7 = t2.rows[3].cells[7]
if cell_r3_c7.text.strip() == "Kocuria palustris":
    cell_r3_c7.text = "Kocuria palustris,\nGram (+) cocci"
    for p in cell_r3_c7.paragraphs:
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.0)

# Row 4: 15MAY 2026 Active Air
cell_r4_c7 = t2.rows[4].cells[7]
if "Gram (+) cocci Gram-variable coccobacilli" in cell_r4_c7.text:
    cell_r4_c7.text = "Gram (+) cocci\nGram-variable coccobacilli"
    for p in cell_r4_c7.paragraphs:
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.0)

# Row 8: 15MAY 2026 Surface
cell_r8_c7 = t2.rows[8].cells[7]
if "Sporosarcina" in cell_r8_c7.text:
    cell_r8_c7.text = "Gram (+) cocci\nSporosarcina psychrophila\nBacillus pumilus"
    for p in cell_r8_c7.paragraphs:
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.0)

out_docx_desktop = os.path.join(desktop_dir, "EM table OOS-261186 07MAY2026.docx")
out_docx_scratch = os.path.join(docs_dir, "scratch", "EM table OOS-261186 07MAY2026.docx")
doc.save(out_docx_desktop)
doc.save(out_docx_scratch)
print(f"Saved corrected Word table to: {out_docx_desktop}")

# 3. Convert DOCX to PDF via Word COM
out_pdf_desktop = os.path.join(desktop_dir, "EM table OOS-261186 07MAY2026.pdf")
out_pdf_scratch = os.path.join(docs_dir, "scratch", "EM table OOS-261186 07MAY2026.pdf")

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
try:
    doc_com = word.Documents.Open(os.path.abspath(out_docx_desktop))
    doc_com.SaveAs(os.path.abspath(out_pdf_desktop), FileFormat=17)
    doc_com.Close()
    shutil.copy2(out_pdf_desktop, out_pdf_scratch)
    print(f"Exported pixel-perfect 1-page Table PDF to: {out_pdf_desktop}")
finally:
    word.Quit()

# 4. Attach Table PDF to OOS-261186.pdf
reader_form = PdfReader(src_pdf)
reader_table = PdfReader(out_pdf_desktop)

writer = PdfWriter()
# Add all pages from src_pdf (Pages 1 to 7)
for page in reader_form.pages:
    writer.add_page(page)

# Append Table PDF page(s)
for page in reader_table.pages:
    writer.add_page(page)

print(f"Total pages in assembled package: {len(writer.pages)}")

# Save to full package and overwrite original on Desktop
out_full_pdf = os.path.join(desktop_dir, "OOS-261186 (Full Package with EM Table).pdf")
with open(out_full_pdf, "wb") as f:
    writer.write(f)

with open(src_pdf, "wb") as f:
    writer.write(f)

print(f"Saved Full Package PDF to: {out_full_pdf}")
print(f"Updated original Desktop PDF: {src_pdf}")

# 5. Render fitz verification of Page 7 and Page 8
doc_fitz = fitz.open(out_full_pdf)
print(f"Fitz opened full PDF, {len(doc_fitz)} pages.")
p7_pix = doc_fitz[6].get_pixmap(dpi=150)
p7_pix.save(os.path.join(docs_dir, "scratch", "assembled_261186_p7.png"))
p8_pix = doc_fitz[7].get_pixmap(dpi=150)
p8_pix.save(os.path.join(docs_dir, "scratch", "assembled_261186_p8.png"))
print("Saved Page 7 (Version History) and Page 8 (EM Table) verification images.")
