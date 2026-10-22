import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
import win32com.client as win32
import fitz
import pypdf
import os
import shutil

sp_04sep = [
    "Staphylococcus capitis", "Micrococcus luteus", "Moraxella osloensis",
    "Corynebacterium segmentosum", "Kytococcus sedentarius", "Staphylococcus hominis",
    "Paracoccus yeei", "Brevibacterium sp. CS2", "Corynebacterium ureicelerivorans",
    "Micrococcus sp.", "Corynebacterium sanguinis", "Staphylococcus haemolyticus",
    "Corynebacterium genitalium", "Cordyceps memorabilis", "Allophoma tropica",
    "Corynebacterium jeikeium", "Pseudomonas fulva", "Lophiotremataceae",
    "Acrodontium salmoneum", "Quambalaria cyanescens", "Sordariales", "Streptococcus sp. A12"
]

sp_09sep = [
    "Corynebacterium simulans", "Staphylococcus epidermidis", "Staphylococcus pasteuri",
    "Micrococcus luteus", "Staphylococcus hominis", "Rothia dentocariosa",
    "Staphylococcus pasteuri", "Corynebacterium imitans", "Staphylococcus sp. M0911",
    "Moraxella osloensis", "Streptococcus thermophilus", "Aeromicrobium erythreum",
    "Roseomonas sp. FDAARGOS_362", "Agrococcus carbonis", "Corynebacterium segmentosum",
    "Streptomyces xiamenensis"
]

non_italic = {"sp.", "sp", "CS2", "A12", "M0911", "FDAARGOS_362", "Lophiotremataceae", "Sordariales", "Gram", "(+)", "(-)", "cocci", "rods"}

def format_microbial_cell(cell, species_list):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.90
    
    for idx, sp in enumerate(species_list):
        if idx > 0:
            p.add_run("\n")
        words = sp.split()
        for w_idx, w in enumerate(words):
            r = p.add_run(w + (" " if w_idx < len(words)-1 else ""))
            r.font.name = "Times New Roman"
            r.font.size = Pt(7.0)
            if w not in non_italic and not w.startswith("sp"):
                r.italic = True
            else:
                r.italic = False

def build_docx_table(oos_id, sample_mode, out_docx_path):
    src_template = r"scratch\EM table OOS-262133 and OOS-262134 04SEP2026 (1-Page Master).docx"
    doc = docx.Document(src_template)
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.25)
        sec.bottom_margin = Inches(0.25)
        sec.left_margin = Inches(0.35)
        sec.right_margin = Inches(0.35)
        
    # Paragraph titles (NO '&')
    doc.paragraphs[0].text = "Table 1: Read Dates and Incubation Observation"
    doc.paragraphs[2].text = "Table 2: Environmental Monitoring for Analyst and Cleanroom"
    
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(8.0)
            r.bold = True
            
    # Table 1 Setup
    t1 = doc.tables[0]
    t1.rows[0].cells[1].text = "Set-up Analyst and Date"
    t1.rows[0].cells[2].text = "Sampling Site and Location"
    t1.rows[0].cells[3].text = "Plate Reading Analyst (≥ 48H)"
    
    if sample_mode == "262133":
        t1.rows[1].cells[0].text = "ETX-260914-0493\nETX-260914-0496"
        t1.rows[1].cells[1].text = "SMO\n04SEP\n2026"
        t1.rows[1].cells[2].text = "Active air sampling plate"
        t1.rows[1].cells[3].text = "MC"
        t1.rows[1].cells[4].text = "53 CFUs on 143 sec I ISO 8, and 44 CFUs on 143 sec II ISO 8"
        t1.rows[1].cells[5].text = "MC"
        t1.rows[1].cells[6].text = "53 CFUs on 143 sec I ISO 8, and 44 CFUs on 143 sec II ISO 8"
    else:
        t1.rows[1].cells[0].text = "ETX-260914-0498"
        t1.rows[1].cells[1].text = "SMO\n04SEP\n2026"
        t1.rows[1].cells[2].text = "Active air sampling plate"
        t1.rows[1].cells[3].text = "MC"
        t1.rows[1].cells[4].text = "17 CFUs on 142 ISO 8"
        t1.rows[1].cells[5].text = "MC"
        t1.rows[1].cells[6].text = "17 CFUs on 142 ISO 8"
        
    format_microbial_cell(t1.rows[1].cells[7], sp_04sep)
    
    # Table 2 Setup
    t2 = doc.tables[1]
    t2.rows[0].cells[2].text = "Date\n(DDMMMYYYY)"
    t2.rows[1].cells[0].text = "Weekly Active Air Sampling of Cleanroom 145, 144, 143, 142 (L-Suite) with Processing BSCs"
    t2.rows[5].cells[0].text = "Weekly Surface Sampling of Cleanroom 145, 144, 143, 142 (L-Suite) with Processing BSCs"
    
    # Day/Weeks
    t2.rows[2].cells[4].text = "Week (Before Testing Date)"
    t2.rows[3].cells[4].text = "Week (Of Testing Date)"
    t2.rows[4].cells[4].text = "Week (After Testing Date)"
    t2.rows[6].cells[4].text = "Week (Before Testing Date)"
    t2.rows[7].cells[4].text = "Week (Of Testing Date)"
    t2.rows[8].cells[4].text = "Week (After Testing Date)"
    
    # 04SEP active air
    format_microbial_cell(t2.rows[3].cells[7], sp_04sep)
    # 09SEP active air
    format_microbial_cell(t2.rows[4].cells[7], sp_09sep)
    
    # Passing surface observation casing & Not Applicable
    t2.rows[7].cells[5].text = "No Growth"
    t2.rows[7].cells[6].text = "Not Applicable"
    t2.rows[7].cells[7].text = "Not Applicable"
    t2.rows[7].cells[8].text = "None"
    
    t2.rows[8].cells[5].text = "No Growth"
    t2.rows[8].cells[6].text = "Not Applicable"
    t2.rows[8].cells[7].text = "Not Applicable"
    t2.rows[8].cells[8].text = "None"
    
    # Widths & 7.0pt font formatting across all cells
    t1_widths = [Inches(0.95), Inches(0.70), Inches(0.80), Inches(0.60), Inches(1.10), Inches(0.55), Inches(1.10), Inches(1.90)]
    t2_widths = [Inches(1.10), Inches(0.45), Inches(0.55), Inches(0.40), Inches(0.75), Inches(1.20), Inches(1.10), Inches(1.80), Inches(0.35)]
    
    for t_idx, table in enumerate(doc.tables):
        widths = t1_widths if t_idx == 0 else t2_widths
        table.autofit = False
        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
            for c_idx, cell in enumerate(row.cells):
                cell.width = widths[c_idx]
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 0.90
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        if r_idx == 0:
                            r.font.size = Pt(7.0)
                            r.bold = True
                        elif t_idx == 1 and r_idx in [1, 5]:
                            r.font.size = Pt(7.0)
                            r.bold = True
                        else:
                            r.font.size = Pt(7.0)
                            
    doc.save(out_docx_path)
    print(f"Saved: {out_docx_path}")

print("build_docx_table defined.")
