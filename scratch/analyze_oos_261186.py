import os
from pypdf import PdfReader
from docx import Document

pdf_path = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\OOS-261186.pdf"
docx_path = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\EM table OOS-261186 07MAY2026 (2).docx"

print("PDF exists:", os.path.exists(pdf_path))
print("DOCX exists:", os.path.exists(docx_path))

# 1. Analyze PDF
if os.path.exists(pdf_path):
    r = PdfReader(pdf_path)
    print(f"\n=== PDF Analysis ({len(r.pages)} pages) ===")
    fields = r.get_fields()
    print("PDF Fields count:", len(fields) if fields else 0)
    if fields:
        for k, v in sorted(fields.items()):
            val = v.get('/V')
            if val not in [None, '', '/Off']:
                val_str = str(val).replace('\n', ' \\n ')
                if len(val_str) > 100: val_str = val_str[:100] + '...'
                print(f"{k:20s}: {val_str}")
    
    print("\n--- Extracted Text per page ---")
    for idx, page in enumerate(r.pages):
        text = page.extract_text() or ""
        print(f"\n>>> Page {idx+1} ({len(text)} chars):")
        for line in text.splitlines()[:15]:
            print("  ", line)

# 2. Analyze DOCX
if os.path.exists(docx_path):
    doc = Document(docx_path)
    print(f"\n=== DOCX Analysis ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables) ===")
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"P[{idx}]: {p.text}")
    for t_idx, table in enumerate(doc.tables):
        print(f"\n--- Table {t_idx+1} ({len(table.rows)} rows, {len(table.columns)} cols) ---")
        for r_idx, row in enumerate(table.rows):
            cells_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
            # deduplicate merged cells text for readability
            dedup_cells = []
            for c in cells_text:
                if not dedup_cells or c != dedup_cells[-1]:
                    dedup_cells.append(c)
            print(f"  R{r_idx}: {' | '.join(dedup_cells[:8])}")
