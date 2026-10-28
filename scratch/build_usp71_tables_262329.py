import os, sys, shutil, docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

sys.stdout.reconfigure(encoding='utf-8')

DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
SCRATCH_DIR = r"scratch"

OOS_ID = "262329"
CLIENT_NAME = "GoGoMeds Select (E10747)"
TEST_NAME = "USP71"
SAMPLE_ID = "ETX-260624-0251 (BUD-T90)"
SAMPLE_URL = "https://etrax.eagleanalytical.com/Submission/Details/nIsCJb3y7xy7m1FK8luTrg__"
SUBCULTURE_ID = "ETX-261012-0367"
SUBCULTURE_URL = "https://etrax.eagleanalytical.com/Submission/Details/e6Jo3DyLpUEd3QnYm5Uttw__"
EM_PLATE_ID = "ETX-261001-0625"
EM_PLATE_URL = "https://etrax.eagleanalytical.com/Submission/Details/yp3RLi-9BGnGhzxnYnqk4Q__"

ANALYTE = "Sermorelin 0.9mg/ml"
LOT_NUMBER = "17009"
MEDIA = "2 x 200mL TSB (Original and Subculture Bottles)"
MICROBIAL_ID = "Bacillus megaterium"
MICROBIAL_STAIN = "Gram (+) rods"

PRO_DATE = "23Sep26"
PRO_ANALYST_FULL = "Guanchen Li"
PRO_INITIAL = "GL"

READ_ANALYST_FULL = "Andrew Carrillo"
READ_INITIAL = "ANC"

AIR_ANALYST = "SMO"
OTHER_BSC_ANALYST = "AOD"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=25, bottom=25, left=40, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_hyperlink(paragraph, url, text, font_size=Pt(7), color_rgb="0000FF"):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color_rgb}"/>'))
    rPr.append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
    new_run.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def format_cell(cell, text, bold=False, italic=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = v_align
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = font_size
    r.bold = bold
    r.italic = italic
    return r

def build_tables_doc():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Table 1: Information for Sample under investigation
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"Table 1: Information for {SAMPLE_ID} under investigation")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(8.5)
    r_t1.underline = True

    t1 = doc.add_table(rows=2, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = [
        "Processing Analyst",
        "Reading Analyst",
        "Sample ID",
        "Related Microbial ID",
        "Media with microbial growth",
        "Microbial ID"
    ]
    t1_widths = [Inches(1.10), Inches(1.10), Inches(1.25), Inches(1.10), Inches(1.50), Inches(1.45)]

    for j, h in enumerate(t1_headers):
        format_cell(t1.cell(0, j), h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    format_cell(t1.cell(1, 0), PRO_ANALYST_FULL, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 1), READ_ANALYST_FULL, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    # Sample ID Hyperlink
    c_sid = t1.cell(1, 2)
    c_sid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sid)
    p_sid = c_sid.paragraphs[0]
    p_sid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sid.paragraph_format.space_before = Pt(0)
    p_sid.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sid, SAMPLE_URL, SAMPLE_ID, font_size=Pt(7))

    # Subculture ID Hyperlink
    c_sub = t1.cell(1, 3)
    c_sub.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sub)
    p_sub = c_sub.paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sub, SUBCULTURE_URL, SUBCULTURE_ID, font_size=Pt(7))

    format_cell(t1.cell(1, 4), MEDIA, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Microbial ID with Gram stain
    c_mid = t1.cell(1, 5)
    c_mid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_mid)
    p_mid = c_mid.paragraphs[0]
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.paragraph_format.space_before = Pt(0)
    p_mid.paragraph_format.space_after = Pt(0)
    p_mid.paragraph_format.line_spacing = 1.0
    r_m1 = p_mid.add_run(MICROBIAL_ID)
    r_m1.font.name = 'Times New Roman'
    r_m1.font.size = Pt(7)
    r_m1.italic = True
    r_m_br = p_mid.add_run("\n")
    r_m_br.font.name = 'Times New Roman'
    r_m_br.font.size = Pt(7)
    r_m2 = p_mid.add_run(f"({MICROBIAL_STAIN})")
    r_m2.font.name = 'Times New Roman'
    r_m2.font.size = Pt(7)
    r_m2.italic = False

    for row in t1.rows:
        for j, w in enumerate(t1_widths):
            row.cells[j].width = w

    # Spacing between tables
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(4)
    p_space.paragraph_format.space_after = Pt(0)

    # Table 2: Environmental Monitoring
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run(f"Table 2: Environmental Monitoring from Testing Performed on {PRO_DATE}")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(8.5)
    r_t2.underline = True

    t2 = doc.add_table(rows=1, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    t2_headers = [
        "Environmental\nMonitoring (EM)\nSampling Site",
        "Frequency",
        "Date\n(DDMMMYY)",
        "Analyst",
        "Day /Week(s)",
        "Observation",
        "EM Plate ETX Number",
        "Microbial ID",
        "Notes"
    ]
    t2_widths = [
        Inches(1.50), # Site
        Inches(0.55), # Freq
        Inches(0.60), # Date
        Inches(0.50), # Analyst
        Inches(0.85), # Day/Week
        Inches(0.85), # Obs
        Inches(0.85), # ETX
        Inches(1.35), # Micro ID
        Inches(0.45)  # Notes
    ]

    for j, h in enumerate(t2_headers):
        format_cell(t2.cell(0, j), h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    def add_section_divider(table, title):
        row = table.add_row()
        set_cell_background(row.cells[0], "D9D9D9")
        set_cell_margins(row.cells[0], top=20, bottom=20, left=40, right=40)
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(7)
        r.bold = True
        a = row.cells[0]._tc
        for k in range(1, 9):
            row.cells[k]._tc.get_or_add_tcPr().append(parse_xml(f'<w:hMerge {nsdecls("w")} w:val="continue"/>'))
        row.cells[0]._tc.get_or_add_tcPr().append(parse_xml(f'<w:hMerge {nsdecls("w")} w:val="restart"/>'))
        return row

    def add_data_row(table, site, freq, date_str, analyst, day_week, obs, etx, micro_content_fn, notes="None", etx_url=None):
        row = table.add_row()
        format_cell(row.cells[0], site, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], freq, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[2], date_str, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[3], analyst, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[4], day_week, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[5], obs, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        
        c_etx = row.cells[6]
        c_etx.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(c_etx)
        p_etx = c_etx.paragraphs[0]
        p_etx.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_etx.paragraph_format.space_before = Pt(0)
        p_etx.paragraph_format.space_after = Pt(0)
        if etx_url and etx != "Not Applicable":
            add_hyperlink(p_etx, etx_url, etx, font_size=Pt(7))
        else:
            r = p_etx.add_run(etx)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(7)
            
        c_micro = row.cells[7]
        c_micro.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(c_micro)
        p_micro = c_micro.paragraphs[0]
        p_micro.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_micro.paragraph_format.space_before = Pt(0)
        p_micro.paragraph_format.space_after = Pt(0)
        p_micro.paragraph_format.line_spacing = 1.0
        
        if callable(micro_content_fn):
            micro_content_fn(p_micro)
        else:
            r = p_micro.add_run(str(micro_content_fn))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(7)
            
        format_cell(row.cells[8], notes, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        return row

    # --- Section 1: Personnel EM ---
    add_section_divider(t2, "Personnel EM")
    add_data_row(t2, "Personnel (Left Touch and Right Touch)", "Daily", "22Sep26", PRO_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personnel (Left Touch and Right Touch)", "Daily", "23Sep26", PRO_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personnel (Left Touch and Right Touch)", "Daily", "24Sep26", PRO_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t2, "Biological Safety Cabinet EM for E001314")
    add_data_row(t2, "Surface Sampling of ISO 5 (4 locations)", "Daily", "22Sep26", OTHER_BSC_ANALYST, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Surface Sampling of ISO 5 (4 locations)", "Daily", "23Sep26", PRO_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Surface Sampling of ISO 5 (4 locations)", "Daily", "24Sep26", OTHER_BSC_ANALYST, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of ISO 5 (2 locations)", "Daily", "22Sep26", OTHER_BSC_ANALYST, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of ISO 5 (2 locations)", "Daily", "23Sep26", PRO_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of ISO 5 (2 locations)", "Daily", "24Sep26", OTHER_BSC_ANALYST, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 3: Cleanroom Active Air ---
    add_section_divider(t2, "Weekly Active Air Sampling of Cleanroom Suite 115 (E001737) with Processing BSC")
    
    def fill_active_air_micro(p):
        # 1. Micrococcus luteus (4 CFUs)
        r1 = p.add_run("Micrococcus luteus")
        r1.font.name = 'Times New Roman'; r1.font.size = Pt(7); r1.italic = True
        r1_c = p.add_run(" (4 CFUs)\n\n")
        r1_c.font.name = 'Times New Roman'; r1_c.font.size = Pt(7); r1_c.italic = False
        
        # 2. Kocuria indica and Kocuria rhizophila (1 CFU)
        r2a = p.add_run("Kocuria indica")
        r2a.font.name = 'Times New Roman'; r2a.font.size = Pt(7); r2a.italic = True
        r2_and = p.add_run(" and ")
        r2_and.font.name = 'Times New Roman'; r2_and.font.size = Pt(7); r2_and.italic = False
        r2b = p.add_run("Kocuria rhizophila")
        r2b.font.name = 'Times New Roman'; r2b.font.size = Pt(7); r2b.italic = True
        r2_c = p.add_run(" (1 CFU)\n\n")
        r2_c.font.name = 'Times New Roman'; r2_c.font.size = Pt(7); r2_c.italic = False
        
        # 3. Pending FunITS (1 CFU, Hyphae)
        r3 = p.add_run("Pending FunITS (1 CFU, Hyphae)")
        r3.font.name = 'Times New Roman'; r3.font.size = Pt(7); r3.italic = False

    add_data_row(t2, "Active Air Sampling of Cleanrooms", "Weekly", "23Sep26", AIR_ANALYST, "Week of Testing", "6 CFUs (ISO 8 115)", EM_PLATE_ID, fill_active_air_micro, "None", etx_url=EM_PLATE_URL)

    # --- Section 4: Cleanroom Surface ---
    add_section_divider(t2, "Weekly Surface Sampling of Cleanroom Suite 115 (E001737) with Processing BSC")
    add_data_row(t2, "Surface Sampling of Cleanrooms", "Weekly", "23Sep26", AIR_ANALYST, "Week of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    for row in t2.rows:
        for j, w in enumerate(t2_widths):
            row.cells[j].width = w

    # Uniform XML 14 half-points for all hyperlinks and text runs
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.size = Pt(7) if r.font.size != Pt(8.5) else Pt(8.5)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(7)

    # Save documents
    filename = f"Tables OOS-{OOS_ID} {CLIENT_NAME} - {TEST_NAME}.docx"
    doc_path_local = os.path.join(DOCUMENTS_DIR, filename)
    doc_path_desktop = os.path.join(DESKTOP_DIR, filename)

    doc.save(doc_path_local)
    shutil.copy2(doc_path_local, doc_path_desktop)
    print("Saved Word table to Documents and Desktop.")

    # Convert to PDF via Word COM to check exact page count and verify visual appearance
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    pdf_temp_path = os.path.join(SCRATCH_DIR, "verify_tables_262329.pdf")
    try:
        wdoc = word.Documents.Open(os.path.abspath(doc_path_local))
        wdoc.SaveAs(os.path.abspath(pdf_temp_path), FileFormat=17) # wdFormatPDF
        wdoc.Close()
    finally:
        word.Quit()

    pdf_doc = fitz.open(pdf_temp_path)
    page_count = len(pdf_doc)
    print(f"Standalone Tables PDF generated: {page_count} page(s).")
    assert page_count == 1, f"CRITICAL ERROR: Tables document spilled over to {page_count} pages! Must be strictly 1 page."
    
    # Save preview image
    pix = pdf_doc[0].get_pixmap(dpi=150)
    pix_path = os.path.join(SCRATCH_DIR, "preview_tables_262329_p1.png")
    pix.save(pix_path)
    print(f"Rendered Page 1 visual inspection image: {pix_path}")
    pdf_doc.close()

if __name__ == "__main__":
    build_tables_doc()
