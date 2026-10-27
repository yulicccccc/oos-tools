import os
import fitz
import docx
from pypdf import PdfReader, PdfWriter

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docs_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"

target_pdf = os.path.join(desktop_dir, "OOS-262289 - ETX-260929-0931 - Link Pharmacy (E19112) - ScanRDI  - QYC -RS.pdf")
print("Target PDF:", target_pdf)

# 1. Update target PDF
doc = fitz.open(target_pdf)
updated_49 = False
updated_50 = False

for page in doc:
    for w in page.widgets():
        if w.field_name == "Text Field49":
            val = w.field_value
            if "15 total events" in val:
                val = val.replace("15 total events", "19 total events")
                w.field_value = val
                w.text_fontsize = 8.2
                w.text_font = "Helv"
                w.update()
                updated_49 = True
                print("Updated Text Field49: 15 -> 19 total events")
            elif "19 total events" in val:
                print("Text Field49 already has 19 total events")
                updated_49 = True
        elif w.field_name == "Text Field50":
            val = w.field_value
            if "BSC E001938 (Cleanroom 144, L-Suite)" in val:
                val = val.replace("BSC E001938 (Cleanroom 144, L-Suite)", "BSC E001938 (Cleanroom 145, L-Suite)")
                w.field_value = val
                w.text_fontsize = 7.3
                w.text_font = "Helv"
                w.update()
                updated_50 = True
                print("Updated Text Field50: Cleanroom 144 -> Cleanroom 145")
            elif "BSC E001938 (Cleanroom 145, L-Suite)" in val:
                print("Text Field50 already has Cleanroom 145")
                updated_50 = True
        elif w.field_name == "Text Field43":
            val = w.field_value
            if "(Sensor E001505) \r \r" in val:
                w.field_value = val.replace("(Sensor E001505) \r \r", "(Sensor E001505)\r \r")
                w.update()

assert updated_49, "Failed to update Text Field49"
assert updated_50, "Failed to update Text Field50"

# Save updated target PDF
temp_save = "scratch/temp_target_updated.pdf"
doc.save(temp_save)
doc.close()

import shutil
shutil.copy2(temp_save, target_pdf)
print("Successfully saved updated target PDF:", target_pdf)

# Also copy to OOS-262289 Link Pharmacy (E19112) - ScanRDI.pdf on Desktop & Documents
shutil.copy2(target_pdf, os.path.join(desktop_dir, "OOS-262289 Link Pharmacy (E19112) - ScanRDI.pdf"))
shutil.copy2(target_pdf, os.path.join(docs_dir, "OOS-262289 Link Pharmacy (E19112) - ScanRDI.pdf"))
shutil.copy2(target_pdf, os.path.join(desktop_dir, "CORP-FORM-21- 01 Oct 2026 OOS-262289 Link Pharmacy (E19112) - ScanRDI.pdf"))
shutil.copy2(target_pdf, os.path.join(docs_dir, "CORP-FORM-21- 01 Oct 2026 OOS-262289 Link Pharmacy (E19112) - ScanRDI.pdf"))
print("Synchronized all master deliverable PDFs on Desktop and Documents.")

# 2. Update Word document Phase I summary block
for d in [desktop_dir, docs_dir]:
    w_path = os.path.join(d, "OOS-262289 Link Pharmacy (E19112) - ScanRDI.docx")
    if os.path.exists(w_path):
        w_doc = docx.Document(w_path)
        modified = False
        for t in w_doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    if "15 total events" in cell.text:
                        for p in cell.paragraphs:
                            if "15 total events" in p.text:
                                p.text = p.text.replace("15 total events", "19 total events")
                                modified = True
                    if "BSC E001938 (Cleanroom 144, L-Suite)" in cell.text:
                        for p in cell.paragraphs:
                            if "BSC E001938 (Cleanroom 144, L-Suite)" in p.text:
                                p.text = p.text.replace("BSC E001938 (Cleanroom 144, L-Suite)", "BSC E001938 (Cleanroom 145, L-Suite)")
                                modified = True
        if modified:
            w_doc.save(w_path)
            print("Successfully updated Word doc:", w_path)
