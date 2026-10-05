import fitz
import docx
import os
import shutil
import gc
import win32com.client

SCRATCH_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\scratch"
REPO_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"

print("=== STEP 1: UPDATE Celsis table OOS-262080.docx (Table 1: 2 x 300mL FTM) ===")
tbl_docx_desktop = os.path.join(DESKTOP_DIR, "Celsis table OOS-262080.docx")
tbl_docx_scratch = os.path.join(SCRATCH_DIR, "Celsis table OOS-262080.docx")

def update_table1_docx(path):
    d = docx.Document(path)
    t1 = d.tables[0]
    c4 = t1.rows[1].cells[4]
    print(f"  Old Table 1 Media cell: {repr(c4.text)}")
    c4.paragraphs[0].text = "2 x 300mL FTM"
    if len(c4.paragraphs[0].runs) > 0:
        c4.paragraphs[0].runs[0].font.size = docx.shared.Pt(7)
    d.save(path)
    print(f"  Updated Table 1 Media cell in {path} to: 2 x 300mL FTM")

update_table1_docx(tbl_docx_desktop)
shutil.copy2(tbl_docx_desktop, tbl_docx_scratch)

print("\n=== STEP 2: EXPORT TABLES DOCX TO PDF VIA WORD COM ===")
tbl_pdf_desktop = os.path.join(DESKTOP_DIR, "Celsis table OOS-262080.pdf")
tbl_pdf_scratch = os.path.join(SCRATCH_DIR, "Celsis table OOS-262080.pdf")

word = None
try:
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_word = word.Documents.Open(tbl_docx_desktop)
    doc_word.ExportAsFixedFormat(tbl_pdf_scratch, 17) # wdExportFormatPDF = 17
    doc_word.Close(False)
    print("  Word COM export successful to scratch.")
    try:
        shutil.copy2(tbl_pdf_scratch, tbl_pdf_desktop)
        print("  Copied updated tables PDF to Desktop.")
    except Exception as e:
        print(f"  Notice: Desktop tables PDF locked: {e}")
except Exception as e:
    print(f"  Word COM error: {e}")
finally:
    if word:
        try: word.Quit()
        except: pass
        del word
        gc.collect()

print("\n=== STEP 3: UPDATE MASTER REPORT DOCX ===")
def update_master_docx(path):
    d = docx.Document(path)
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    txt = p.text
                    if "one of the 300 mL FTM media jars" in txt:
                        p.text = txt.replace(
                            "one of the 300 mL FTM media jars",
                            "two of the 300 mL FTM media jars (the 3rd and 5th jars)"
                        )
                    if "originating from the FTM sample bottle" in p.text:
                        p.text = p.text.replace(
                            "originating from the FTM sample bottle",
                            "originating from the positive FTM sample bottles"
                        )
                    if "positive FTM bottle for ETX-260828-0527 was submitted" in p.text:
                        p.text = p.text.replace(
                            "positive FTM bottle for ETX-260828-0527 was submitted",
                            "positive FTM bottles for ETX-260828-0527 were submitted"
                        )
                    if "All other samples processed by the same analyst on the day of testing were found to test negative" in p.text:
                        if "It was the 1st sample" not in p.text:
                            p.text = p.text.replace(
                                "All other samples processed by the same analyst on the day of testing were found to test negative",
                                "It was the 1st sample processed from the batch of samples (with the 3rd and 5th jars found positive). All other samples processed by the same analyst on the day of testing were found to test negative"
                            )
    d.save(path)
    print(f"  Updated master report docx: {path}")

master_docx_desktop = os.path.join(DESKTOP_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis.docx")
master_docx_scratch = os.path.join(SCRATCH_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis.docx")
if os.path.exists(master_docx_desktop):
    update_master_docx(master_docx_desktop)
if os.path.exists(master_docx_scratch):
    update_master_docx(master_docx_scratch)

print("\n=== STEP 4: UPDATE CORP-FORM-21 PDF ===")
src_clean_pdf = os.path.join(SCRATCH_DIR, "CORP-FORM-21_clean.pdf")
doc_corp = fitz.open(src_clean_pdf)

# Update Page 4 (Text Field50)
p4 = doc_corp[3]
for w in p4.widgets():
    if w.field_name == 'Text Field50':
        val = w.field_value
        val = val.replace("one of the 300 mL FTM media jars", "two of the 300 mL FTM media jars (the 3rd and 5th jars)")
        val = val.replace("originating from the FTM sample bottle", "originating from the positive FTM sample bottles")
        val = val.replace("positive FTM bottle for ETX-260828-0527 was submitted", "positive FTM bottles for ETX-260828-0527 were submitted")
        w.field_value = val
        w.update()
        print("  Updated Page 4 Text Field50 (2 FTM jars).")

# Update Page 5 (Text Field51)
p5 = doc_corp[4]
for w in p5.widgets():
    if w.field_name == 'Text Field51':
        val = w.field_value
        target_sub = "All other samples processed by the same analyst on the day of testing were found to test negative"
        if target_sub in val and "It was the 1st sample" not in val:
            val = val.replace(
                target_sub,
                "It was the 1st sample processed from the batch of samples (with the 3rd and 5th jars found positive). All other samples processed by the same analyst on the day of testing were found to test negative"
            )
            w.field_value = val
            w.update()
            print("  Updated Page 5 Text Field51 (batch processing sequence).")

out_clean_v3 = os.path.join(SCRATCH_DIR, "CORP-FORM-21_clean_v3.pdf")
doc_corp.save(out_clean_v3)
doc_corp.close()

shutil.copy2(out_clean_v3, src_clean_pdf)

# Copy to repo and desktop
for target in [
    os.path.join(REPO_DIR, "CORP-FORM-21 Laboratory OOS Investigation Form (v11.1).pdf"),
    os.path.join(DESKTOP_DIR, "CORP-FORM-21 Laboratory OOS Investigation Form (v11.1).pdf")
]:
    try:
        shutil.copy2(src_clean_pdf, target)
        print(f"  Deployed 6-page form to: {target}")
    except Exception as e:
        print(f"  Notice: copy to {target} locked: {e}")

print("\n=== STEP 5: ASSEMBLE 8-PAGE COMPLETE PACKET ===")
doc_complete = fitz.open()
doc_6 = fitz.open(src_clean_pdf)
doc_complete.insert_pdf(doc_6, links=True)
doc_6.close()

source_tables = tbl_pdf_scratch if os.path.exists(tbl_pdf_scratch) else tbl_pdf_desktop
doc_tbl = fitz.open(source_tables)
doc_complete.insert_pdf(doc_tbl, links=True)
doc_tbl.close()

out_scratch_complete = os.path.join(SCRATCH_DIR, "OOS-262080_complete_final_clean.pdf")
doc_complete.save(out_scratch_complete)
print(f"  Saved 8-page packet to {out_scratch_complete} (pages: {len(doc_complete)})")

for name in [
    "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (Complete).pdf",
    "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (2).pdf",
    "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis - QYC (Updated).pdf",
    "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis - QYC (Latest).pdf",
    "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis - QYC.pdf"
]:
    target = os.path.join(DESKTOP_DIR, name)
    try:
        shutil.copy2(out_scratch_complete, target)
        print(f"  Deployed 8-page packet to: {target}")
    except Exception as e:
        print(f"  Notice: {name} locked by user: {e}")

print("\n=== DONE: All files updated for 2 positive FTM jars ===")
