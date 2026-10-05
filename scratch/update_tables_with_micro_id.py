import os
import sys
import docx
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import fitz

sys.stdout.reconfigure(encoding='utf-8')

DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
DOCS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
SCRATCH_DIR = os.path.join(DOCS_DIR, "scratch")

tables_docx_path = os.path.join(DESKTOP_DIR, "Celsis table OOS-262080.docx")
tables_pdf_path = os.path.join(DESKTOP_DIR, "Celsis table OOS-262080.pdf")
main_docx_path = os.path.join(DESKTOP_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis.docx")

URL_0520 = "https://etrax.eagleanalytical.com/SubmissionTest/Details/qnwcLQO5BWeJhWCBG7jl8Q__"
SAMPLE_0520 = "ETX-260921-0520"

print("=== STEP 1: UPDATE Celsis table OOS-262080.docx ===")
doc_tbl = docx.Document(tables_docx_path)
t3 = doc_tbl.tables[2]

# Find row for 0520
target_row = None
for r in t3.rows:
    text = ' '.join(c.text for c in r.cells)
    if SAMPLE_0520 in text:
        target_row = r
        break

assert target_row is not None, "0520 row not found in Table 3!"

# Cell 8 (or 9) is ETX ID, Cell 10 (or 11) is Microbial ID
# Let's inspect cells in target_row
print("Target row cells:")
for i, c in enumerate(target_row.cells):
    print(f"  [{i}]: {repr(c.text.strip())}")

# Update ETX ID cell with clickable link
for i, c in enumerate(target_row.cells):
    if SAMPLE_0520 in c.text:
        part = c.part
        r_id = part.relate_to(URL_0520, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        tc_xml = (
            f'<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<w:tcPr>'
            f'<w:tcW w:w="1444" w:type="dxa"/>'
            f'<w:tcBorders>'
            f'<w:top w:val="single" w:sz="8" w:space="0" w:color="auto"/>'
            f'<w:left w:val="single" w:sz="8" w:space="0" w:color="auto"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="auto"/>'
            f'<w:right w:val="single" w:sz="8" w:space="0" w:color="auto"/>'
            f'</w:tcBorders>'
            f'<w:vAlign w:val="center"/>'
            f'</w:tcPr>'
            f'<w:p>'
            f'<w:pPr>'
            f'<w:spacing w:line="360" w:lineRule="auto"/>'
            f'<w:jc w:val="center"/>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:sz w:val="14"/>'
            f'<w:szCs w:val="14"/>'
            f'</w:rPr>'
            f'</w:pPr>'
            f'<w:hyperlink r:id="{r_id}" w:history="1">'
            f'<w:r>'
            f'<w:rPr>'
            f'<w:rStyle w:val="Hyperlink"/>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:color w:val="0000FF"/>'
            f'<w:sz w:val="14"/>'
            f'<w:szCs w:val="14"/>'
            f'<w:u w:val="single"/>'
            f'</w:rPr>'
            f'<w:t>{SAMPLE_0520}</w:t>'
            f'</w:r>'
            f'</w:hyperlink>'
            f'</w:p>'
            f'</w:tc>'
        )
        new_tc = parse_xml(tc_xml)
        c._tc.getparent().replace(c._tc, new_tc)
        print(f"  Inserted hyperlink for {SAMPLE_0520} at cell [{i}]")

# Update Microbial ID cell to 'Micrococcus luteus'
for i, c in enumerate(target_row.cells):
    if 'Gram (+) cocci' in c.text:
        # Update text to Micrococcus luteus
        for p in c.paragraphs:
            p.text = ""
        run = c.paragraphs[0].add_run("Micrococcus luteus")
        run.font.name = "Times New Roman"
        run.font.size = docx.shared.Pt(7)
        c.paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        print(f"  Updated Microbial ID to 'Micrococcus luteus' at cell [{i}]")

scratch_tables_docx = os.path.join(SCRATCH_DIR, "Celsis table OOS-262080.docx")
doc_tbl.save(scratch_tables_docx)
doc_tbl.save(tables_docx_path)
print(f"Saved updated tables DOCX to: {tables_docx_path}")

print("\n=== STEP 2: CONVERT TABLES DOCX TO PDF VIA WORD COM ===")
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = 0

doc_com = word.Documents.Open(tables_docx_path)
scratch_tables_pdf = os.path.join(SCRATCH_DIR, "Celsis table OOS-262080.pdf")
doc_com.SaveAs2(scratch_tables_pdf, FileFormat=17)
doc_com.Close()
word.Quit()

import shutil
shutil.copy2(scratch_tables_pdf, tables_pdf_path)
print(f"Converted and deployed tables PDF to: {tables_pdf_path}")

print("\n=== STEP 3: VERIFY NEW PDF LINKS & TEXT ===")
doc_pdf = fitz.open(tables_pdf_path)
print(f"Pages: {len(doc_pdf)}")
for i, p in enumerate(doc_pdf):
    print(f"Page {i+1} links: {p.get_links()}")
    t = p.get_text()
    if 'Micrococcus' in t:
        print(f"  Found 'Micrococcus' on Page {i+1}!")
doc_pdf.close()
