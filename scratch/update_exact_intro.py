import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client as win32

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_path = os.path.join(desktop_dir, "OOS-261242 Review Email.docx")
html_path = os.path.join(desktop_dir, "OOS-261242_Review_Email_Robin_Format.html")

new_intro = "I have reviewed OOS-261242. Please find the following edits made in the summary table below:"

# 1. Update DOCX
doc = Document(docx_path)
for p in doc.paragraphs:
    if "I have reviewed" in p.text:
        p.text = new_intro
        p.runs[0].font.name = "Calibri"
        from docx.shared import Pt
        p.runs[0].font.size = Pt(11)

tbl = doc.tables[0]
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
tblPr = tbl._element.xpath('w:tblPr')
if tblPr:
    for jc in tblPr[0].xpath('w:jc'):
        tblPr[0].remove(jc)
    tblPr[0].append(parse_xml(f'<w:jc {nsdecls("w")} w:val="left"/>'))
    
    for ind in tblPr[0].xpath('w:tblInd'):
        tblPr[0].remove(ind)
    tblPr[0].append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>'))

doc.save(docx_path)
print("Updated DOCX intro to exact user-requested sentence.")

# 2. Update HTML
if os.path.exists(html_path):
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()
    # Replace paragraph
    import re
    html = re.sub(r'<p style="margin: 0 0 12px 0;[^>]*>I have reviewed[^<]*</p>',
                  f'<p style="margin: 0 0 12px 0; font-family: Calibri, sans-serif; font-size: 11pt;">{new_intro}</p>',
                  html)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    print("Updated HTML intro.")

# 3. Copy via Word COM to clipboard
word_app = win32.Dispatch("Word.Application")
word_app.Visible = False
try:
    d = word_app.Documents.Open(os.path.abspath(docx_path))
    d.Content.Copy()
    d.Close(False)
    print("SUCCESS: Copied refined email to Windows Clipboard via Word COM!")
finally:
    word_app.Quit()
