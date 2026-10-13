import docx

doc = docx.Document('scratch/user_latest_tables_262237.docx')
t = doc.tables[1]
print(f"Table 2 total rows: {len(t.rows)}")
for i, r in enumerate(t.rows):
    txt0 = r.cells[0].text.replace('\n', ' ')
    txt5 = r.cells[5].text.replace('\n', ' ')
    print(f"Row {i}: {txt0[:40]} | Obs: {txt5[:40]}")

print("\n--- Inspecting Row 9 (or L-Suite Active Air row) ---")
# find row with L-suite hits
for i, r in enumerate(t.rows):
    if "CFUs" in r.cells[5].text and "143" in r.cells[5].text:
        target_row = r
        print(f"Found target row at index {i}")
        for j, cell in enumerate(target_row.cells):
            print(f"\n=== Col {j} ===")
            for p in cell.paragraphs:
                xml_hyper = 'w:hyperlink' in p._p.xml
                print(f"  P: text='{p.text}' | hyper={xml_hyper}")
                if xml_hyper:
                    for child in p._p:
                        if 'hyperlink' in child.tag:
                            print(f"    Hyperlink XML: {child.xml}")
