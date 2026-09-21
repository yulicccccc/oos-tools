import os
from pypdf import PdfReader
from docx import Document

pdf_path = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\OOS-261186.pdf"
docx_path = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\EM table OOS-261186 07MAY2026 (2).docx"

out_txt = r"scratch\oos_261186_clean_summary.txt"

with open(out_txt, "w", encoding="utf-8") as out:
    out.write("=====================================================\n")
    out.write("OOS-261186 PDF CLEAN FIELDS DUMP\n")
    out.write("=====================================================\n\n")

    r = PdfReader(pdf_path)
    fields = r.get_fields() or {}
    for k, v in sorted(fields.items()):
        ft = v.get('/FT')
        val = v.get('/V')
        if ft == '/Sig':
            continue
        if val not in [None, '', '/Off']:
            out.write(f"[{k}] ({ft}):\n{val}\n\n")

    out.write("\n=====================================================\n")
    out.write("DOCX CONTENT DUMP\n")
    out.write("=====================================================\n\n")
    doc = Document(docx_path)
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            out.write(f"P[{p_idx}]: {p.text}\n")
    
    for t_idx, table in enumerate(doc.tables):
        out.write(f"\n--- TABLE {t_idx+1} ({len(table.rows)} rows, {len(table.columns)} cols) ---\n")
        for r_idx, row in enumerate(table.rows):
            row_vals = [c.text.strip().replace('\n', ' ') for c in row.cells]
            d_vals = []
            for c in row_vals:
                if not d_vals or c != d_vals[-1]:
                    d_vals.append(c)
            out.write(f"Row {r_idx:2d}: {' | '.join(d_vals)}\n")

print("Clean dump saved to:", out_txt)
