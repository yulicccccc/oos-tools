import docx
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE

def set_cell_hyperlink(cell, url, text):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    part = p.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink_xml = (
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rStyle w:val="Hyperlink"/>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="467886"/>'
        f'<w:sz w:val="14"/>'
        f'<w:szCs w:val="14"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    p._p.append(parse_xml(hyperlink_xml))

doc = docx.Document('tables for scan.docx')
t = doc.tables[0]
set_cell_hyperlink(t.rows[1].cells[2], 'https://etrax.eagleanalytical.com/Submission/Details/test', 'ETX-TEST')
doc.save('scratch/test_link.docx')
print('Successfully tested set_cell_hyperlink!')
