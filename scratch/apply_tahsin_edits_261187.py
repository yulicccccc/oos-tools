import os, shutil
import fitz
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader, PdfWriter
import win32com.client as win32

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docs_dir = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
history_dir = os.path.join(docs_dir, ".history")

src_pdf_name = "OOS-261187 ScanC_O CGS E001309 S1 11MAY2026.pdf"
src_pdf_path = os.path.join(desktop_dir, src_pdf_name)

print("=== 1. BACKUP CURRENT FILE ===")
shutil.copy2(src_pdf_path, os.path.join(history_dir, "OOS_261187_before_tahsin_edits.pdf"))

# === 2. UPDATE FORM FIELDS (PAGES 1-6) ===
doc_pdf = fitz.open(src_pdf_path)

# Page 1: Text Field7
for w in doc_pdf[0].widgets():
    if w.field_name == "Text Field7":
        w.field_value = "The CFU count for the environmental monitoring plate met the action level."
        w.update()
        print("Updated P1 Text Field7: 'met the action level'")

# Page 5: Text Field51
p5_text = (
    "Based on the available evidence, one colony-like artifact was observed on ISO 5 BSC E001309 "
    "Surface Sampling Plate# 1 on the date of testing (11 May 2026). However, the observation could not be "
    "confirmed as a viable microbial CFU because no growth was obtained following inoculation onto fresh media. "
    "Therefore, the observation represents a nonviable or non-microbial artifact.\r \r"
    "Even if the observed artifact had represented a viable CFU, it would appear to be an isolated event. "
    "This assessment is supported by the absence of microbial recovery from Clea S. Garza's personnel-monitoring "
    "plates on the date before testing (08 May 2026), the date of testing (11 May 2026), and the date after testing "
    "(12 May 2026); the absence of growth from BSC E001309 surface samples collected on the date before testing "
    "(09 May 2026) and the date after testing (12 May 2026); and the absence of growth from ISO 5 settling plates "
    "throughout the bracketing period. Additionally, it is important to note that no growth was observed on the other "
    "three surface sampling plates from the date of testing.\r \r"
    "The negative ISO 5 settling, personnel, and bracketing surface monitoring results demonstrate that the "
    "BSC E001309 critical environment remained in a state of control. The available data do not support migration, "
    "persistence, or recurrence of contamination within ISO 5 BSC E001309. Available data indicate that the analyst "
    "adhered to all approved aseptic protocols with no documented deviations, and no specific or assignable laboratory "
    "discrepancy was identified. Collectively, the evidence supports that the recovery was an isolated, transient, "
    "and non-recurring event.\r \r"
    "Accordingly, no systemic environmental control deficiencies were identified, and no additional corrective or "
    "preventive actions are warranted at this time beyond continued routine environmental monitoring and adherence to "
    "approved cleaning, disinfection, and aseptic procedures."
)

for w in doc_pdf[4].widgets():
    if w.field_name == "Text Field51":
        w.field_value = p5_text
        w.update()
        print("Updated P5 Text Field51: removed speculative laboratory error statement")

temp_form_pdf = os.path.join(docs_dir, "scratch", "temp_261187_form_tahsin.pdf")
doc_pdf.save(temp_form_pdf)
doc_pdf.close()

# === 3. UPDATE TABLE 1 SETUP DATE ON PAGE 7 ===
# Let's inspect how Page 7 was built.
# Let's check EM table OOS-261187 11MAY2026.docx on Desktop
docx_tbl_path = os.path.join(desktop_dir, "EM table OOS-261187 11MAY2026.docx")
if os.path.exists(docx_tbl_path):
    doc_t = Document(docx_tbl_path)
    t1 = doc_t.tables[0]
    # Check row 1 cell 1
    c1 = t1.rows[1].cells[1]
    print("Original Table 1 cell text:", repr(c1.text))
    c1.text = "CGS\n11MAY 2026"
    for p in c1.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(7.0)
    
    # Save updated docx
    doc_t.save(docx_tbl_path)
    print("Updated Table 1 DOCX setup date to CGS 11MAY 2026")
    
    # Export to PDF via Word COM
    out_table_pdf = os.path.join(desktop_dir, "EM table OOS-261187 11MAY2026.pdf")
    word_app = win32.Dispatch("Word.Application")
    word_app.Visible = False
    try:
        d_com = word_app.Documents.Open(os.path.abspath(docx_tbl_path))
        d_com.SaveAs(os.path.abspath(out_table_pdf), FileFormat=17)
        d_com.Close(False)
        print("Exported updated Table PDF via Word COM.")
    finally:
        word_app.Quit()
        
    # Reassemble full 7-page PDF (Pages 1-6 form + Page 7 Table)
    reader_form = PdfReader(temp_form_pdf)
    reader_table = PdfReader(out_table_pdf)
    
    writer = PdfWriter()
    for p in reader_form.pages[:6]:
        writer.add_page(p)
    for p in reader_table.pages:
        writer.add_page(p)
        
    target_out_pdf = os.path.join(desktop_dir, "OOS-261187 ScanC_O CGS E001309 S1 11MAY2026.pdf")
    with open(target_out_pdf, "wb") as f:
        writer.write(f)
    print("SUCCESS: Reassembled finalized 7-page PDF with all Tahsin edits:", target_out_pdf)

    # Verify Page count
    doc_v = fitz.open(target_out_pdf)
    print(f"Final PDF Page Count: {len(doc_v)}")
    p7_t = doc_v[6].get_text().encode('ascii', 'replace').decode('ascii')
    print("Page 7 sample:", p7_t[:120])
