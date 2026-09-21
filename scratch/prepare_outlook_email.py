import os
import win32clipboard
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_out = os.path.join(desktop_dir, "OOS-261186 Review Email.docx")
html_out = os.path.join(desktop_dir, "OOS-261186_Review_Email_Robin_Format.html")

# --- 1. GENERATE WORD DOCX (Best for copy-pasting to Outlook) ---
doc = Document()

# Set page margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Greeting
p1 = doc.add_paragraph()
r1 = p1.add_run("Good morning @Simin Mohammad,")
r1.font.name = "Calibri"
r1.font.size = Pt(11)

p2 = doc.add_paragraph()
r2 = p2.add_run("I have reviewed this OOS. Please find the following edits made in the summary table below. Please note that the ")
r2.font.name = "Calibri"
r2.font.size = Pt(11)
r2_b = p2.add_run("EM table will also need to be amended to ensure uniformity, (see highlighted) sections and reflect updated gram stain results.")
r2_b.font.name = "Calibri"
r2_b.font.size = Pt(11)
r2_b.bold = True

# Helper functions for Word table
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_borders(cell):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcBorders'):
            tcPr.remove(child)
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

table_data = [
    # Headers
    ["Section", "Unreviewed", "Reviewed", "Impact"],
    
    # Row 1
    [
        "EM Table Attachment & Formatting",
        "The EM Table was omitted from the PDF submission. In the draft Word document:\n• Spacing error in Table 1 date (07MAY   2026)\n• Missing Gram stain in Microbial ID (Kocuria palustris)\n• Table 2 title non-standard\n• Multi-organism entries running together without line breaks",
        "Amended and attached as Page 8 (following Page 7 Version History) to create the complete 8-page package:\n• Revised date to 07MAY 2026\n• Revised Microbial ID to Kocuria palustris, Gram (+) cocci\n• Revised title to Table 2: Environmental Monitoring Plates for Analyst and Cleanroom Bracketing\n• Formatted multi-organism entries with clean line breaks",
        "Document Assembly & Technical Uniformity Correction.\nCompleted on analyst's behalf."
    ],
    
    # Row 2
    [
        "Name of Analyst who Performed the Test",
        "Listed as:\nSimin Mohammad(Weekly Active Air Sampling Plate Setup)Sophia Santamaria (Weekly Active Air Sampling Plate Readers)",
        "Revised to:\nSimin Mohammad\n(Weekly Active Air Sampling Analyst)\n\nSophia Santamaria\n(Weekly Active Air Sampling Plate Reader)",
        "Formatting & Grammatical Correction."
    ],
    
    # Row 3
    [
        "SOP Reference & Effective Date",
        "Listed as MICRO-SOP-2 Rev 15, Effective Date 05-AUG-2026.",
        "Revised to MICRO-SOP-2 Rev 16, Effective Date 23-Jul-2026.",
        "Compliance Correction."
    ],
    
    # Row 4
    [
        "Limits / Specification",
        "Action level: >10",
        "Revised to Action level: >= 10 CFU/m³ (or Action level: >= 10 CFU/Plate)",
        "Technical Specification Correction."
    ],
    
    # Row 5
    [
        "Analyst interviewed? (Comments)",
        '"Yes, analyst Simin Mohammad, and Sophia Santamaria were comprehensively interviewed."',
        '"Yes, analysts Simin Mohammad and Sophia Santamaria were interviewed comprehensively."',
        "Grammatical Improvement."
    ],
    
    # Row 6
    [
        "Equipment & Calibration Details",
        "Equipment IDs and calibration dates running together without line breaks:\nIncubator E001034 (Sensor E001501)Incubator E001031 (Sensor E001505)\nAug-2026Feb-2027Aug-2026Feb-2027",
        "Revised with clean line breaks:\nIncubator E001034 (Sensor E001501)\nIncubator E001031 (Sensor E001505)\n\nAug 2026 / Feb 2027",
        "Critical Readability & Formatting Correction."
    ],
    
    # Row 7
    [
        "Phase I Summary – Plate Setup & Incubation Narrative",
        "Multiple missing spaces after periods (document.The, system.Active, Facility.The), missing period after Gram stain (cocci)To observe), unnecessary comma (clean room, were), double hyphen (MICRO-SOP--9), and typo in reader's name (Sophia Sanatamaria).",
        "Revised all punctuation, inserted required spaces, corrected spelling to Sophia Santamaria, removed double hyphen, and refined sentence flow.",
        "Minor Grammatical & Typographical Correction."
    ],
    
    # Row 8
    [
        "Phase I Summary – Cleanroom Suite Reference",
        '"It is also important to note that no samples processed in suite 115 for the week of testing (03-May-2026 to 09-May-2026) failed testing."',
        '"It is also important to note that no samples processed in Suite 116 for the week of testing (03-May-2026 to 09-May-2026) failed testing."',
        "Critical Assessment Correction."
    ],
    
    # Row 9
    [
        "Phase I Summary – Root Cause Statement",
        '"Based on the findings outlined in the preceding sections, the Out-Of-Specification (OOS) result observed for the Environmental Monitoring (EM) Active Air Sampling plate may be attributed to a potential analyst error."',
        "Revised to remove speculative analyst error: Available data indicate that the analyst adhered to all aseptic protocols with no documented deviations. The recovery represents an isolated, transient environmental event, and laboratory error is highly unlikely.",
        "Regulatory & cGMP Compliance Correction."
    ]
]

col_widths = [1.5, 2.3, 2.5, 1.4]  # Total 7.7 inches
table = doc.add_table(rows=len(table_data), cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

for r_idx, row_content in enumerate(table_data):
    row = table.rows[r_idx]
    is_header = (r_idx == 0)
    
    for c_idx, cell_text in enumerate(row_content):
        cell = row.cells[c_idx]
        cell.width = Inches(col_widths[c_idx])
        set_cell_borders(cell)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        
        if is_header:
            set_cell_background(cell, "FFFF00")  # Yellow Header
        else:
            set_cell_background(cell, "FFFFFF")
            
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        if is_header:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(cell_text)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # Format text cleanly without asterisks
            lines = cell_text.split('\n')
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                
                # Highlight key phrases with real bolding, NO asterisks
                if c_idx == 0:
                    r = p.add_run(line)
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.bold = True
                elif c_idx == 2 and (line.startswith("•") or line.startswith('"') or "Revised to" in line or "Amended and attached" in line):
                    r = p.add_run(line)
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.bold = True
                else:
                    r = p.add_run(line)
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)

p3 = doc.add_paragraph()
p3.paragraph_format.space_before = Pt(12)
r3 = p3.add_run("@Simin Mohammad please review and update the OOS form text fields accordingly. You may directly use the newly compiled 8-page PDF package attached here.")
r3.font.name = "Calibri"
r3.font.size = Pt(11)

p4 = doc.add_paragraph()
r4 = p4.add_run("Thanks,")
r4.font.name = "Calibri"
r4.font.size = Pt(11)

doc.save(docx_out)
print(f"Saved native Word table to: {docx_out}")

# --- 2. GENERATE CLEAN HTML (ZERO ASTERISKS) ---
html_rows = ""
for r_idx, row in enumerate(table_data[1:]):
    sec, unrev, rev, imp = row
    
    # Clean up formatting for HTML
    unrev_html = unrev.replace('\n', '<br>')
    rev_html = rev.replace('\n', '<br>')
    imp_html = imp.replace('\n', '<br>')
    
    html_rows += f"""
    <tr>
      <td style="border: 1px solid #B0B0B0; padding: 8px 10px; font-weight: bold; vertical-align: top;">{sec}</td>
      <td style="border: 1px solid #B0B0B0; padding: 8px 10px; vertical-align: top;">{unrev_html}</td>
      <td style="border: 1px solid #B0B0B0; padding: 8px 10px; vertical-align: top;"><b>{rev_html}</b></td>
      <td style="border: 1px solid #B0B0B0; padding: 8px 10px; vertical-align: top;">{imp_html}</td>
    </tr>
    """

full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #000; line-height: 1.4; }}
  table {{ border-collapse: collapse; width: 100%; font-family: Calibri, Arial, sans-serif; font-size: 10pt; }}
  th {{ background-color: #FFFF00 !important; color: #000; font-weight: bold; text-align: center; border: 1px solid #B0B0B0; padding: 8px 10px; }}
  td {{ border: 1px solid #B0B0B0; padding: 8px 10px; vertical-align: top; }}
</style>
</head>
<body>
<p>Good morning @Simin Mohammad,</p>
<p>I have reviewed this OOS. Please find the following edits made in the summary table below. Please note that the <b>EM table will also need to be amended to ensure uniformity, (see highlighted) sections and reflect updated gram stain results.</b></p>
<table>
  <thead>
    <tr>
      <th style="background-color: #FFFF00; width: 20%;">Section</th>
      <th style="background-color: #FFFF00; width: 30%;">Unreviewed</th>
      <th style="background-color: #FFFF00; width: 32%;">Reviewed</th>
      <th style="background-color: #FFFF00; width: 18%;">Impact</th>
    </tr>
  </thead>
  <tbody>
    {html_rows}
  </tbody>
</table>
<p>@Simin Mohammad please review and update the OOS form text fields accordingly. You may directly use the newly compiled 8-page PDF package attached here.</p>
<p>Thanks,</p>
</body>
</html>
"""

with open(html_out, "w", encoding="utf-8") as f:
    f.write(full_html)
print(f"Saved clean HTML to: {html_out}")

# --- 3. COPY TO WINDOWS CLIPBOARD AS CF_HTML ---
def put_html_to_clipboard(html_fragment):
    HEADER = """Version:0.9
StartHTML:{start_html:08d}
EndHTML:{end_html:08d}
StartFragment:{start_fragment:08d}
EndFragment:{end_fragment:08d}
"""
    fragment_wrapper = f"<html><body><!--StartFragment-->{html_fragment}<!--EndFragment--></body></html>"
    dummy_header = HEADER.format(start_html=0, end_html=0, start_fragment=0, end_fragment=0)
    
    start_html = len(dummy_header)
    start_fragment = start_html + fragment_wrapper.find("<!--StartFragment-->") + len("<!--StartFragment-->")
    end_fragment = start_html + fragment_wrapper.find("<!--EndFragment-->")
    end_html = start_html + len(fragment_wrapper)
    
    final_header = HEADER.format(
        start_html=start_html,
        end_html=end_html,
        start_fragment=start_fragment,
        end_fragment=end_fragment
    )
    payload = (final_header + fragment_wrapper).encode('utf-8')
    
    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        CF_HTML = win32clipboard.RegisterClipboardFormat("HTML Format")
        win32clipboard.SetClipboardData(CF_HTML, payload)
        # Also set plain text fallback
        win32clipboard.SetClipboardText("OOS-261186 Review Email Table")
        print("SUCCESS: Copied rich HTML table directly into Windows Clipboard!")
    finally:
        win32clipboard.CloseClipboard()

put_html_to_clipboard(full_html)
