import os, sys, copy, shutil
OOS_ROOT = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
sys.path.insert(0, OOS_ROOT)
os.chdir(OOS_ROOT)

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfReader
import win32com.client

OUTPUT_DIR = os.path.join(OOS_ROOT, "scratch")
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"

data = {
    "oos_id": "262064",
    "client_name": "Blue Ocean Rx, LLC (E75110)",
    "sample_id": "ETX-260804-0335",
    "sample_url": "https://etrax.eagleanalytical.com/Submission/Details/ETX-260804-0335",
    "sample_name": "Tirzepatide 10mg/mL",
    "lot_number": "080326-TIR-10",
    "test_name": "USP <71> / EP 2.6.1 Sterility Test",
    "process_date": "25 Aug 2026",
    "process_date_bracket_pre": "24 Aug 2026",
    "process_date_bracket_post": "26 Aug 2026",
    "observation_date": "04 Sep 2026",
    "incubation_day": "Day 10",
    "media_growth": "1 x 100mL TSB bottle",
    "microbial_id": "ETX-260904-0331",
    "microbial_url": "https://etrax.eagleanalytical.com/Submission/Details?id=qa0T20auUnDSY-CQtpZQpw__",
    "microbial_org": "Pending",
    "processing_analyst": "[Processing Analyst]",
    "reading_analyst": "Qiyue Chen",
    "analyst_initial": "TBD",
    "bsc_id": "E001316",
    "cr_name": "CR114 (E001736)"
}

out_doc_tables = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - USP71.docx")
out_pdf_tables = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - USP71.pdf")

doc = docx.Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.4)
sec.bottom_margin = Inches(0.4)
sec.left_margin = Inches(0.45)
sec.right_margin = Inches(0.45)

def update_cell_text(cell, text, bold=False, italic=False, font_size=Pt(7.0), align=docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = font_size
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_cell_hyperlink(cell, url, text):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    part = p.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink_xml = (
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rStyle w:val="Hyperlink"/>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="467886"/>'
        f'<w:sz w:val="15"/>'
        f'<w:szCs w:val="15"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    p._p.append(parse_xml(hyperlink_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

# --- Table 1 ---
p_t1 = doc.add_paragraph()
p_t1.paragraph_format.space_before = Pt(0)
p_t1.paragraph_format.space_after = Pt(3)
r_t1 = p_t1.add_run(f"Table 1: Information for {data['sample_id']} under investigation")
r_t1.bold = True
r_t1.font.name = "Times New Roman"
r_t1.font.size = Pt(9.5)

t1 = doc.add_table(rows=2, cols=6)
t1.style = 'Table Grid'
t1.autofit = False

t1_headers = [
    "Processing Analyst",
    "Reading Analyst",
    "Sample ID",
    "Related Microbial ID",
    "Media with microbial growth",
    "Microbial ID"
]
for c_idx, h_text in enumerate(t1_headers):
    update_cell_text(t1.rows[0].cells[c_idx], h_text, bold=True, font_size=Pt(7.5))

r1 = t1.rows[1]
update_cell_text(r1.cells[0], data["processing_analyst"], font_size=Pt(7.5))
update_cell_text(r1.cells[1], data["reading_analyst"], font_size=Pt(7.5))
set_cell_hyperlink(r1.cells[2], data["sample_url"], data["sample_id"])
set_cell_hyperlink(r1.cells[3], data["microbial_url"], data["microbial_id"])
update_cell_text(r1.cells[4], data["media_growth"], font_size=Pt(7.5))
update_cell_text(r1.cells[5], data["microbial_org"], font_size=Pt(7.5))

# --- Table 2 ---
p_t2 = doc.add_paragraph()
p_t2.paragraph_format.space_before = Pt(8)
p_t2.paragraph_format.space_after = Pt(3)
r_t2 = p_t2.add_run("Table 2: Environmental Monitoring for Analyst and Cleanroom")
r_t2.bold = True
r_t2.font.name = "Times New Roman"
r_t2.font.size = Pt(9.5)

t2 = doc.add_table(rows=16, cols=9)
t2.style = 'Table Grid'
t2.autofit = False

t2_headers = [
    "Environmental Monitoring\n(EM) Sampling Site",
    "Frequency",
    "Date\n(DDMMM\nYYYY)",
    "Analyst\n(Initials)",
    "Day /Week(s)",
    "Observation",
    "Environmental\nMonitoring Plate\nETX ID",
    "Microbial ID",
    "Notes"
]
for c_idx, h_text in enumerate(t2_headers):
    update_cell_text(t2.rows[0].cells[c_idx], h_text, bold=True, font_size=Pt(7.0))

# Row 1: Merged header
t2.rows[1].cells[0].merge(t2.rows[1].cells[-1])
update_cell_text(t2.rows[1].cells[0], f"Personnel EM Bracketing {data['process_date']}", bold=True, font_size=Pt(7.5))

# Personnel rows (2, 3, 4)
em_personnel_data = [
    ("Personal (Left Touch and Right Touch)", "Daily", data['process_date_bracket_pre'], data['analyst_initial'], "Date Before Testing", "No growth", "N/A", "N/A", "None"),
    ("Personal (Left Touch and Right Touch)", "Daily", data['process_date'], data['analyst_initial'], "Date of Testing", "No growth", "N/A", "N/A", "None"),
    ("Personal (Left Touch and Right Touch)", "Daily", data['process_date_bracket_post'], data['analyst_initial'], "Date After Testing", "No growth", "N/A", "N/A", "None")
]
for r_offset, row_vals in enumerate(em_personnel_data):
    for c_idx, val in enumerate(row_vals):
        align = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        update_cell_text(t2.rows[2 + r_offset].cells[c_idx], val, font_size=Pt(7.0), align=align)

# Row 5: Merged BSC header
t2.rows[5].cells[0].merge(t2.rows[5].cells[-1])
update_cell_text(t2.rows[5].cells[0], f"Biological Safety Cabinet EM Bracketing Biological Safety Cabinet (BSC) {data['bsc_id']}", bold=True, font_size=Pt(7.5))

# Surface sampling (6, 7, 8)
em_surf_data = [
    ("Surface Sampling of ISO 5 (4 locations)", "Daily", data['process_date_bracket_pre'], data['analyst_initial'], "Date Before Testing", "No growth", "N/A", "N/A", "None"),
    ("Surface Sampling of ISO 5 (4 locations)", "Daily", data['process_date'], data['analyst_initial'], "Date of Testing", "No growth", "N/A", "N/A", "None"),
    ("Surface Sampling of ISO 5 (4 locations)", "Daily", data['process_date_bracket_post'], data['analyst_initial'], "Date After Testing", "No growth", "N/A", "N/A", "None")
]
for r_offset, row_vals in enumerate(em_surf_data):
    for c_idx, val in enumerate(row_vals):
        align = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        update_cell_text(t2.rows[6 + r_offset].cells[c_idx], val, font_size=Pt(7.0), align=align)

# Settling sampling (9, 10, 11)
em_sett_data = [
    ("Settling Sampling of ISO 5 (2 locations)", "Daily", data['process_date_bracket_pre'], data['analyst_initial'], "Date Before Testing", "No growth", "N/A", "N/A", "None"),
    ("Settling Sampling of ISO 5 (2 locations)", "Daily", data['process_date'], data['analyst_initial'], "Date of Testing", "No growth", "N/A", "N/A", "None"),
    ("Settling Sampling of ISO 5 (2 locations)", "Daily", data['process_date_bracket_post'], data['analyst_initial'], "Date After Testing", "No growth", "N/A", "N/A", "None")
]
for r_offset, row_vals in enumerate(em_sett_data):
    for c_idx, val in enumerate(row_vals):
        align = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        update_cell_text(t2.rows[9 + r_offset].cells[c_idx], val, font_size=Pt(7.0), align=align)

# Row 12: Weekly Active Air header
t2.rows[12].cells[0].merge(t2.rows[12].cells[-1])
update_cell_text(t2.rows[12].cells[0], f"Weekly Active Air Sampling of Cleanroom - {data['cr_name']}", bold=True, font_size=Pt(7.5))

# Row 13: Active Air data
air_vals = ("Active Air Sampling of Cleanrooms", "Weekly", data['process_date'], "SMO", "Week of Testing", "4 CFU in ISO 8 Room 114", "ETX-260901-0112", "3 Gram (+) cocci,\n1 Hyphae", "None")
for c_idx, val in enumerate(air_vals):
    align = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    update_cell_text(t2.rows[13].cells[c_idx], val, font_size=Pt(7.0), align=align)

# Row 14: Surface Sampling Anteroom header
t2.rows[14].cells[0].merge(t2.rows[14].cells[-1])
update_cell_text(t2.rows[14].cells[0], f"Surface Sampling of Anteroom and Cleanroom – {data['cr_name']}", bold=True, font_size=Pt(7.5))

# Row 15: Cleanroom surface data
clean_vals = ("Surface Sampling of Cleanrooms", "Weekly", data['process_date'], "SMO", "Week of Testing", "No growth", "N/A", "N/A", "None")
for c_idx, val in enumerate(clean_vals):
    align = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    update_cell_text(t2.rows[15].cells[c_idx], val, font_size=Pt(7.0), align=align)

# Row heights
for t in [t1, t2]:
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trHeight = parse_xml(f'<w:trHeight xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="220" w:hRule="atLeast"/>')
        trPr.append(trHeight)

doc.save(out_doc_tables)
print("Saved Docx to:", out_doc_tables)

# Export to PDF via Word COM
try:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(os.path.abspath(out_doc_tables), ReadOnly=True)
        d.SaveAs(os.path.abspath(out_pdf_tables), FileFormat=17)
        d.Close(SaveChanges=False)
        print("Saved PDF via Word to:", out_pdf_tables)
    finally:
        word.Quit()
except Exception as e:
    print(f"Word COM error: {e}")

if os.path.exists(out_pdf_tables):
    r_check = PdfReader(out_pdf_tables)
    print(f"Verified Tables PDF Page Count: {len(r_check.pages)} (Must be 1)")

# Sync to Desktop and Documents
for dest_dir in [DESKTOP_DIR, DOCUMENTS_DIR]:
    shutil.copy2(out_doc_tables, os.path.join(dest_dir, os.path.basename(out_doc_tables)))
    if os.path.exists(out_pdf_tables):
        shutil.copy2(out_pdf_tables, os.path.join(dest_dir, os.path.basename(out_pdf_tables)))
print("Synced files successfully.")
