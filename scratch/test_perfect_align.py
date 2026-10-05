import os
import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import fitz

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_path = os.path.join(desktop_dir, "Celsis table OOS-262080.docx")

doc = docx.Document(docx_path)
t0 = doc.tables[0]
cell = t0.rows[1].cells[2]

# URL and Text
url = "https://etrax.eagleanalytical.com/SubmissionTest/Details/jzraMOOTFYUFD%24erkwgubw__"
text = "ETX-260828-0527"

# Inspect Col 1 XML to replicate EXACT structure
col1 = t0.rows[1].cells[1]
print("Col 1 paragraph XML:\n", col1.paragraphs[0]._p.xml)

# Clear cell completely
cell._tc.clear_content()

# Reconstruct paragraph and hyperlink with EXACT font size 14 (7pt) and line spacing 360
part = doc.part
r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

tc_xml = (
    f'<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    f'<w:tcPr>'
    f'<w:tcW w:w="1400" w:type="dxa"/>'
    f'<w:vAlign w:val="center"/>'
    f'</w:tcPr>'
    f'<w:p>'
    f'<w:pPr>'
    f'<w:spacing w:line="360" w:lineRule="auto"/>'
    f'<w:jc w:val="center"/>'
    f'<w:rPr>'
    f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
    f'<w:sz w:val="14"/>'
    f'<w:szCs w:val="14"/>'
    f'</w:rPr>'
    f'</w:pPr>'
    f'<w:hyperlink r:id="{r_id}" w:history="1">'
    f'<w:r>'
    f'<w:rPr>'
    f'<w:rStyle w:val="Hyperlink"/>'
    f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
    f'<w:color w:val="0000FF"/>'
    f'<w:sz w:val="14"/>'
    f'<w:szCs w:val="14"/>'
    f'<w:u w:val="single"/>'
    f'</w:rPr>'
    f'<w:t>{text}</w:t>'
    f'</w:r>'
    f'</w:hyperlink>'
    f'</w:p>'
    f'</w:tc>'
)

new_tc = parse_xml(tc_xml)
cell._tc.getparent().replace(cell._tc, new_tc)

out_docx = "scratch/test_perfect_align.docx"
doc.save(out_docx)
print(f"Saved {out_docx}")

# Export via Word COM
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    doc_word = word.Documents.Open(os.path.abspath(out_docx))
    out_pdf = os.path.abspath("scratch/test_perfect_align.pdf")
    doc_word.ExportAsFixedFormat(out_pdf, 17)
    doc_word.Close(False)
    print("Exported PDF via Word COM")
finally:
    word.Quit()

d = fitz.open("scratch/test_perfect_align.pdf")
p0 = d[0]
words = p0.get_text('words')
print("\n=== WORDS ON ROW 1 ===")
for w in words:
    if 115 <= w[1] <= 135:
        print(f"{w[4]:20} y0={w[1]:.3f} y1={w[3]:.3f}")

pix = p0.get_pixmap(dpi=150, clip=fitz.Rect(50, 40, 560, 160))
pix.save("scratch/test_perfect_align.png")
