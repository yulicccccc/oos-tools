import os
from docx import Document
from docx.shared import Pt
import win32com.client as win32

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_path = os.path.join(desktop_dir, "OOS-261242 Review Email.docx")
html_path = os.path.join(desktop_dir, "OOS-261242_Review_Email_Robin_Format.html")

# 1. Update DOCX
doc = Document(docx_path)
tbl = doc.tables[0]

# Find the EM Table row
for row in tbl.rows:
    sec_text = row.cells[0].text
    if "EM Table" in sec_text:
        row.cells[0].text = "EM Table Formatting"
        row.cells[0].paragraphs[0].runs[0].font.name = "Calibri"
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(10)
        
        # Remove "The EM Table was omitted from the PDF submission."
        row.cells[1].text = (
            "In the draft Word document, Table 1 date had excessive spacing (\"14MAY   2026\"), "
            "CFU count lacked spacing (\"11CFUs\"), Table 2 Row 7 had plural typo "
            "(\"1 CFUs on table  in 115 ISO 8 and 1 CFUs on cart in 115 ISO 8\"), and multi-organism entries lacked line breaks."
        )
        row.cells[1].paragraphs[0].runs[0].font.name = "Calibri"
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
        
        row.cells[2].text = (
            "Amended both tables to correct spacing (14MAY 2026, 11 CFUs), corrected plural typos "
            "(1 CFU on table in 115 ISO 8 and 1 CFU on cart in 115 ISO 8), formatted multi-organism line breaks, "
            "standardized bracketing section headers, and attached as Page 8 to complete the finalized 8-page package."
        )
        row.cells[2].paragraphs[0].runs[0].font.name = "Calibri"
        row.cells[2].paragraphs[0].runs[0].font.size = Pt(10)

doc.save(docx_path)
print("Updated DOCX: removed 'The EM Table was omitted from the PDF submission.'")

# 2. Update HTML
if os.path.exists(html_path):
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()
    html = html.replace("The EM Table was omitted from the PDF submission. ", "")
    html = html.replace("The EM Table was omitted from the PDF submission.", "")
    html = html.replace("EM Table Attachment & Formatting", "EM Table Formatting")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    print("Updated HTML.")

# 3. Copy via Word COM to clipboard
word_app = win32.Dispatch("Word.Application")
word_app.Visible = False
try:
    d = word_app.Documents.Open(os.path.abspath(docx_path))
    d.Content.Copy()
    d.Close(False)
    print("SUCCESS: Copied updated email to Windows Clipboard via Word COM!")
finally:
    word_app.Quit()
