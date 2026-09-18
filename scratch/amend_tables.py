import docx
from docx.shared import Pt
import copy
import os

DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"
TABLES_DOCX = os.path.join(DOCUMENTS_DIR, "Tables OOS-261814 GoGoMeds Select (E10747) - ScanRDI.docx")

doc = docx.Document(TABLES_DOCX)
t = doc.tables[1] # Table 2 (EM Table)

def update_cell(cell, text, bold=False, font_size=Pt(7), align=docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = font_size
    run.font.bold = bold
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

# 1. Highlight 6: Row 0 Header Date (DDMMMYY) - fix extra space
update_cell(t.rows[0].cells[2], "Date\n(DDMMMYY)", bold=True, font_size=Pt(6.5))

# 2. Highlight 7: Row 7 SU Settling BSC 1313 - formalize observation and add quality record references to Notes
update_cell(t.rows[7].cells[5], "2 CFUs on Left Settling Plate (S1)", font_size=Pt(6.5))
update_cell(t.rows[7].cells[8], "Plate was desiccated on 5 day read\n(OOS-261891, NCR-N26483)", font_size=Pt(6.0))

# 3. Highlight 4 & 1: Row 9 & Row 13 - Suite 115 Cleanroom nomenclature
update_cell(t.rows[9].cells[0], "Weekly Active Air Sampling Bracketing - Cleanroom CR115 (E001737)", bold=True, font_size=Pt(7))
update_cell(t.rows[13].cells[0], "Surface Sampling of Anteroom and Cleanroom Bracketing - Cleanroom CR115 (E001737)", bold=True, font_size=Pt(7))

# 4. Highlight 3 & 5 & 2: Row 11 & Row 15 - L-Suite Cleanroom nomenclature
update_cell(t.rows[11].cells[0], "Weekly Active Air Sampling Bracketing - Cleanroom CR144 (E001978, L-Suite)", bold=True, font_size=Pt(7))
update_cell(t.rows[15].cells[0], "Surface Sampling of Anteroom and Cleanroom Bracketing - Cleanroom CR144 (E001978, L-Suite)", bold=True, font_size=Pt(7))

# 5. Highlight 8: Row 12 L-Suite Active Air Observation - include CR144 no growth and CR-room nomenclature
air_obs = (
    "No growth in ISO 7 CR144;\n"
    "6 CFUs in ISO 8 CR143 Sec I,\n"
    "32 CFUs in ISO 8 CR143 Sec II, and\n"
    "12 CFUs in ISO 8 CR142"
)
update_cell(t.rows[12].cells[5], air_obs, font_size=Pt(6.0))

# Ensure ALL cells are centered and formatted
for row in t.rows:
    for cell in row.cells:
        cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

out_docx_path = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\scratch\test_amended_tables.docx"
doc.save(out_docx_path)
print("Saved test_amended_tables.docx")
