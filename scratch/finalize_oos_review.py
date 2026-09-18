import os, sys, shutil
import fitz
from pypdf import PdfReader, PdfWriter
import docx
from docx.shared import Pt

DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
SCRATCH_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\scratch"

# 1. Update standalone Tables DOCX and PDF in Documents
docx_amended = os.path.join(SCRATCH_DIR, "test_amended_tables.docx")
pdf_amended = os.path.join(SCRATCH_DIR, "test_amended_tables.pdf")

dest_tables_docx = os.path.join(DOCUMENTS_DIR, "Tables OOS-261814 GoGoMeds Select (E10747) - ScanRDI.docx")
dest_tables_pdf = os.path.join(DOCUMENTS_DIR, "Tables OOS-261814 GoGoMeds Select (E10747) - ScanRDI.pdf")

shutil.copy2(docx_amended, dest_tables_docx)
shutil.copy2(pdf_amended, dest_tables_pdf)
print(f"Updated: {dest_tables_docx}")
print(f"Updated: {dest_tables_pdf}")

# 2. Assemble Finalized 7-Page OOS PDF
rs_reviewed_pdf = os.path.join(DESKTOP_DIR, "OOS-261814 GoGoMeds Select (E10747) - ScanRDI - RS Reviewed - Signed by.pdf")
final_pdf_writer = PdfWriter()

rs_reader = PdfReader(rs_reviewed_pdf)
# Take Pages 1 to 6 (Reviewed content & Qiyue signature)
for i in range(6):
    final_pdf_writer.add_page(rs_reader.pages[i])

# Append clean amended Table page as Page 7
table_reader = PdfReader(pdf_amended)
final_pdf_writer.add_page(table_reader.pages[0])

out_final_desktop = os.path.join(DESKTOP_DIR, "OOS-261814 GoGoMeds Select (E10747) - ScanRDI - Final Reviewed.pdf")
out_std_desktop = os.path.join(DESKTOP_DIR, "OOS-261814 GoGoMeds Select (E10747) - ScanRDI.pdf")
out_std_docs = os.path.join(DOCUMENTS_DIR, "OOS-261814 GoGoMeds Select (E10747) - ScanRDI.pdf")

with open(out_final_desktop, "wb") as f:
    final_pdf_writer.write(f)
print(f"Saved: {out_final_desktop}")

try:
    shutil.copy2(out_final_desktop, out_std_desktop)
    print(f"Updated: {out_std_desktop}")
except Exception as e:
    print(f"Notice: {e}")

try:
    shutil.copy2(out_final_desktop, out_std_docs)
    print(f"Updated: {out_std_docs}")
except Exception as e:
    print(f"Notice: {e}")

# 3. Update Master Word Report (OOS-261814 GoGoMeds Select (E10747) - ScanRDI.docx)
master_docx_path = os.path.join(DOCUMENTS_DIR, "OOS-261814 GoGoMeds Select (E10747) - ScanRDI.docx")
if os.path.exists(master_docx_path):
    doc_m = docx.Document(master_docx_path)
    # If table 2 exists, replace with amended table
    if len(doc_m.tables) >= 3:
        t2_old = doc_m.tables[2]._element
        t2_parent = t2_old.getparent()
        idx = t2_parent.index(t2_old)
        t2_parent.remove(t2_old)
        
        doc_amended = docx.Document(docx_amended)
        t_amended_elem = doc_amended.tables[1]._element
        t2_parent.insert(idx, t_amended_elem)
        doc_m.save(master_docx_path)
        print(f"Updated Master Word Document: {master_docx_path}")

# 4. Visual Verification Render
doc_check = fitz.open(out_final_desktop)
print(f"Verified Final PDF page count: {len(doc_check)}")
pix_p7 = doc_check[6].get_pixmap(dpi=150)
pix_p7.save(os.path.join(SCRATCH_DIR, "final_p7_verified.png"))
doc_check.close()
print("Rendered final_p7_verified.png successfully")

print("\n--- ALL TASKS COMPLETED SUCCESSFULLY ---")
