import os, sys, shutil, datetime, copy
import docx
from docx.shared import Pt, Inches
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
from docxtpl import DocxTemplate
import win32com.client
from pypdf import PdfReader
import fitz

OOS_ROOT = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
sys.path.insert(0, OOS_ROOT)
os.chdir(OOS_ROOT)

OUTPUT_DIR = os.path.join(OOS_ROOT, "scratch")
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"

# Base template from proven pure Celsis style without aliquoting
TPL_BASE = os.path.join(OUTPUT_DIR, "tables_71_pure_celsis.docx")
if not os.path.exists(TPL_BASE):
    TPL_BASE = os.path.join(OOS_ROOT, "tables for celsis.docx")

data = {
    "oos_id": "262098",
    "client_name": "Solyn LLC (E75000)",
    "sample_id": "ETX-260902-0505 and ETX-260914-0470",
    "sample_id_1": "ETX-260902-0505",
    "sample_url": "https://etrax.eagleanalytical.com/Submission/Details/%247Ydn%24JOiYgmsnaKjbxn-g__",
    "sample_name": "GLP3R/Cagrilinitide",
    "lot_number": "2608-216",
    "analyst_name": "Alex Saravia",
    "reading_name": "Elysse Nioupin",
    "positive_id": "ETX-260910-0290",
    "positive_url": "https://etrax.eagleanalytical.com/Submission/Details?id=oZa2TD1UI-0eAJOFysBJ3A__",
    "positive_media": "1 x 100mL TSB bottle",
    "positive_org": "Microbacterium sp. PM5\n(Gram (+) rods)",

    # Table 2 Dates & Analysts (04Sep26)
    "process_date": "04Sep26",
    "pro_before_test": "03Sep26",
    "pro_test_date": "04Sep26",
    "pro_after_test": "08Sep26",
    "pro_analyst_initial": "ES",
    "pro_date_of_weekly": "04Sep26",

    # Personnel EM
    "pro_be_obs_pers_dur_pro": "No growth", "pro_be_etx_pers_dur_pro": "N/A", "pro_be_id_pers_dur_pro": "N/A",
    "pro_obs_pers_dur_pro": "No growth",    "pro_etx_pers_dur_pro": "N/A", "pro_id_pers_dur_pro": "N/A",
    "pro_af_obs_pers_dur_pro": "No growth", "pro_af_etx_pers_dur_pro": "N/A", "pro_af_id_pers_dur_pro": "N/A",

    # Surface ISO 5
    "pro_be_obs_surf_dur_pro": "No growth", "pro_be_etx_surf_dur_pro": "N/A", "pro_be_id_surf_dur_pro": "N/A",
    "pro_obs_surf_dur_pro": "No growth",    "pro_etx_surf_dur_pro": "N/A", "pro_id_surf_dur_pro": "N/A",
    "pro_af_obs_surf_dur_pro": "No growth", "pro_af_etx_surf_dur_pro": "N/A", "pro_af_id_surf_dur_pro": "N/A",

    # Settling ISO 5
    "pro_be_obs_sett_dur_pro": "No growth", "pro_be_etx_sett_dur_pro": "N/A", "pro_be_id_sett_dur_pro": "N/A",
    "pro_obs_sett_dur_pro": "No growth",    "pro_etx_sett_dur_pro": "N/A", "pro_id_sett_dur_pro": "N/A",
    "pro_af_obs_sett_dur_pro": "No growth", "pro_af_etx_sett_dur_pro": "N/A", "pro_af_id_sett_dur_pro": "N/A",

    # Weekly Active Air & Surface
    "pro_obs_air_wk_of": "2 CFU (ISO 8 114)",
    "pro_etx_air_wk_of": "ETX-260914-0487",
    "pro_id_air_wk_of": "Corynebacterium ureicelerivorans\nMycobacterium grossiae",
    "pro_obs_room_wk_of": "No growth",
    "pro_etx_room_wk_of": "N/A",
    "pro_id_room_wk_of": "N/A"
}

OUT_DOCX = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - USP71.docx")
OUT_PDF = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - USP71.pdf")

tpl = DocxTemplate(TPL_BASE)
tpl.render(data)
tpl.save(OUT_DOCX)

# Post-processing: Add Native Hyperlinks in Table 1 with strict Times New Roman
doc_rendered = docx.Document(OUT_DOCX)

URL_0487 = "https://etrax.eagleanalytical.com/SubmissionTest/Details/fd3G2StZClcy1TP2ES6BLw__"
URL_0520 = "https://etrax.eagleanalytical.com/SubmissionTest/Details/qnwcLQO5BWeJhWCBG7jl8Q__#TestDetails"
URL_RETEST = "https://etrax.eagleanalytical.com/Submission/Details/18rFOYVWqrO0Afk1OwB2hg__"
URL_RETEST_ID = "https://etrax.eagleanalytical.com/Submission/Details/7K%24JgQe6bduNBFY3pBa03w__"

def set_cell_clean_text(cell, text, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_hyperlink_to_cell(cell, url, text, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    part = cell.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:hyperlink r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="0000FF"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_dual_microbial_id_cell(cell, org1, org2, font_size_pt=6.5):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{org1}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    p2_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{org2}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    p_empty_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="120" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_empty_xml))
    tc.append(parse_xml(p2_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_single_microbial_id_cell(cell, species, gram_stain=None, font_size_pt=6.5):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{species}</w:t>'
        f'</w:r>'
    )
    if gram_stain:
        p_xml += (
            f'<w:r>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:sz w:val="{int(font_size_pt*2)}"/>'
            f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
            f'</w:rPr>'
            f'<w:br/>'
            f'<w:t>{gram_stain}</w:t>'
            f'</w:r>'
        )
    p_xml += '</w:p>'
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_table1_microbial_id_cell(cell, species, gram_stain, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{species}</w:t>'
        f'</w:r>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:br/>'
        f'<w:t>{gram_stain}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

# 1. Finalize Table 1 (Both Samples)
t1_rendered = doc_rendered.tables[0]
add_hyperlink_to_cell(t1_rendered.rows[1].cells[2], data["sample_url"], data["sample_id_1"], font_size_pt=7)
add_hyperlink_to_cell(t1_rendered.rows[1].cells[3], data["positive_url"], data["positive_id"], font_size_pt=7)
set_table1_microbial_id_cell(t1_rendered.rows[1].cells[5], "Microbacterium sp. PM5", "(Gram (+) rods)", font_size_pt=7)

# Add Row 2 for ETX-260914-0470
r_new = t1_rendered.add_row()
cells_r2 = r_new.cells
set_cell_clean_text(cells_r2[0], "Abayomi Odugbesi", font_size_pt=7)
set_cell_clean_text(cells_r2[1], "Andrew Carrillo", font_size_pt=7)
add_hyperlink_to_cell(cells_r2[2], URL_RETEST, "ETX-260914-0470", font_size_pt=7)
add_hyperlink_to_cell(cells_r2[3], URL_RETEST_ID, "ETX-260921-0498", font_size_pt=7)
set_cell_clean_text(cells_r2[4], "1 x 100mL TSB bottle", font_size_pt=7)
set_table1_microbial_id_cell(cells_r2[5], "Microbacterium sp. PM5", "(Gram (+) short rods)", font_size_pt=7)

# 2. Finalize Table 2 (04Sep26 Processing)
t2_rendered = doc_rendered.tables[1]
# Row 13: Weekly Active Air
set_cell_clean_text(t2_rendered.rows[13].cells[3], "04Sep26", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[13].cells[4], "ISS", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[13].cells[8], "2 CFU (ISO 8 114)", font_size_pt=7)
add_hyperlink_to_cell(t2_rendered.rows[13].cells[9], URL_0487, "ETX-260914-0487", font_size_pt=7)
set_dual_microbial_id_cell(t2_rendered.rows[13].cells[10], "Corynebacterium ureicelerivorans", "Mycobacterium grossiae", font_size_pt=6.5)

# Row 15: Weekly Surface
set_cell_clean_text(t2_rendered.rows[15].cells[3], "04Sep26", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[4], "ISS", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[8], "No growth", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[10], "N/A", font_size_pt=7)

# 3. Create Table 3 (15Sep26 Processing for ETX-260914-0470)
t3_tbl = copy.deepcopy(t2_rendered._tbl)

# Add Page Break before Table 3
p_br = doc_rendered.add_paragraph()
p_br_pr = p_br._p.get_or_add_pPr()
p_br_sp = parse_xml('<w:spacing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:before="0" w:after="0"/>')
p_br_pr.append(p_br_sp)
r_br = p_br.add_run()
r_br.add_break(docx.enum.text.WD_BREAK.PAGE)

# Add Table 3 Heading
p_t3 = doc_rendered.add_paragraph()
p_t3_xml = (
    f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    f'<w:pPr>'
    f'<w:spacing w:before="0" w:after="160"/>'
    f'<w:rPr>'
    f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
    f'<w:bCs/>'
    f'<w:sz w:val="18"/>'
    f'<w:szCs w:val="18"/>'
    f'<w:u w:val="single"/>'
    f'</w:rPr>'
    f'</w:pPr>'
    f'<w:r>'
    f'<w:rPr>'
    f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
    f'<w:bCs/>'
    f'<w:sz w:val="18"/>'
    f'<w:szCs w:val="18"/>'
    f'<w:u w:val="single"/>'
    f'</w:rPr>'
    f'<w:t>Table 3: Environmental Monitoring from Processing Performed on 15Sep26</w:t>'
    f'</w:r>'
    f'</w:p>'
)
p_t3._p.getparent().replace(p_t3._p, parse_xml(p_t3_xml))

# Append Table 3
doc_rendered._body._body.append(t3_tbl)
t3_rendered = doc_rendered.tables[2]

# Populate Table 3 for 15Sep26 processing by AO in BSC 1314
# Rows:
# 2: Personnel Pre (10Sep26, AO)
set_cell_clean_text(t3_rendered.rows[2].cells[3], "10Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[2].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[2].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[2].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[2].cells[11], "None", font_size_pt=7)

# 3: Personnel Test Date (15Sep26, AO)
set_cell_clean_text(t3_rendered.rows[3].cells[3], "15Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[3].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[3].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[3].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[3].cells[11], "None", font_size_pt=7)

# 4: Personnel Post (16Sep26, AO)
set_cell_clean_text(t3_rendered.rows[4].cells[3], "16Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[4].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[4].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[4].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[4].cells[11], "None", font_size_pt=7)

# 6: Surface Pre (10Sep26, AO)
set_cell_clean_text(t3_rendered.rows[6].cells[3], "10Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[6].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[6].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[6].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[6].cells[11], "None", font_size_pt=7)

# 7: Surface Test Date (15Sep26, AO)
set_cell_clean_text(t3_rendered.rows[7].cells[3], "15Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[7].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[7].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[7].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[7].cells[11], "None", font_size_pt=7)

# 8: Surface Post (16Sep26, AO)
set_cell_clean_text(t3_rendered.rows[8].cells[3], "16Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[8].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[8].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[8].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[8].cells[11], "None", font_size_pt=7)

# 9: Settling Pre (10Sep26, AO)
set_cell_clean_text(t3_rendered.rows[9].cells[3], "10Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[9].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[9].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[9].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[9].cells[11], "None", font_size_pt=7)

# 10: Settling Test Date (15Sep26, AO)
set_cell_clean_text(t3_rendered.rows[10].cells[3], "15Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[10].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[10].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[10].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[10].cells[11], "None", font_size_pt=7)

# 11: Settling Post (16Sep26, AO)
set_cell_clean_text(t3_rendered.rows[11].cells[3], "16Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[11].cells[4], "AO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[11].cells[7], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[11].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[11].cells[11], "None", font_size_pt=7)

# 13: Weekly Active Air (10Sep26, SMO) - Recovered 1 CFU in ISO 8 Anteroom 114
set_cell_clean_text(t3_rendered.rows[13].cells[3], "10Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[13].cells[4], "SMO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[13].cells[8], "1 CFU (ISO 8 114)", font_size_pt=7)
add_hyperlink_to_cell(t3_rendered.rows[13].cells[9], URL_0520, "ETX-260921-0520", font_size_pt=7)
set_single_microbial_id_cell(t3_rendered.rows[13].cells[10], "Micrococcus luteus", "(Gram (+) cocci)", font_size_pt=6.5)

# 15: Weekly Surface (10Sep26, SMO) - No growth
set_cell_clean_text(t3_rendered.rows[15].cells[3], "10Sep26", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[15].cells[4], "SMO", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[15].cells[8], "No growth", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[15].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t3_rendered.rows[15].cells[10], "N/A", font_size_pt=7)

# Strict 100% Times New Roman Enforcement across all runs in all tables
for t_idx, table in enumerate(doc_rendered.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    rPr = run._r.get_or_add_rPr()
                    rFonts = parse_xml('<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>')
                    for existing in rPr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts'):
                        rPr.remove(existing)
                    rPr.append(rFonts)

# Row heights
for t in doc_rendered.tables:
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr()
        trHeight = parse_xml(f'<w:trHeight xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="220" w:hRule="atLeast"/>')
        trPr.append(trHeight)

doc_rendered.save(OUT_DOCX)
print("Saved DOCX:", OUT_DOCX)

# Convert to PDF via Word COM
try:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(os.path.abspath(OUT_DOCX), ReadOnly=True)
        d.SaveAs(os.path.abspath(OUT_PDF), FileFormat=17)
        d.Close(SaveChanges=False)
        print("Exported PDF via Word COM:", OUT_PDF)
    finally:
        word.Quit()
except Exception as e:
    print(f"Word COM error: {e}")

if os.path.exists(OUT_PDF):
    reader = PdfReader(OUT_PDF)
    print(f"Verified PDF Page Count: {len(reader.pages)} (Must be 2)")

# Also produce Phase II tables with identical content
OUT_P2_DOCX = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - Phase II.docx")
OUT_P2_PDF = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - Phase II.pdf")
shutil.copy2(OUT_DOCX, OUT_P2_DOCX)
shutil.copy2(OUT_PDF, OUT_P2_PDF)

# Sync to Desktop and Documents safely
for dest_dir in [DESKTOP_DIR, DOCUMENTS_DIR]:
    # USP71 tables
    dest_docx = os.path.join(dest_dir, os.path.basename(OUT_DOCX))
    dest_pdf = os.path.join(dest_dir, os.path.basename(OUT_PDF))
    try:
        shutil.copy2(OUT_DOCX, dest_docx)
        print("Synced DOCX to:", dest_docx)
    except Exception as e:
        print(f"Notice: {dest_docx} copy skipped ({e})")
        revised_docx = os.path.join(dest_dir, f"Tables OOS-{data['oos_id']} {data['client_name']} - USP71 - Revised.docx")
        try:
            shutil.copy2(OUT_DOCX, revised_docx)
            print("Saved Revised DOCX fallback to:", revised_docx)
        except Exception:
            pass

    if os.path.exists(OUT_PDF):
        try:
            shutil.copy2(OUT_PDF, dest_pdf)
            print("Synced PDF to:", dest_pdf)
        except Exception as e:
            print(f"Notice: {dest_pdf} copy skipped ({e})")
            revised_pdf = os.path.join(dest_dir, f"Tables OOS-{data['oos_id']} {data['client_name']} - USP71 - Revised.pdf")
            try:
                shutil.copy2(OUT_PDF, revised_pdf)
                print("Saved Revised PDF fallback to:", revised_pdf)
            except Exception:
                pass

    # Phase II tables
    dest_p2_docx = os.path.join(dest_dir, os.path.basename(OUT_P2_DOCX))
    dest_p2_pdf = os.path.join(dest_dir, os.path.basename(OUT_P2_PDF))
    try:
        shutil.copy2(OUT_P2_DOCX, dest_p2_docx)
        print("Synced Phase II DOCX to:", dest_p2_docx)
    except Exception as e:
        print(f"Notice: {dest_p2_docx} copy skipped ({e})")
        revised_p2_docx = os.path.join(dest_dir, f"Tables OOS-{data['oos_id']} {data['client_name']} - Phase II - Revised.docx")
        try:
            shutil.copy2(OUT_P2_DOCX, revised_p2_docx)
            print("Saved Revised Phase II DOCX fallback to:", revised_p2_docx)
        except Exception:
            pass

    try:
        shutil.copy2(OUT_P2_PDF, dest_p2_pdf)
        print("Synced Phase II PDF to:", dest_p2_pdf)
    except Exception as e:
        print(f"Notice: {dest_p2_pdf} copy skipped ({e})")
        revised_p2_pdf = os.path.join(dest_dir, f"Tables OOS-{data['oos_id']} {data['client_name']} - Phase II - Revised.pdf")
        try:
            shutil.copy2(OUT_P2_PDF, revised_p2_pdf)
            print("Saved Revised Phase II PDF fallback to:", revised_p2_pdf)
        except Exception:
            pass

print("All standalone table files generated successfully.")
