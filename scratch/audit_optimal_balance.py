import openpyxl
from datetime import datetime

excel_path = r'scratch\Sterile Lab - OOS Tracking Log.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("="*80)
print("COMPREHENSIVE AUDIT OF 'OPTIMAL BALANCE PHARMACY' (E19193) ACROSS ALL SHEETS")
print("="*80)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    
    header = None
    for idx, r in enumerate(rows[:5]):
        r_str = " ".join([str(x) for x in r if x is not None])
        if "OOS" in r_str or "ETX" in r_str or "Client" in r_str:
            header = [str(x).strip() if x is not None else f"Col{i}" for i, x in enumerate(r)]
            header_idx = idx
            break
            
    matches = []
    for r_idx, r in enumerate(rows):
        r_str = " ".join([str(x) for x in r if x is not None])
        if "optimal balance" in r_str.lower() or "e19193" in r_str.lower():
            matches.append((r_idx + 1, r))
            
    if matches:
        print(f"\nSheet: [{sheet_name}] - {len(matches)} matching rows found:")
        for r_num, r_vals in matches:
            cols_with_val = []
            for c_idx, val in enumerate(r_vals):
                if val is not None and str(val).strip() != "":
                    col_name = header[c_idx] if header and c_idx < len(header) else f"Col{c_idx}"
                    cols_with_val.append(f"{col_name}: {val}")
            print(f"  Row {r_num} -> " + " | ".join(cols_with_val[:8]))
