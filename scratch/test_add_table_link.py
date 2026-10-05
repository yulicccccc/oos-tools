import os
import shutil
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import fitz

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_path = os.path.join(desktop_dir, "Celsis table OOS-262080.docx")
pdf_path = os.path.join(desktop_dir, "Celsis table OOS-262080.pdf")

url = "https://etrax.eagleanalytical.com/SubmissionTest/Details/jzraMOOTFYUFD%24erkwgubw__"
sample_id = "ETX-260828-0527"

def set_cell_hyperlink(cell, url, text):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    part = p.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink_xml = (
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rStyle w:val="Hyperlink"/>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="0000FF"/>'
        f'<w:sz w:val="15"/>'
        f'<w:szCs w:val="15"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    p._p.append(parse_xml(hyperlink_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

doc = docx.Document(docx_path)
t0 = doc.tables[0]
# Cell 2 is Sample ID
set_cell_hyperlink(t0.rows[1].cells[2], url, sample_id)

out_docx_test = "scratch/test_table_link.docx"
doc.save(out_docx_test)
print(f"Saved test docx with hyperlink to {out_docx_test}")

# Export via Word COM
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    doc_word = word.Documents.Open(os.path.abspath(out_docx_test))
    out_pdf_test = os.path.abspath("scratch/test_table_link.pdf")
    doc_word.ExportAsFixedFormat(out_pdf_test, 17) # wdExportFormatPDF = 17
    doc_word.Close(False)
    print(f"Exported test PDF via Word COM to {out_pdf_test}")
finally:
    word.Quit()

# Inspect links in exported PDF
d = fitz.open("scratch/test_table_link.pdf")
print("Total pages in exported test PDF:", len(d))
p0 = d[0]
links = list(p0.get_links())
print(f"Links found on Page 1: {len(links)}")
for l in links:
    print("  Link:", l)
