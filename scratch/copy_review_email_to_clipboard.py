import os, sys
import win32clipboard
import docx

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
html_path = os.path.join(desktop_dir, "OOS-261242_Review_Email_Robin_Format.html")
docx_path = os.path.join(desktop_dir, "OOS-261242 Review Email.docx")

# Read HTML
with open(html_path, "r", encoding="utf-8") as f:
    raw_html = f.read()

# Build standard Windows CF_HTML payload
# Format required:
# Version:0.9
# StartHTML:xxxxxxxx
# EndHTML:xxxxxxxx
# StartFragment:xxxxxxxx
# EndFragment:xxxxxxxx
# <html>...<!--StartFragment-->...<!--EndFragment-->...</html>

fragment_start_tag = "<!--StartFragment-->"
fragment_end_tag = "<!--EndFragment-->"

idx_start = raw_html.find(fragment_start_tag)
idx_end = raw_html.find(fragment_end_tag)

if idx_start == -1 or idx_end == -1:
    body_content = raw_html
    raw_html = f"<html><body><!--StartFragment-->{body_content}<!--EndFragment--></body></html>"
    idx_start = raw_html.find(fragment_start_tag)
    idx_end = raw_html.find(fragment_end_tag)

header_template = (
    "Version:0.9\r\n"
    "StartHTML:00000000\r\n"
    "EndHTML:00000000\r\n"
    "StartFragment:00000000\r\n"
    "EndFragment:00000000\r\n"
)

# Dummy offset length
header_len = len(header_template)

# Everything in bytes (UTF-8)
html_bytes = raw_html.encode("utf-8")
start_html = header_len
end_html = header_len + len(html_bytes)
start_fragment = header_len + len(raw_html[:idx_start + len(fragment_start_tag)].encode("utf-8"))
end_fragment = header_len + len(raw_html[:idx_end].encode("utf-8"))

header = (
    f"Version:0.9\r\n"
    f"StartHTML:{start_html:08d}\r\n"
    f"EndHTML:{end_html:08d}\r\n"
    f"StartFragment:{start_fragment:08d}\r\n"
    f"EndFragment:{end_fragment:08d}\r\n"
)

cf_html_data = header.encode("utf-8") + html_bytes

# Extract Plain Text from docx
doc = docx.Document(docx_path)
text_parts = []
for p in doc.paragraphs:
    if p.text.strip():
        text_parts.append(p.text)
for t in doc.tables:
    for row in t.rows:
        text_parts.append("\t".join([c.text.strip().replace("\n", " ") for c in row.cells]))
plain_text = "\n\n".join(text_parts)

# Write to Clipboard
CF_HTML = win32clipboard.RegisterClipboardFormat("HTML Format")

win32clipboard.OpenClipboard(0)
try:
    win32clipboard.EmptyClipboard()
    # Set CF_HTML
    win32clipboard.SetClipboardData(CF_HTML, cf_html_data)
    # Set Unicode Plain Text
    win32clipboard.SetClipboardData(win32clipboard.CF_UNICODETEXT, plain_text)
    print("SUCCESS: Copied rich HTML and Plain Text to Windows Clipboard!")
finally:
    win32clipboard.CloseClipboard()

# Verify clipboard contents
win32clipboard.OpenClipboard(0)
try:
    has_html = win32clipboard.IsClipboardFormatAvailable(CF_HTML)
    has_text = win32clipboard.IsClipboardFormatAvailable(win32clipboard.CF_UNICODETEXT)
    print(f"Verification: has_html={has_html}, has_text={has_text}")
finally:
    win32clipboard.CloseClipboard()
