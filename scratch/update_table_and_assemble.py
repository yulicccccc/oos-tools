import os, sys, shutil
from datetime import datetime
import fitz
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client as win32

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docs_dir = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
history_dir = os.path.join(docs_dir, ".history")
scratch_dir = os.path.join(docs_dir, "scratch")

# Original backup of analyst's unedited table docx
orig_table_docx = os.path.join(history_dir, "EM_table_261242_orig_backup.docx")

print("=== STEP 1: PREPARE AMENDED TABLE DOCX ===")
doc_tbl = Document(orig_table_docx)

def update_cell_text_preserving_style(cell, new_text):
    p = cell.paragraphs[0]
    font_name = p.runs[0].font.name if p.runs and p.runs[0].font.name else 'Times New Roman'
    font_size = p.runs[0].font.size if p.runs and p.runs[0].font.size else Pt(6.0)
    bold = p.runs[0].bold if p.runs else False
    italic = p.runs[0].italic if p.runs else False
    align = p.alignment
    # remove extra paragraphs in cell
    for extra_p in cell.paragraphs[1:]:
        p_elem = extra_p._p
        p_elem.getparent().remove(p_elem)
    # reset first paragraph
    p.text = ''
    p.alignment = align
    lines = new_text.split('\n')
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.alignment = align
        r = p.add_run(line)
        r.font.name = font_name
        r.font.size = font_size
        r.bold = bold
        r.italic = italic

# Table 0 (Table 1): Read Dates & Incubation Observation
t0 = doc_tbl.tables[0]
# R1 C1: SMO 14MAY 2026 (remove excessive spaces)
update_cell_text_preserving_style(t0.rows[1].cells[1], 'SMO\n14MAY 2026')
# R1 C4: 11 CFUs on active air sample plate (fix missing space)
update_cell_text_preserving_style(t0.rows[1].cells[4], '11 CFUs on\nactive air sample\nplate')

# Table 1 (Table 2): Environmental Monitoring Plates for Analyst and Cleanroom Bracketing
t1 = doc_tbl.tables[1]
# R4 C4: Week After Testing Date (fix extra space)
update_cell_text_preserving_style(t1.rows[4].cells[4], 'Week After\nTesting Date')
# R7 C5: 1 CFU on table in 115 ISO 8 and 1 CFU on cart in 115 ISO 8 (fix plural typos & space)
update_cell_text_preserving_style(t1.rows[7].cells[5], '1 CFU on table in 115 ISO 8 and 1 CFU on cart in 115 ISO 8')
# R8 C4: Week After Testing Date (fix extra space)
update_cell_text_preserving_style(t1.rows[8].cells[4], 'Week After\nTesting Date')

# Save amended Word table
out_table_docx_1 = os.path.join(desktop_dir, "EM table OOS-261242 14MAY2026.docx")
out_table_docx_2 = os.path.join(desktop_dir, "Tables OOS-261242 EM SMO 115B Air 14MAY2026 - EM.docx")
doc_tbl.save(out_table_docx_1)
doc_tbl.save(out_table_docx_2)
print("Saved amended table DOCX files.")

# Convert to 1-page PDF via Word COM
out_table_pdf_1 = os.path.join(desktop_dir, "EM table OOS-261242 14MAY2026.pdf")
out_table_pdf_2 = os.path.join(desktop_dir, "Tables OOS-261242 EM SMO 115B Air 14MAY2026 - EM.pdf")

word_app = win32.Dispatch("Word.Application")
word_app.Visible = False
try:
    doc_com = word_app.Documents.Open(os.path.abspath(out_table_docx_1))
    doc_com.SaveAs(os.path.abspath(out_table_pdf_1), FileFormat=17)
    doc_com.Close(False)
    shutil.copy2(out_table_pdf_1, out_table_pdf_2)
    print("Converted amended table to PDF via Word COM.")
finally:
    word_app.Quit()

doc_check = fitz.open(out_table_pdf_1)
print(f"Amended Table PDF page count: {len(doc_check)} (Must be 1)")
assert len(doc_check) == 1, "Table PDF must be exactly 1 page!"

# === STEP 2: LOAD USER'S LATEST EDITED PDF ===
print("\n=== STEP 2: LOAD USER'S EDITED PDF ===")
user_pdf_path = os.path.join(desktop_dir, "OOS-261242 EM SMO 115B Air 14MAY2026 - EM.pdf")

# Open user's PDF in memory
user_doc = fitz.open(user_pdf_path)
print(f"User PDF page count: {len(user_doc)}")

# Create a clean new PDF containing Pages 1 to 7 from user's PDF
final_doc = fitz.open()
final_doc.insert_pdf(user_doc, from_page=0, to_page=min(6, len(user_doc)-1))
print(f"Pages 1-7 inserted. Count: {len(final_doc)}")

# Append Page 8: The newly converted 1-page amended table PDF
table_doc = fitz.open(out_table_pdf_1)
final_doc.insert_pdf(table_doc)
print(f"Table inserted as Page 8. Total final pages: {len(final_doc)}")
assert len(final_doc) == 8, f"Expected 8 pages, got {len(final_doc)}"

# Save to scratch and writable desktop targets
scratch_final_pdf = os.path.join(scratch_dir, "assembled_261242_final.pdf")
final_doc.save(scratch_final_pdf)

std_pdf_1 = os.path.join(desktop_dir, "OOS-261242.pdf")
std_pdf_2 = os.path.join(desktop_dir, "OOS-261242 (1).pdf")
shutil.copy2(scratch_final_pdf, std_pdf_1)
shutil.copy2(scratch_final_pdf, std_pdf_2)
print(f"Saved synchronized 8-page PDFs to: {std_pdf_1} and {std_pdf_2}")

# Try to overwrite user's original PDF if unlocked
user_doc.close()
table_doc.close()
final_doc.close()

try:
    shutil.copy2(scratch_final_pdf, user_pdf_path)
    print(f"SUCCESS: Directly overwritten user PDF: {user_pdf_path}")
except Exception as e:
    alt_user_pdf = os.path.join(desktop_dir, "OOS-261242 EM SMO 115B Air 14MAY2026 - EM_Updated.pdf")
    shutil.copy2(scratch_final_pdf, alt_user_pdf)
    print(f"Notice: User PDF is open in Foxit/Reader ({e}). Saved updated version to: {alt_user_pdf}")

# Render preview of Page 8 for verification
verify_doc = fitz.open(scratch_final_pdf)
p8 = verify_doc[7]
pix = p8.get_pixmap(dpi=150)
pix_path = os.path.join(scratch_dir, "assembled_261242_p8_fixed.png")
pix.save(pix_path)
print(f"Saved Page 8 verification image to: {pix_path}")
