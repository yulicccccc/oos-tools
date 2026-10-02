import os
import sys
import shutil
import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import fitz

sys.stdout.reconfigure(encoding='utf-8')

DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
DOCS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
SCRATCH_DIR = os.path.join(DOCS_DIR, "scratch")
HISTORY_DIR = os.path.join(DOCS_DIR, ".history")

os.makedirs(SCRATCH_DIR, exist_ok=True)
os.makedirs(HISTORY_DIR, exist_ok=True)

URL = "https://etrax.eagleanalytical.com/SubmissionTest/Details/jzraMOOTFYUFD%24erkwgubw__"
SAMPLE_ID = "ETX-260828-0527"

def set_cell_hyperlink(cell, url, text):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    part = p.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink_xml = (
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rStyle w:val="Hyperlink"/>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="0000FF"/>'
        f'<w:sz w:val="15"/>'
        f'<w:szCs w:val="15"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    p._p.append(parse_xml(hyperlink_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

print("=== STEP 1: UPDATE STANDALONE TABLES DOCX ===")
tables_docx_path = os.path.join(DESKTOP_DIR, "Celsis table OOS-262080.docx")
tables_docx_scratch = os.path.join(SCRATCH_DIR, "Celsis table OOS-262080.docx")

doc_tbl = docx.Document(tables_docx_path if os.path.exists(tables_docx_path) else tables_docx_scratch)
t0 = doc_tbl.tables[0]
print(f"Table 0 header: {[c.text.strip() for c in t0.rows[0].cells]}")
print(f"Table 0 row 1 before: {[c.text.strip() for c in t0.rows[1].cells]}")

# Cell 2 is Sample ID
set_cell_hyperlink(t0.rows[1].cells[2], URL, SAMPLE_ID)
print(f"Set hyperlink in Table 0, Cell 2 -> {SAMPLE_ID} ({URL})")

doc_tbl.save(tables_docx_scratch)
try:
    doc_tbl.save(tables_docx_path)
    print(f"Saved: {tables_docx_path}")
except PermissionError:
    print(f"Notice: {tables_docx_path} is locked by Word, saved to scratch.")

print("\n=== STEP 2: UPDATE MASTER WORD REPORT DOCX ===")
main_docx_path = os.path.join(DESKTOP_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis.docx")
main_docx_scratch = os.path.join(SCRATCH_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis.docx")

if os.path.exists(main_docx_path):
    doc_main = docx.Document(main_docx_path)
    t1_main = doc_main.tables[1]
    print(f"Master Doc Table 1 header: {[c.text.strip() for c in t1_main.rows[0].cells]}")
    set_cell_hyperlink(t1_main.rows[1].cells[2], URL, SAMPLE_ID)
    print(f"Set hyperlink in Master Doc Table 1, Cell 2 -> {SAMPLE_ID}")
    doc_main.save(main_docx_scratch)
    try:
        doc_main.save(main_docx_path)
        print(f"Saved: {main_docx_path}")
    except PermissionError:
        print(f"Notice: {main_docx_path} is locked, saved to scratch.")

print("\n=== STEP 3: EXPORT TABLES DOCX TO PDF VIA WORD COM ===")
tables_pdf_scratch = os.path.join(SCRATCH_DIR, "Celsis table OOS-262080.pdf")
tables_pdf_path = os.path.join(DESKTOP_DIR, "Celsis table OOS-262080.pdf")

word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    source_to_open = os.path.abspath(tables_docx_scratch if os.path.exists(tables_docx_scratch) else tables_docx_path)
    doc_word = word.Documents.Open(source_to_open)
    doc_word.ExportAsFixedFormat(os.path.abspath(tables_pdf_scratch), 17) # wdExportFormatPDF = 17
    doc_word.Close(False)
    print(f"Exported PDF via Word COM to {tables_pdf_scratch}")
finally:
    word.Quit()

# Verify link in tables_pdf_scratch
chk_tbl = fitz.open(tables_pdf_scratch)
print(f"Tables PDF page count: {len(chk_tbl)}")
links_p0 = list(chk_tbl[0].get_links())
print(f"Links on Page 1: {len(links_p0)}")
for l in links_p0:
    print(f"  URI: {l.get('uri')}, Rect: {l.get('from')}")
assert len(links_p0) >= 1, "Expected at least 1 hyperlink on Page 1 of Table PDF!"
chk_tbl.close()

try:
    shutil.copy2(tables_pdf_scratch, tables_pdf_path)
    print(f"Saved: {tables_pdf_path}")
except PermissionError:
    print(f"Notice: {tables_pdf_path} is locked.")

print("\n=== STEP 4: UPDATE COMBINED DELIVERABLE PDFS ===")
# 1. Update OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis - QYC.pdf
qyc_pdf_path = os.path.join(DESKTOP_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis - QYC.pdf")
if os.path.exists(qyc_pdf_path):
    doc_qyc_orig = fitz.open(qyc_pdf_path)
    doc_qyc_new = fitz.open()
    # Keep Pages 1 to 6 from user's edited QYC.pdf
    doc_qyc_new.insert_pdf(doc_qyc_orig, from_page=0, to_page=5)
    doc_qyc_orig.close()
    
    # Append the 2 table pages with active link
    doc_tbl_pdf = fitz.open(tables_pdf_scratch)
    doc_qyc_new.insert_pdf(doc_tbl_pdf, annots=True)
    doc_tbl_pdf.close()
    
    scratch_qyc = os.path.join(SCRATCH_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis - QYC.pdf")
    doc_qyc_new.save(scratch_qyc)
    doc_qyc_new.close()
    
    try:
        shutil.copy2(scratch_qyc, qyc_pdf_path)
        print(f"Updated QYC.pdf cleanly to 8 pages with active link: {qyc_pdf_path}")
    except PermissionError:
        print(f"Notice: {qyc_pdf_path} is locked.")

# 2. Update OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (2).pdf
c2_pdf_path = os.path.join(DESKTOP_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (2).pdf")
if os.path.exists(c2_pdf_path):
    doc_c2_orig = fitz.open(c2_pdf_path)
    doc_c2_new = fitz.open()
    doc_c2_new.insert_pdf(doc_c2_orig, from_page=0, to_page=5)
    doc_c2_orig.close()
    
    doc_tbl_pdf = fitz.open(tables_pdf_scratch)
    doc_c2_new.insert_pdf(doc_tbl_pdf, annots=True)
    doc_tbl_pdf.close()
    
    scratch_c2 = os.path.join(SCRATCH_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (2).pdf")
    doc_c2_new.save(scratch_c2)
    doc_c2_new.close()
    try:
        shutil.copy2(scratch_c2, c2_pdf_path)
        print(f"Updated (2).pdf to 8 pages with active link: {c2_pdf_path}")
    except PermissionError:
        print(f"Notice: {c2_pdf_path} is locked.")

# 3. Update OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (Complete).pdf
complete_pdf_path = os.path.join(DESKTOP_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (Complete).pdf")
base_corp = os.path.join(DESKTOP_DIR, "CORP-FORM-21 Laboratory OOS Investigation Form (v11.1).pdf")
doc_comp_new = fitz.open()
doc_base = fitz.open(base_corp)
doc_comp_new.insert_pdf(doc_base)
doc_base.close()

doc_tbl_pdf = fitz.open(tables_pdf_scratch)
doc_comp_new.insert_pdf(doc_tbl_pdf, annots=True)
doc_tbl_pdf.close()

scratch_comp = os.path.join(SCRATCH_DIR, "OOS-262080 Optimal Balance Pharmacy (E19193) - Celsis (Complete).pdf")
doc_comp_new.save(scratch_comp)
doc_comp_new.close()
try:
    shutil.copy2(scratch_comp, complete_pdf_path)
    print(f"Updated Complete.pdf to 8 pages with active link: {complete_pdf_path}")
except PermissionError:
    print(f"Notice: {complete_pdf_path} is locked.")

print("\n=== STEP 5: VERIFY ACTIVE LINKS ACROSS DELIVERABLES ===")
for p_check in [tables_pdf_path, qyc_pdf_path, complete_pdf_path]:
    if os.path.exists(p_check):
        d = fitz.open(p_check)
        # Table 1 is on page 1 of table PDF or page 7 of 8-page PDF
        tbl_page_idx = 0 if len(d) == 2 else 6
        links = list(d[tbl_page_idx].get_links())
        print(f"File: {os.path.basename(p_check)} (Pages: {len(d)}) -> Table 1 Page {tbl_page_idx+1} Links: {len(links)}")
        for l in links:
            print(f"   Target URI: {l.get('uri')}")
        d.close()

# Render verification preview of Table 1
d_vis = fitz.open(tables_pdf_scratch)
p_vis = d_vis[0]
pix = p_vis.get_pixmap(dpi=150, clip=fitz.Rect(50, 40, 560, 160))
pix_path = os.path.join(SCRATCH_DIR, "verified_table1_hyperlink.png")
pix.save(pix_path)
print(f"Rendered Table 1 preview image to: {pix_path}")

print("\n=== HYPERLINK INJECTION COMPLETE! ===")
