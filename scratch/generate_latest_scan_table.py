import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
SCRATCH_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\scratch"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=50, bottom=50, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_hyperlink(paragraph, url, text, font_size=Pt(7), color_rgb="2980B9"):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color_rgb}"/>'))
    rPr.append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
    new_run.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def format_cell(cell, text, bold=False, italic=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if fill_hex:
        set_cell_background(cell, fill_hex)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = font_size
    r.bold = bold
    r.italic = italic
    return r

def build_tables():
    doc = docx.Document()
    
    # Page setup: Letter, Portrait, Margins = 0.5 in
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # ---------------- Table 1 ----------------
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(4)
    r_t1 = p_t1.add_run("Table 1: Information for ETX-260813-0778 under investigation")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(9)
    r_t1.underline = True

    t1 = doc.add_table(rows=2, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = [
        "Processing Analyst",
        "Reading Analyst",
        "Sample ID",
        "Events",
        "Confirmed Microbial\nEvents",
        "Morphology Description"
    ]
    t1_widths = [Inches(1.25), Inches(1.25), Inches(1.1), Inches(0.7), Inches(1.2), Inches(2.0)]

    for j, h in enumerate(t1_headers):
        cell = t1.cell(0, j)
        format_cell(cell, h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None)

    # Row 1 values
    format_cell(t1.cell(1, 0), "Varsha Subramanian", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 1), "Sonal Uprety", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Sample ID with hyperlink
    c_sid = t1.cell(1, 2)
    c_sid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sid)
    p_sid = c_sid.paragraphs[0]
    p_sid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sid.paragraph_format.space_before = Pt(0)
    p_sid.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sid, "https://etrax.eagleanalytical.com/SubmissionTest/Details/XgZ%241BCFiwnbIs4lw0bNgw__", "ETX-260813-0778", font_size=Pt(7))

    format_cell(t1.cell(1, 3), "31", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 4), "4", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 5), "Short rod-shaped morphology", bold=True, italic=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in t1.rows:
        for j, w in enumerate(t1_widths):
            row.cells[j].width = w

    # Spacing between tables
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(12)
    p_space.paragraph_format.space_after = Pt(0)

    # ---------------- Table 2 ----------------
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(4)
    # Using our agreed standard format DDMMMYY: 25Aug26
    r_t2 = p_t2.add_run("Table 2: Environmental Monitoring from Testing Performed on 25Aug26")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(9)
    r_t2.underline = True

    t2 = doc.add_table(rows=0, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    t2_widths = [
        Inches(1.05), # Sampling site
        Inches(0.65), # Frequency
        Inches(0.70), # Date (DDMMMYY)
        Inches(0.55), # Analyst
        Inches(1.15), # Day / Week(s)
        Inches(1.10), # Observation
        Inches(0.85), # EM Plate ETX Number
        Inches(0.85), # Microbial ID
        Inches(0.60)  # Notes
    ]

    # Header row
    hdr_row = t2.add_row()
    t2_headers = [
        "Environmental\nMonitoring (EM)\nSampling Site",
        "Frequency",
        "Date\n(DDMMMYY)",
        "Analyst",
        "Day /Week(s)",
        "Observation",
        "EM Plate ETX\nNumber",
        "Microbial ID",
        "Notes"
    ]
    for j, h in enumerate(t2_headers):
        cell = hdr_row.cells[j]
        format_cell(cell, h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None)

    def add_section_divider(table, title_text):
        row = table.add_row()
        # Merge all 9 cells
        cell = row.cells[0]
        for c in row.cells[1:]:
            cell.merge(c)
        format_cell(cell, title_text, bold=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex="E8E8E8")
        return row

    def add_data_row(table, site, freq, date_str, analyst, day_week, obs, etx_num, micro_id, notes):
        row = table.add_row()
        format_cell(row.cells[0], site, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], freq, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[2], date_str, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[3], analyst, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[4], day_week, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[5], obs, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[6], etx_num, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[7], micro_id, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[8], notes, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        return row

    # --- Section 1: Personnel EM ---
    add_section_divider(t2, "Personnel EM for 25Aug26")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t2, "Biological Safety Cabinet (BSC) EM for BSC E001319 for 25Aug26")
    add_data_row(t2, "Surface Sampling of\nISO 5 BSC E001319\n(4 Locations)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Surface Sampling of\nISO 5 BSC E001319\n(4 Locations)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of\nISO 5 BSC E001319", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan)", "[Pending]", "[Pending]", "Not Applicable", "Pending scan")
    add_data_row(t2, "Settling Sampling of\nISO 5 BSC E001319", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan C/O)", "[Pending]", "[Pending]", "Not Applicable", "Pending scan")

    # --- Section 3: Weekly Active Air ---
    add_section_divider(t2, "Weekly Active Air Sampling of CR 145 (E001979) with Processing BSC for 25Aug26")
    add_data_row(t2, "Active Air Sampling\nof Cleanrooms", "Weekly", "28Aug26", "SMO", "Week of Testing", "5 CFU\n(ISO 8 142)", "ETX-260908-0584", "Pending", "None")

    # --- Section 4: Weekly Surface ---
    add_section_divider(t2, "Weekly Surface Sampling of CR 145 (E001979) with Processing BSC for 25Aug26")
    add_data_row(t2, "Surface Sampling of\nCleanrooms", "Weekly", "28Aug26", "SMO", "Week of Testing", "1 CFU\n(ISO 8 143)", "ETX-260908-0580", "Pending", "None")

    # Set column widths across all rows
    for row in t2.rows:
        if len(row.cells) == 9:
            for j, w in enumerate(t2_widths):
                row.cells[j].width = w

    return doc

def export_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
    finally:
        word.Quit()
    print(f"Exported to PDF: {pdf_path}")

if __name__ == "__main__":
    doc = build_tables()
    
    out_docx_desktop = os.path.join(DESKTOP_DIR, "Tables OOS-261967 TAM Pharmacy (E74685) - ScanRDI.docx")
    out_pdf_desktop = os.path.join(DESKTOP_DIR, "Tables OOS-261967 TAM Pharmacy (E74685) - ScanRDI.pdf")
    
    out_docx_docs = os.path.join(DOCUMENTS_DIR, "Tables OOS-261967 TAM Pharmacy (E74685) - ScanRDI.docx")
    out_pdf_docs = os.path.join(DOCUMENTS_DIR, "Tables OOS-261967 TAM Pharmacy (E74685) - ScanRDI.pdf")

    doc.save(out_docx_desktop)
    print(f"Saved DOCX: {out_docx_desktop}")
    export_to_pdf(out_docx_desktop, out_pdf_desktop)

    doc.save(out_docx_docs)
    print(f"Saved DOCX: {out_docx_docs}")
    import shutil
    shutil.copy2(out_pdf_desktop, out_pdf_docs)
    print(f"Saved PDF: {out_pdf_docs}")

    # Render preview image
    doc_fitz = fitz.open(out_pdf_desktop)
    pix = doc_fitz[0].get_pixmap(dpi=150)
    preview_img = os.path.join(SCRATCH_DIR, "latest_tables_preview.png")
    pix.save(preview_img)
    doc_fitz.close()
    print(f"Preview image saved: {preview_img}")
