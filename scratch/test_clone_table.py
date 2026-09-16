import docx
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
import win32com.client
from pypdf import PdfReader

doc_src = docx.Document(rC:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\Tables OOS-261814 GoGoMeds Select (E10747) - ScanRDI.docx)
doc_em = docx.Document(rC:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\EM Table OOS-261814 07AUG2026.docx)

# Remove existing Table 2 from doc_src
t2_old = doc_src.tables[1]
t2_old._element.getparent().remove(t2_old._element)

# Clone Table from doc_em into doc_src
t_new_elem = doc_em.tables[0]._element
import copy
t_copy = copy.deepcopy(t_new_elem)
doc_src._body._element.append(t_copy)

# Save test docx
test_docx = rscratch\test_tables_17rows.docx
test_pdf = rscratch\test_tables_17rows.pdf
doc_src.save(test_docx)
print(Saved cloned test docx)

# Convert to PDF
word = win32com.client.DispatchEx(Word.Application)
word.Visible = False
try:
    d = word.Documents.Open(rC:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS" + test_docx, ReadOnly=True)
 d.SaveAs(rC:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS" + test_pdf, FileFormat=17)
    d.Close(SaveChanges=False)
    print(Converted test docx to PDF)
finally:
    word.Quit()

reader = PdfReader(test_pdf)
print(fTest PDF page count: {len(reader.pages)})

