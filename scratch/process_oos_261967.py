"""
Script: process_oos_261967.py
Purpose: Complete generation of Scan RDI OOS-261967 report files, tables, and 7-page PDF package.
Client: TAM Pharmacy (E74685)
Sample ID: ETX-260813-0778
Lot: LG403000087
Test Date: 25Aug26
"""

import os
import sys
import shutil
import json
from datetime import datetime
import fitz  # PyMuPDF
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docxtpl import DocxTemplate
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"

OOS_ID = "261967"
CLIENT_NAME = "TAM Pharmacy (E74685)"
SAMPLE_ID = "ETX-260813-0778"
SAMPLE_URL = "https://etrax.eagleanalytical.com/SubmissionTest/Details/XgZ%241BCFiwnbIs4lw0bNgw__"
SAMPLE_NAME = "Semaglutide/Pyridoxine 2.5mg/10mg/mL"
LOT_NUMBER = "LG403000087"
DOSAGE_FORM = "Injectable"
TEST_DATE = "25Aug26"
TEST_DATE_LONG = "25-Aug-2026"

# ==========================================
# 1. BUILD WORD & PDF TABLES (PAGE 7)
# ==========================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=50, bottom=50, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_hyperlink(paragraph, url, text, font_size=Pt(7), color_rgb="2980B9"):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color_rgb}"/>'))
    rPr.append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
    new_run.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def format_cell(cell, text, bold=False, italic=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if fill_hex:
        set_cell_background(cell, fill_hex)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = font_size
    r.bold = bold
    r.italic = italic
    return r

def build_standalone_tables_doc():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # ---------------- Table 1 ----------------
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(4)
    r_t1 = p_t1.add_run(f"Table 1: Information for {SAMPLE_ID} under investigation")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(9)
    r_t1.underline = True

    t1 = doc.add_table(rows=2, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = [
        "Processing Analyst",
        "Reading Analyst",
        "Sample ID",
        "Events",
        "Confirmed Microbial\nEvents",
        "Morphology Description"
    ]
    t1_widths = [Inches(1.25), Inches(1.25), Inches(1.1), Inches(0.7), Inches(1.2), Inches(2.0)]

    for j, h in enumerate(t1_headers):
        cell = t1.cell(0, j)
        format_cell(cell, h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    format_cell(t1.cell(1, 0), "Varsha Subramanian", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 1), "Sonal Uprety", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    c_sid = t1.cell(1, 2)
    c_sid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sid)
    p_sid = c_sid.paragraphs[0]
    p_sid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sid.paragraph_format.space_before = Pt(0)
    p_sid.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sid, SAMPLE_URL, SAMPLE_ID, font_size=Pt(7))

    format_cell(t1.cell(1, 3), "31", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 4), "4", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 5), "Short rod-shaped morphology", bold=True, italic=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in t1.rows:
        for j, w in enumerate(t1_widths):
            row.cells[j].width = w

    # Spacing between tables
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(12)
    p_space.paragraph_format.space_after = Pt(0)

    # ---------------- Table 2 ----------------
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(4)
    r_t2 = p_t2.add_run(f"Table 2: Environmental Monitoring from Testing Performed on {TEST_DATE}")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(9)
    r_t2.underline = True

    t2 = doc.add_table(rows=0, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    t2_widths = [
        Inches(1.05), # Sampling site
        Inches(0.65), # Frequency
        Inches(0.70), # Date (DDMMMYY)
        Inches(0.55), # Analyst
        Inches(1.15), # Day / Week(s)
        Inches(1.10), # Observation
        Inches(0.85), # EM Plate ETX Number
        Inches(0.85), # Microbial ID
        Inches(0.60)  # Notes
    ]

    hdr_row = t2.add_row()
    t2_headers = [
        "Environmental\nMonitoring (EM)\nSampling Site",
        "Frequency",
        "Date\n(DDMMMYY)",
        "Analyst",
        "Day /Week(s)",
        "Observation",
        "EM Plate ETX\nNumber",
        "Microbial ID",
        "Notes"
    ]
    for j, h in enumerate(t2_headers):
        format_cell(hdr_row.cells[j], h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    def add_section_divider(table, title_text):
        row = table.add_row()
        cell = row.cells[0]
        for c in row.cells[1:]:
            cell.merge(c)
        format_cell(cell, title_text, bold=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex="E8E8E8")
        return row

    def add_data_row(table, site, freq, date_str, analyst, day_week, obs, etx_num, micro_id, notes):
        row = table.add_row()
        format_cell(row.cells[0], site, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], freq, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[2], date_str, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[3], analyst, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[4], day_week, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[5], obs, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[6], etx_num, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[7], micro_id, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[8], notes, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        return row

    # --- Section 1: Personnel EM ---
    add_section_divider(t2, f"Personnel EM for {TEST_DATE}")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t2, f"Biological Safety Cabinet (BSC) EM for BSC E001319 for {TEST_DATE}")
    add_data_row(t2, "Surface Sampling of\nISO 5 BSC E001319\n(4 Locations)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Surface Sampling of\nISO 5 BSC E001319\n(4 Locations)", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of\nISO 5 BSC E001319", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of\nISO 5 BSC E001319", "Daily", "25Aug26", "VV", "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 3: Weekly Active Air ---
    add_section_divider(t2, f"Weekly Active Air Sampling of CR 145 (E001979) with Processing BSC for {TEST_DATE}")
    add_data_row(t2, "Active Air Sampling\nof Cleanrooms", "Weekly", "28Aug26", "SMO", "Week of Testing", "5 CFU\n(ISO 8 142)", "ETX-260908-0584", "Pending", "None")

    # --- Section 4: Weekly Surface ---
    add_section_divider(t2, f"Weekly Surface Sampling of CR 145 (E001979) with Processing BSC for {TEST_DATE}")
    add_data_row(t2, "Surface Sampling of\nCleanrooms", "Weekly", "28Aug26", "SMO", "Week of Testing", "1 CFU\n(ISO 8 143)", "ETX-260908-0580", "Pending", "None")

    for row in t2.rows:
        if len(row.cells) == 9:
            for j, w in enumerate(t2_widths):
                row.cells[j].width = w

    return doc, t2

def export_docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
    finally:
        word.Quit()
    print(f"Exported to PDF: {pdf_path}")


# ==========================================
# 2. GENERATE COMPLETE INVESTIGATION NARRATIVE
# ==========================================
p1 = (
    "All analysts involved in the prepping, processing, changeover and reading of the sample – "
    "Min Jang, Varsha Subramanian, and Sonal Uprety – were interviewed and their answers are recorded throughout this document."
)

p2 = (
    "The sample was stored upon arrival according to the Client’s instructions. "
    "Analysts Min Jang and Varsha Subramanian confirmed the integrity of the sample vials throughout both the preparation and "
    "processing stages. No leaks or turbidity were observed at any point, verifying the integrity of the sample."
)

p3 = (
    "All reagents and supplies mentioned in the material section above were stored according to the suppliers’ recommendations, "
    "and their integrity was visually verified before utilization. Moreover, each reagent and supply had valid expiration dates. "
    "During the preparation phase on 24Aug26, Min Jang disinfected the samples using acidified bleach and placed them into a "
    "pre-disinfected storage bin."
)

p4 = (
    "On 25Aug26, prior to sample processing, Varsha Subramanian performed a second disinfection with acidified bleach, allowing a "
    "minimum contact time of 10 minutes before transferring the samples into the ISO 8 cleanroom (143). A final disinfection step was "
    "completed immediately before the samples were introduced into the ISO 5 Biological Safety Cabinet (BSC), E001319, located "
    "within the innermost ISO 7 cleanroom (145). All activities were performed in accordance with MICRO-SOP-12, Rapid Scan RDI® Test using FIFU Method."
)

p5 = (
    "The cleanroom suite used for testing (L-Suite, CR145, E001979) consists of four interconnected sections: the innermost ISO 7 cleanroom (145), "
    "which opens into the adjacent ISO 7 buffer cleanroom (144), followed by ISO 8 anteroom (143) and the outermost ISO 8 room (142). "
    "A positive air pressure system is maintained throughout the suite to ensure controlled, unidirectional airflow cascading outward from "
    "145 through 144 and 143 into 142. The ISO 5 BSC E001319, located in the innermost ISO 7 room (145), was used for both the pre-labelling "
    "and filtration steps and for the changeover procedure. It was thoroughly cleaned and disinfected prior to procedures in accordance "
    "with MICRO-SOP-9, Cleaning and Disinfecting Procedure for Microbiology. Furthermore, the biosafety cabinet (BSC E001319) utilized "
    "during testing was current in certification and had been approved for use by both the Engineering and Quality Assurance teams."
)

p6 = (
    "Pre-labeling and filtration activities were performed by Varsha Subramanian on 25Aug26 within the ISO 5 biosafety cabinet (BSC E001319) "
    "located in the innermost ISO 7 cleanroom (145). The changeover procedure was subsequently performed by the same analyst on the same date "
    "within the same ISO 5 biosafety cabinet (BSC E001319). The reading analyst, Sonal Uprety, confirmed that the cytometer Cs2-105 (E001040) "
    "was set up as per ENG-SOP-4 (Scan RDI® System – Operations (Standard C3 Quality Check and Microscope Setup and Maintenance), and the "
    "negative control (0 events) and positive control (Clostridium sporogenes, Lot 05282026-19404-CS) for analyst Varsha Subramanian yielded "
    "expected results. On 25Aug26, a rapid sterility test was performed on the sample using the ScanRDI method. The sample was initially prepared "
    "by analyst Min Jang, processed by Varsha Subramanian and subsequently read by Sonal Uprety. The test revealed 31 total events and 4 confirmed "
    "microbial events exhibiting short rod-shaped morphology, see Table 1."
)

p7 = (
    f"Table 2 (see attached tables) presents the environmental monitoring results for {SAMPLE_ID}. The environmental monitoring (EM) plates "
    "were incubated for no less than 48 hours at 30–35°C and for no less than an additional five days at 20–25°C, as per MICRO-SOP-2 "
    "(Environmental Monitoring of the Cleanroom Facility). Upon review of the environmental monitoring data associated with the sterility test, "
    "no microbial growth was observed on the left and right personnel touch plates for processor and changeover analyst Varsha Subramanian. "
    "Additionally, no microbial growth was recovered from the surface contact plates (4 locations) or settling plates collected from BSC E001319 during both "
    "testing and changeover."
)

p8 = (
    "Weekly surface monitoring of the cleanroom suite conducted on 28Aug26 demonstrated no microbial recovery in the ISO 7 cleanroom (145) or "
    "ISO 7 buffer room (144); however, 1 CFU (ETX-260908-0580) was recovered from Table 1 with Scan Unit in the ISO 8 anteroom (143). "
    "Weekly active air monitoring conducted on 28Aug26 demonstrated no microbial recovery in ISO 7 cleanroom 145, ISO 7 buffer room 144, or "
    "ISO 8 anteroom 143; however, 5 CFUs (ETX-260908-0584) were recovered from the outermost ISO 8 room (142). It is important to note that "
    "all sample processing activities were performed strictly within the validated ISO 5 BSC E001319 located in the innermost ISO 7 cleanroom (145). "
    "The test samples do not come into contact with ambient ISO 8 air, as samples and supplies are transferred in disinfected, closed containers "
    "on carts through the layered cleanroom suites. Furthermore, the recoveries occurred in the lower-classified ISO 8 anteroom and outer room environments, "
    "which are physically segregated from the ISO 5 processing zone by closed doors and an outward-cascading positive air pressure gradient."
)

p9 = (
    "Furthermore, the absence of microbial recovery from analyst glove touch plates, settling plates, and BSC work surfaces confirms that no viable contamination "
    "transfer pathway existed from the room environment into the ISO 5 BSC. Based on the lack of detectable environmental contamination on critical "
    "surfaces and the controlled processing conditions, the cleanroom environment is not considered a likely source of contamination for the test sample."
)

p10 = (
    "Monthly cleaning and disinfection of the ISO 7 cleanrooms and the associated ISO 5 Biological Safety Cabinets (BSCs) were performed "
    "using hydrogen peroxide (H2O2) on 16 Aug 2026 in accordance with MICRO-SOP-9, Cleaning and Disinfecting Procedure for Microbiology. "
    "During both disinfection cycles, all H2O2 indicators met acceptance criteria and were documented as passing. These results confirm the "
    "effectiveness of the monthly cleaning and disinfection program for all areas of L-Suite (CR145), including the associated BSCs."
)

p11 = (
    f"Analyzing a 6-month sample history for {CLIENT_NAME}, this specific analyte '{SAMPLE_NAME}' has had no prior failures using the "
    "Scan RDI method during this period. To evaluate the potential for sample-to-sample contamination, all samples processed on the same "
    f"day were reviewed. The sample {SAMPLE_ID} was the 45th and final sample processed in a batch of 45 samples during session 082526-1040-2. "
    "All other 44 samples processed by the analyst in this session tested negative (0 confirmed microbial events), indicating that "
    f"cross-contamination is unlikely. Review of sample lot history indicates that there have been no other submissions of specific sample lot "
    f"{LOT_NUMBER} from the client for ScanRDI or any other sterility tests."
)

p12 = (
    "Based on the cumulative evidence, it is highly unlikely that the failing results were due to reagents, supplies, the cleanroom environment, "
    "the process, or analyst involvement. Consequently, the possibility of laboratory error contributing to this failure is minimal. "
    "Therefore, the original test result is deemed valid."
)

# Text Field 49 (Page 3/4)
smart_phase1_part1 = "\r \r".join([p1, p2, p3, p4, p5])
# Text Field 50 (Page 4/5)
smart_phase1_part2_page5 = "\r \r".join([p6, p7, p8])
# Text Field 51 (Page 5/6)
smart_phase1_part2_page6 = "\r \r".join([p9, p10, p11, p12])

# Complete Word narrative
smart_phase1_full = "\n\n".join([p1, p2, p3, p4, p5, p6, p7, p8, p9, p10, p11, p12])


# ==========================================
# 3. BUILD MASTER WORD INVESTIGATION REPORT
# ==========================================
def generate_master_word_report(tables_docx_path):
    tpl_path = os.path.join(ROOT_DIR, "ScanRDI OOS P1 template.docx")
    tpl = DocxTemplate(tpl_path)

    word_context = {
        'oos_id': OOS_ID,
        'client_name': CLIENT_NAME,
        'sample_id': SAMPLE_ID,
        'sample_name': SAMPLE_NAME,
        'lot_number': LOT_NUMBER,
        'dosage_form': DOSAGE_FORM,
        'test_date': TEST_DATE,
        'analyst_name': 'Varsha Subramanian',
        'analyst_initial': 'VV',
        'reader_name': 'Sonal Uprety',
        'changeover_initial': 'VV',
        'report_header': f"{SAMPLE_ID}\n\n{CLIENT_NAME}",
        'analyst_signature': 'Varsha Subramanian (Written by: Qiyue Chen)',
        'smart_personnel_block': 'Prepping Analyst: Min Jang (MJ)\nProcessing Analyst: Varsha Subramanian (VV)\nChangeover Analyst: Varsha Subramanian (VV)\nReading Analyst: Sonal Uprety (SU)',
        'smart_incident_opening': f"On {TEST_DATE}, sample {SAMPLE_ID} was found positive for viable microorganisms after ScanRDI testing.",
        'smart_comment_interview': "Yes, analysts Min Jang, Varsha Subramanian, and Sonal Uprety were interviewed comprehensively.",
        'smart_comment_samples': f"Yes, Sample ID: {SAMPLE_ID}",
        'smart_comment_records': "Yes, See 082526-1040-2 for more information.",
        'smart_comment_storage': f"Yes, Information is available in Eagle Trax Sample Location History under {SAMPLE_ID}",
        'bsc_id': '1319',
        'chgbsc_id': '1319',
        'cr_id': '1979',
        'cr_suit': '145',
        'smart_cr_id': 'CR145 (E001979) in L-Suite',
        'smart_scan_id': 'Cs2-105 (E001040)',
        'smart_bsc_bracketing_header': f"Biological Safety Cabinet (BSC) EM for BSC E001319 for {TEST_DATE}",
        'event_number': '31',
        'confirm_number': '4',
        'organism_morphology': 'Short rod-shaped morphology',
        'control_positive': 'C. sporogenes',
        'control_lot': '05282026-19404-CS',
        'control_data': '28May28',
        'date_of_weekly': '28Aug26',
        'weekly_initial': 'SMO',
        'obs_pers_dur': 'No Growth',
        'etx_pers_dur': 'Not Applicable',
        'id_pers_dur': 'Not Applicable',
        'note_pers': 'None',
        'obs_surf_dur': 'No Growth',
        'etx_surf_dur': 'Not Applicable',
        'id_surf_dur': 'Not Applicable',
        'note_surf': 'None',
        'obs_sett_dur': 'No Growth',
        'etx_sett_dur': 'Not Applicable',
        'id_sett_dur': 'Not Applicable',
        'note_sett': 'None',
        'note_sett_chg': 'None',
        'obs_air_wk_of': '5 CFU (ISO 8 142)',
        'etx_air_wk_of': 'ETX-260908-0584',
        'id_air_wk_of': 'Pending',
        'note_air': 'None',
        'obs_room_wk_of': '1 CFU (ISO 8 143)',
        'etx_room_wk_of': 'ETX-260908-0580',
        'id_room_wk_of': 'Pending',
        'note_room': 'None',
        'smart_phase1_summary': smart_phase1_full,
        'smart_phase1_continued': '',
        'oos1_analyst_name': 'N/A',
        'oos1_sample_id': 'N/A',
        'oos1_sample_name': 'N/A',
        'oos1_organism_morphology': 'N/A'
    }

    tpl.render(word_context)

    # Clean up Table 3 (past OOS results) and its header paragraph since no prior failures exist
    doc_word = tpl.docx
    for t in list(doc_word.tables):
        if len(t.rows) > 0 and 'In Trend of Past OOS' in t.rows[0].cells[0].text:
            p_elem = t._element
            p_elem.getparent().remove(p_elem)
        elif len(t.rows) > 0 and len(t.columns) == 5 and 'Analyst' in t.rows[0].cells[0].text:
            p_elem = t._element
            p_elem.getparent().remove(p_elem)

    for p in list(doc_word.paragraphs):
        if 'In Trend of Past OOS Results' in p.text:
            p._element.getparent().remove(p._element)

    # Replace Table 2 with the newest 9-column format from tables_docx
    tables_doc = docx.Document(tables_docx_path)
    new_t2_element = tables_doc.tables[1]._tbl

    # Find Table 2 in doc_word
    if len(doc_word.tables) >= 3:
        old_t2 = doc_word.tables[2]._tbl
        parent = old_t2.getparent()
        idx = parent.index(old_t2)
        parent.remove(old_t2)
        parent.insert(idx, new_t2_element)
    
    # Format Table 1
    if len(doc_word.tables) >= 2:
        t1 = doc_word.tables[1]
        for row in t1.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Hyperlink on sample ID cell
        add_hyperlink(t1.rows[1].cells[2].paragraphs[0], SAMPLE_URL, SAMPLE_ID, font_size=Pt(7))

    out_docx_desktop = os.path.join(DESKTOP_DIR, f"OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.docx")
    out_docx_docs = os.path.join(DOCUMENTS_DIR, f"OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.docx")

    doc_word.save(out_docx_desktop)
    doc_word.save(out_docx_docs)
    print(f"Generated Master Word Report:\n  {out_docx_desktop}\n  {out_docx_docs}")
    return out_docx_desktop


# ==========================================
# 4. BUILD COMPLETE 7-PAGE ACROFORM PDF
# ==========================================
def generate_7page_pdf(tables_pdf_path):
    pdf_template_path = os.path.join(ROOT_DIR, "ScanRDI OOS P1 template.pdf")
    doc_pdf = fitz.open(pdf_template_path)

    # Page 1
    page1 = doc_pdf[0]
    for w in page1.widgets():
        fn = w.field_name
        if fn == 'Text Field57': w.field_value = OOS_ID
        elif fn == 'Text Field0': w.field_value = 'Varsha Subramanian (Written by: Qiyue Chen)'
        elif fn == 'Date Field0': w.field_value = TEST_DATE_LONG
        elif fn == 'Date Field1': w.field_value = TEST_DATE_LONG
        elif fn == 'Date Field2': w.field_value = TEST_DATE_LONG
        elif fn == 'Text Field1': w.field_value = 'Scan RDI Sterility Test'
        elif fn == 'Text Field2': w.field_value = SAMPLE_ID
        elif fn == 'Text Field3':
            w.text_fontsize = 6.2
            w.field_value = (
                "Prepping Analyst: Min Jang (MJ)\r"
                "Processing Analyst: Varsha Subramanian (VV)\r"
                "Changeover Analyst: Varsha Subramanian (VV)\r"
                "Reading Analyst: Sonal Uprety (SU)"
            )
        elif fn == 'Text Field4': w.field_value = f"{SAMPLE_NAME}\r \r \r \r"
        elif fn == 'Text Field5': w.field_value = DOSAGE_FORM
        elif fn == 'Text Field6': w.field_value = LOT_NUMBER
        elif fn == 'Text Field7':
            w.text_fontsize = 7.5
            w.field_value = f"On {TEST_DATE}, sample {SAMPLE_ID} was found positive for viable microorganisms after ScanRDI testing\r \r"
        elif fn in ['Check Box0', 'Check Box1', 'Check Box2']: w.field_value = 'Yes'
        elif fn == 'Text Field8': w.field_value = 'MICRO-SOP-12 (16)\rENG-SOP-4 (05)'
        elif fn == 'Text Field9': w.field_value = '24Jul26\r24Jul26'
        elif fn == 'Text Field10': w.field_value = 'Rev: 16\rRev: 05'
        elif fn == 'Text Field11': w.field_value = '<1 event'
        elif fn == 'Text Field12': w.field_value = 'Kathan Parikh'
        elif fn == 'Date Field3': w.field_value = TEST_DATE_LONG
        elif fn == 'Text Field13':
            w.text_fontsize = 6.8
            w.field_value = 'Yes, analysts Min Jang, Varsha Subramanian, and Sonal Uprety were interviewed comprehensively.'
        elif fn == 'Text Field14': w.field_value = f"Yes, Sample ID: {SAMPLE_ID}"
        elif fn in ['Text Field15', 'Text Field16']: w.field_value = 'Yes, as per MICRO-SOP-12, ENG-SOP-4'
        elif fn == 'Text Field17': w.field_value = 'Yes, See 082526-1040-2 for more information.'
        elif fn == 'Text Field18': w.field_value = 'Yes, all analysts are trained and qualified by quality to perform the test'
        elif fn in ['Text Field19', 'Text Field20']: w.field_value = 'Not Applicable'
        elif fn == 'Text Field21': w.field_value = f"Yes, Information is available in Eagle Trax Sample Location History under {SAMPLE_ID}"
        elif fn in ['Check Box4', 'Check Box7', 'Check Box10', 'Check Box13', 'Check Box16', 'Check Box19',
                    'Check Box24', 'Check Box27', 'Check Box28', 'Check Box32', 'Check Box34', 'Check Box38']:
            w.field_value = 'Yes'
        w.update()

    # Page 2
    page2 = doc_pdf[1]
    for w in page2.widgets():
        fn = w.field_name
        if fn == 'Text Field57': w.field_value = OOS_ID
        elif fn == 'Text Field22':
            w.field_value = (
                "Scan Consumables: See the attached data packet\r"
                "Environmental Plates: TSA and Surface Plate: see attached environmental logs\r"
                "Monthly Cleaning: H2O2 strips and IPA: See attached monthly cleaning logs"
            )
        elif fn == 'Text Field23':
            w.field_value = (
                "Scan Consumables: See the attached data packet\r"
                "Environmental Plates: TSA and Surface Plate: see attached environmental logs\r"
                "Monthly Cleaning: H2O2 strips and IPA: See attached monthly cleaning logs"
            )
        elif fn == 'Text Field24': w.field_value = 'C. sporogenes'
        elif fn == 'Text Field25': w.field_value = '05282026-19404-CS\r \r \r'
        elif fn == 'Text Field26': w.field_value = '28 May 2028'
        elif fn in ['Text Field27', 'Text Field28', 'Text Field29']: w.field_value = 'Not Applicable'
        elif fn == 'Text Field30': w.field_value = 'E001040 (Cs2-105)'
        elif fn == 'Text Field31': w.field_value = 'Nov 2026'
        elif fn == 'Text Field32': w.field_value = 'CR145 (E001979) in L-Suite'
        elif fn == 'Text Field33': w.field_value = 'Dec 2026'
        elif fn == 'Text Field34': w.field_value = 'E001040 (Cs2-105)'
        elif fn == 'Text Field35': w.field_value = 'Nov 2026'
        elif fn in ['Text Field36', 'Text Field37', 'Text Field38', 'Text Field39']: w.field_value = 'Not Applicable'
        elif fn in ['Text Field40', 'Text Field41', 'Text Field42', 'Text Field45']: w.field_value = 'See Phase I Summary'
        elif fn == 'Text Field43':
            w.field_value = (
                "Incubator E001034 (Monitored by Sensor E001501)\r"
                "Incubator E001031 (Monitored by Sensor E001505)\r"
                "Incubators E001933, E001932 (30°C) & E001033 (33°C)"
            )
        elif fn == 'Text Field44':
            w.field_value = "Aug 2027     Feb 2027      Aug 2027     Feb 2027"
        elif fn in ['Check Box42', 'Check Box43', 'Check Box48', 'Check Box51', 'Check Box52',
                    'Check Box55', 'Check Box58', 'Check Box63', 'Check Box66', 'Check Box67',
                    'Check Box70', 'Check Box73']:
            w.field_value = 'Yes'
        elif fn in ['Check Box46', 'Check Box64']:
            w.field_value = 'Off'
        w.update()

    # Page 3
    page3 = doc_pdf[2]
    for w in page3.widgets():
        fn = w.field_name
        if fn == 'Text Field57': w.field_value = OOS_ID
        elif fn in ['Check Box78', 'Check Box79']: w.field_value = 'Yes'
        elif fn in ['Text Field46', 'Text Field47']: w.field_value = 'Not Applicable'
        elif fn == 'Text Field48': w.field_value = f"N/A QYC {datetime.now().strftime('%d%b%y')}"
        elif fn == 'Text Field49':
            w.text_fontsize = 6.8
            w.field_value = smart_phase1_part1
        w.update()

    # Page 4
    page4 = doc_pdf[3]
    for w in page4.widgets():
        fn = w.field_name
        if fn == 'Text Field57': w.field_value = OOS_ID
        elif fn == 'Text Field50':
            w.text_fontsize = 6.8
            w.field_value = smart_phase1_part2_page5
        w.update()

    # Page 5
    page5 = doc_pdf[4]
    for w in page5.widgets():
        fn = w.field_name
        if fn == 'Text Field57': w.field_value = OOS_ID
        elif fn == 'Text Field51':
            w.text_fontsize = 6.6
            w.field_value = smart_phase1_part2_page6
        w.update()

    # Page 6
    page6 = doc_pdf[5]
    for w in page6.widgets():
        fn = w.field_name
        if fn == 'Text Field57': w.field_value = OOS_ID
        elif fn == 'Check Box88': w.field_value = 'Yes'
        elif fn == 'Text Field53': w.field_value = 'Qiyue Chen'
        elif fn == 'Text Field54': w.field_value = 'Robin Seymour'
        w.update()

    # Append Page 7 (Standalone Tables PDF)
    tbl_doc = fitz.open(tables_pdf_path)
    doc_pdf.insert_pdf(tbl_doc, from_page=0, to_page=0)
    print(f"Appended Table 1 & Table 2 as Page 7. Total pages: {len(doc_pdf)}")

    out_pdf_desktop = os.path.join(DESKTOP_DIR, f"OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.pdf")
    out_pdf_docs = os.path.join(DOCUMENTS_DIR, f"OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.pdf")

    doc_pdf.save(out_pdf_desktop)
    doc_pdf.save(out_pdf_docs)
    print(f"Generated Complete 7-Page PDF Package:\n  {out_pdf_desktop}\n  {out_pdf_docs}")

    # Render preview images of pages 1 to 7 for verification
    preview_dir = os.path.join(SCRIPT_DIR, "preview_oos_261967")
    os.makedirs(preview_dir, exist_ok=True)
    for i, page in enumerate(doc_pdf):
        pix = page.get_pixmap(dpi=150)
        pix.save(os.path.join(preview_dir, f"page_{i+1}.png"))
    print(f"Rendered {len(doc_pdf)} page previews in {preview_dir}")

    doc_pdf.close()
    tbl_doc.close()
    return out_pdf_desktop


if __name__ == "__main__":
    print("=== [Step 1] Building Latest Standalone Tables (Word & PDF) ===")
    tables_doc, _ = build_standalone_tables_doc()
    
    tables_docx_desktop = os.path.join(DESKTOP_DIR, f"Tables OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.docx")
    tables_pdf_desktop = os.path.join(DESKTOP_DIR, f"Tables OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.pdf")
    tables_docx_docs = os.path.join(DOCUMENTS_DIR, f"Tables OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.docx")
    tables_pdf_docs = os.path.join(DOCUMENTS_DIR, f"Tables OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.pdf")

    tables_doc.save(tables_docx_desktop)
    tables_doc.save(tables_docx_docs)
    export_docx_to_pdf(tables_docx_desktop, tables_pdf_desktop)
    shutil.copy2(tables_pdf_desktop, tables_pdf_docs)
    print("Tables DOCX and PDF successfully saved to Desktop and Documents.")

    print("\n=== [Step 2] Building Master Investigation Word Report ===")
    generate_master_word_report(tables_docx_desktop)

    print("\n=== [Step 3] Building Complete 7-Page AcroForm PDF Package ===")
    generate_7page_pdf(tables_pdf_desktop)

    print("\n=== All OOS-261967 Deliverables Generated Successfully! ===")
