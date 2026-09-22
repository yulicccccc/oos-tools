import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client as win32

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
docx_path = os.path.join(desktop_dir, "OOS-261242 Review Email.docx")
html_path = os.path.join(desktop_dir, "OOS-261242_Review_Email_Robin_Format.html")

# 1. Update DOCX
doc = Document(docx_path)

# Ensure intro is exact
for p in doc.paragraphs:
    if "I have reviewed" in p.text:
        p.text = "I have reviewed OOS-261242. Please find the following edits made in the summary table below:"
        from docx.shared import Pt
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(11)

tbl = doc.tables[0]
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

# Row 6 is Phase I Summary: row index 6 (header is 0, rows 1-7)
# Let's find the row that has "Pagination" or "Root Cause"
for row in tbl.rows:
    sec_text = row.cells[0].text
    if "Root Cause" in sec_text or "Pagination" in sec_text:
        # Clean section title
        row.cells[0].text = "Phase I Summary – Root Cause Statement & Continuation Pages"
        row.cells[0].paragraphs[0].runs[0].font.name = "Calibri"
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(10)
        
        # Clean Unreviewed: remove (Text Field50) and (Text Field51)
        row.cells[1].text = (
            "The Phase I Root Cause Statement was placed at the bottom of Page 4, leaving Page 5 marked as "
            "\"N/A SMO 18-AUG-2026\", causing Page 4 to be overcrowded while leaving Page 5 blank."
        )
        row.cells[1].paragraphs[0].runs[0].font.name = "Calibri"
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
        
        # Clean Reviewed: remove (Text Field51)
        row.cells[2].text = (
            "Balanced narrative across continuation pages: Page 4 details bracketing summary, analyst interview, and monthly cleaning; "
            "Page 5 houses the evaluation of concurrent cleanroom BSC samples, taxonomic differentiation, and the finalized Phase I Root Cause Statement "
            "confirming that the recovery represents an isolated, transient environmental event with effective routine disinfection."
        )
        row.cells[2].paragraphs[0].runs[0].font.name = "Calibri"
        row.cells[2].paragraphs[0].runs[0].font.size = Pt(10)

doc.save(docx_path)
print("Updated DOCX: removed all AI 'Text Field' references.")

# 2. Update HTML
if os.path.exists(html_path):
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()
    html = html.replace("(`Text Field50`)", "").replace("(`Text Field51`)", "").replace("`Text Field50`", "Page 4").replace("`Text Field51`", "Page 5")
    html = html.replace("(Text Field50)", "").replace("(Text Field51)", "")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    print("Updated HTML: removed all AI 'Text Field' references.")

# 3. Copy via Word COM to clipboard
word_app = win32.Dispatch("Word.Application")
word_app.Visible = False
try:
    d = word_app.Documents.Open(os.path.abspath(docx_path))
    d.Content.Copy()
    d.Close(False)
    print("SUCCESS: Copied clean human-toned email to Windows Clipboard via Word COM!")
finally:
    word_app.Quit()
