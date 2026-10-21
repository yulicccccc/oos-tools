"""
Script: update_celsis_tables_262333_perfect.py
Purpose: Generate 100% PRD-compliant standalone Celsis tables document (.docx ONLY)
         for OOS-262333 (Optimal Balance Pharmacy, E19193), strictly adhering to:
         - Rule 0A: Ground truth roster (Alex Saravia / ES, Cuong Du / CCD)
         - Rule 0B / Rule 47 / Rule 50.7: Robin Seymour canonical Section Dividers:
             1. 'Personnel EM'
             2. 'Biological Safety Cabinet EM for E001316 and E001798'
             3. 'Weekly Active Air Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs'
             4. 'Weekly Surface Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs'
         - Rule 0E: Strict prohibition of '&' (zero ampersands)
         - Rule 0M: Deliver ONLY Word (.docx) to Desktop, NO standalone Table PDF
         - Rule 0N / Rule 60: Strict 7.0pt Times New Roman, zero font size tampering
         - Rule 34.5 / Rule 50.8: 'Week of Testing' in Day/Week column for weekly rows
         - Rule 35: Complete bracketing resolution ('No growth' / 'N/A' / 'N/A' / 'None')
         - Rule 41: Mandatory italics for 'Bacillus paralicheniformis', regular for '(Gram (+) rods)'
         - Interactive OpenXML hyperlinks for ETX-261002-0288 and ETX-261012-0654
"""

import os
import sys
import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
from docxtpl import DocxTemplate

sys.stdout.reconfigure(encoding='utf-8')

DOCS_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
DESKTOP_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
SCRATCH_DIR = os.path.join(DOCS_DIR, "scratch")
TEMPLATE_PATH = os.path.join(DOCS_DIR, "tables for celsis.docx")

OUT_DOCX_DESKTOP = os.path.join(DESKTOP_DIR, "Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis.docx")
OUT_DOCX_SCRATCH = os.path.join(SCRATCH_DIR, "Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis.docx")

# Clean old draft from desktop if present
for fname in [
    "Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis (Draft).docx",
    "Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis (Draft).pdf",
    "Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis.pdf"
]:
    p = os.path.join(DESKTOP_DIR, fname)
    if os.path.exists(p):
        try:
            os.remove(p)
            print(f"Removed {p}")
        except Exception as e:
            print(f"Could not remove {p}: {e}")

# Base Jinja render with empty strings so template variables don't show raw {{ }}
tpl = DocxTemplate(TEMPLATE_PATH)
expected_vars = tpl.get_undeclared_template_variables()
context = {v: "" for v in expected_vars}
tpl.render(context)
tpl.save(OUT_DOCX_SCRATCH)

doc = docx.Document(OUT_DOCX_SCRATCH)

# Helpers for cell XML manipulation
def set_cell_clean_text(cell, text, font_size_pt=7, bold=False):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    b_xml = "<w:b/><w:bCs/>" if bold else ""
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'{b_xml}'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'{b_xml}'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_hyperlink_to_cell(cell, url, text, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    part = cell.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:hyperlink r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="0000FF"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_table1_microbial_id_cell(cell, species, gram_stain, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{species}</w:t>'
        f'<w:br/>'
        f'</w:r>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{gram_stain}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_section_divider(row, banner_text, font_size_pt=7):
    """Sets section divider banner text across merged row with bold formatting."""
    for cell in row.cells:
        tc = cell._tc
        for p in tc.xpath('w:p'):
            tc.remove(p)
        p_xml = (
            f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:pPr>'
            f'<w:jc w:val="center"/>'
            f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:b/>'
            f'<w:bCs/>'
            f'<w:sz w:val="{int(font_size_pt*2)}"/>'
            f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
            f'</w:rPr>'
            f'</w:pPr>'
            f'<w:r>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:b/>'
            f'<w:bCs/>'
            f'<w:sz w:val="{int(font_size_pt*2)}"/>'
            f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
            f'</w:rPr>'
            f'<w:t>{banner_text}</w:t>'
            f'</w:r>'
            f'</w:p>'
        )
        tc.append(parse_xml(p_xml))
        cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

# ==================== TABLE 1 ====================
URL_SAMPLE = "https://etrax.eagleanalytical.com/Submission/Details/fZlFl0vWDpqR-VWBsez7IA__"
URL_MICRO = "https://etrax.eagleanalytical.com/Submission/Details/-2W3%24lpbkR6si5jTUhP7pg__"

t1 = doc.tables[0]
set_cell_clean_text(t1.rows[1].cells[0], "Alex Saravia (ES)")
set_cell_clean_text(t1.rows[1].cells[1], "Cuong Du (CCD)")
add_hyperlink_to_cell(t1.rows[1].cells[2], URL_SAMPLE, "ETX-261002-0288")
add_hyperlink_to_cell(t1.rows[1].cells[3], URL_MICRO, "ETX-261012-0654")
set_cell_clean_text(t1.rows[1].cells[4], "1 x 300mL FTM")
set_table1_microbial_id_cell(t1.rows[1].cells[5], "Bacillus paralicheniformis", "(Gram (+) rods)")

# ==================== TABLE 2 (Processing Phase) ====================
# Cleanroom: Suite 114 (CR114, CR114A, CR114B)
# Processing BSC: BSC E001316
# Canonical Banners per Rule 0B, Rule 47, Rule 50.7:
BANNER_PERS = "Personnel EM"
BANNER_BSC_PRO = "Biological Safety Cabinet EM for E001316 and E001798"
BANNER_AIR = "Weekly Active Air Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs"
BANNER_SURF = "Weekly Surface Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs"

t2 = doc.tables[1]

# Dividers
set_section_divider(t2.rows[1], BANNER_PERS)
set_section_divider(t2.rows[5], BANNER_BSC_PRO)
set_section_divider(t2.rows[12], BANNER_AIR)
set_section_divider(t2.rows[14], BANNER_SURF)

# Daily Rows: 2, 3, 4, 6, 7, 8, 9, 10, 11
# Row index -> (Date, Analyst, Day/Week)
pro_daily_data = {
    2:  ("02Oct26", "ES", "Date Before Testing"),
    3:  ("05Oct26", "ES", "Date of Testing"),
    4:  ("06Oct26", "ES", "Date After Testing"),
    6:  ("02Oct26", "ES", "Date Before Testing"),
    7:  ("05Oct26", "ES", "Date of Testing"),
    8:  ("06Oct26", "ES", "Date After Testing"),
    9:  ("02Oct26", "ES", "Date Before Testing"),
    10: ("05Oct26", "ES", "Date of Testing"),
    11: ("06Oct26", "ES", "Date After Testing"),
}

for r_idx, (dt, ana, day_desc) in pro_daily_data.items():
    row = t2.rows[r_idx]
    set_cell_clean_text(row.cells[3], dt)
    set_cell_clean_text(row.cells[4], ana)
    set_cell_clean_text(row.cells[6], day_desc)
    set_cell_clean_text(row.cells[7], "No growth")
    set_cell_clean_text(row.cells[9], "N/A")
    set_cell_clean_text(row.cells[10], "N/A")
    set_cell_clean_text(row.cells[11], "None")

# Weekly Rows: 13 (Active Air) and 15 (Surface)
# Week of 05Oct26
set_cell_clean_text(t2.rows[13].cells[3], "05Oct26")
set_cell_clean_text(t2.rows[13].cells[4], "SMO")
set_cell_clean_text(t2.rows[13].cells[5], "Week of Testing")
set_cell_clean_text(t2.rows[13].cells[8], "No growth")
set_cell_clean_text(t2.rows[13].cells[9], "N/A")
set_cell_clean_text(t2.rows[13].cells[10], "N/A")
set_cell_clean_text(t2.rows[13].cells[12], "None")

set_cell_clean_text(t2.rows[15].cells[3], "05Oct26")
set_cell_clean_text(t2.rows[15].cells[4], "SMO")
set_cell_clean_text(t2.rows[15].cells[5], "Week of Testing")
set_cell_clean_text(t2.rows[15].cells[8], "No growth")
set_cell_clean_text(t2.rows[15].cells[9], "N/A")
set_cell_clean_text(t2.rows[15].cells[10], "N/A")
set_cell_clean_text(t2.rows[15].cells[12], "None")

# ==================== TABLE 3 (Aliquoting Phase) ====================
# Aliquoting BSC: BSC E001798
t3 = doc.tables[2]

# Dividers
set_section_divider(t3.rows[1], BANNER_PERS)
set_section_divider(t3.rows[5], BANNER_BSC_PRO)
set_section_divider(t3.rows[12], BANNER_AIR)
set_section_divider(t3.rows[14], BANNER_SURF)

alq_daily_data = {
    2:  ("09Oct26", "CCD", "Date Before Testing"),
    3:  ("12Oct26", "CCD", "Date of Testing"),
    4:  ("13Oct26", "CCD", "Date After Testing"),
    6:  ("09Oct26", "CCD", "Date Before Testing"),
    7:  ("12Oct26", "CCD", "Date of Testing"),
    8:  ("13Oct26", "CCD", "Date After Testing"),
    9:  ("09Oct26", "CCD", "Date Before Testing"),
    10: ("12Oct26", "CCD", "Date of Testing"),
    11: ("13Oct26", "CCD", "Date After Testing"),
}

for r_idx, (dt, ana, day_desc) in alq_daily_data.items():
    row = t3.rows[r_idx]
    set_cell_clean_text(row.cells[3], dt)
    set_cell_clean_text(row.cells[4], ana)
    set_cell_clean_text(row.cells[6], day_desc)
    set_cell_clean_text(row.cells[7], "No growth")
    set_cell_clean_text(row.cells[9], "N/A")
    set_cell_clean_text(row.cells[10], "N/A")
    set_cell_clean_text(row.cells[11], "None")

# Weekly Rows: 13 (Active Air) and 15 (Surface)
# Week of 12Oct26
set_cell_clean_text(t3.rows[13].cells[3], "12Oct26")
set_cell_clean_text(t3.rows[13].cells[4], "SMO")
set_cell_clean_text(t3.rows[13].cells[5], "Week of Testing")
set_cell_clean_text(t3.rows[13].cells[8], "No growth")
set_cell_clean_text(t3.rows[13].cells[9], "N/A")
set_cell_clean_text(t3.rows[13].cells[10], "N/A")
set_cell_clean_text(t3.rows[13].cells[12], "None")

set_cell_clean_text(t3.rows[15].cells[3], "12Oct26")
set_cell_clean_text(t3.rows[15].cells[4], "SMO")
set_cell_clean_text(t3.rows[15].cells[5], "Week of Testing")
set_cell_clean_text(t3.rows[15].cells[8], "No growth")
set_cell_clean_text(t3.rows[15].cells[9], "N/A")
set_cell_clean_text(t3.rows[15].cells[10], "N/A")
set_cell_clean_text(t3.rows[15].cells[12], "None")

# Page break before Table 3
for p in doc.paragraphs:
    if "Table 3:" in p.text:
        p.paragraph_format.page_break_before = True

# Verification: Zero ampersands
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            if '&' in c.text:
                c.text = c.text.replace('&', 'and')

amp_count = 0
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            if '&' in c.text:
                print(f"ERROR: Table {t_idx+1} Row {r_idx} Cell {c_idx} has '&': {c.text}")
                amp_count += 1
assert amp_count == 0, f"Found {amp_count} ampersands in document!"
print(">>> ZERO AMPERSANDS CONFIRMED! <<<")

# Save DOCX
doc.save(OUT_DOCX_SCRATCH)
print(f"Saved scratch docx: {OUT_DOCX_SCRATCH}")
doc.save(OUT_DOCX_DESKTOP)
print(f"Saved desktop docx: {OUT_DOCX_DESKTOP}")

print("PERFECT CELSIS TABLES GENERATED SUCCESSFULLY!")
