import docx

doc = docx.Document('scratch/Celsis table OOS-262017 (User Edited 25th Backup).docx')
t2 = doc.tables[1]
for r_idx in [8, 11]:
    row = t2.rows[r_idx]
    for c_idx, c in enumerate(row.cells):
        if 'GS' in c.text:
            fonts = [r.font.name for p in c.paragraphs for r in p.runs]
            sizes = [r.font.size for p in c.paragraphs for r in p.runs]
            print(f"Row {r_idx} Cell {c_idx}: text='{c.text}', fonts={fonts}, sizes={sizes}")
