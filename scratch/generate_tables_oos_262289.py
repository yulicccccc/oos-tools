import os
import sys
import shutil
import fitz
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"

OOS_ID = "262289"
CLIENT_NAME = "Link Pharmacy (E19112)"
SAMPLE_ID = "ETX-260929-0931"
SAMPLE_URL = "https://etrax.eagleanalytical.com/Submission/Details/sM81r08$r4X9k5LrerlJ7g__"
SAMPLE_NAME = "Tirzepatide/Levocarnitine 17mg/100mg/mL (1ml)"
LOT_NUMBER = "20260928@54"
TEST_DATE = "01Oct26"

PROCESSING_ANALYST = "Ellen Bennett"
PROCESSING_INITIAL = "ELB"
READING_ANALYST = "Varsha Subramanian"
EVENT_COUNT = "15"
CONFIRMED_COUNT = "4"
MORPHOLOGY = "Cocci shaped morphology"

PROCESSING_BSC = "1938"
CHANGEOVER_ANALYST = "Karla Silva"
CHANGEOVER_INITIAL = "KSM"
CHANGEOVER_BSC = "1937"

AIR_ANALYST = "SMO"
AIR_DATE = "01Oct26"

# Active air positive recoveries
ETX_144 = "ETX-261012-0430"
URL_144 = "https://etrax.eagleanalytical.com/Submission/Details/xxCNoUxIy4doEwKyIX-rYg__"
ETX_142 = "ETX-261012-0433"
URL_142 = "https://etrax.eagleanalytical.com/Submission/Details/Qr5FQ8nlCFY1Y9QOAyCjZQ__"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=35, bottom=35, left=45, right=45):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_hyperlink(paragraph, url, text, font_size=Pt(7), color_rgb="0000FF"):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color_rgb}"/>'))
    rPr.append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
    new_run.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def format_cell(cell, text, bold=False, italic=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None, v_align=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = v_align
    set_cell_margins(cell)
    if fill_hex:
        set_cell_background(cell, fill_hex)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = font_size
    r.bold = bold
    r.italic = italic
    return r

def build_standalone_tables_doc():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # ---------------- Table 1 ----------------
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"Table 1: Information for {SAMPLE_ID} under investigation")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(8.5)
    r_t1.underline = True

    t1 = doc.add_table(rows=2, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = [
        "Processing Analyst",
        "Reading Analyst",
        "Sample ID",
        "Events",
        "Confirmed Microbial\nEvents",
        "Morphology Description"
    ]
    t1_widths = [Inches(1.25), Inches(1.25), Inches(1.1), Inches(0.7), Inches(1.2), Inches(2.0)]

    for j, h in enumerate(t1_headers):
        format_cell(t1.cell(0, j), h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    format_cell(t1.cell(1, 0), PROCESSING_ANALYST, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 1), READING_ANALYST, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    c_sid = t1.cell(1, 2)
    c_sid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sid)
    p_sid = c_sid.paragraphs[0]
    p_sid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sid.paragraph_format.space_before = Pt(0)
    p_sid.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sid, SAMPLE_URL, SAMPLE_ID, font_size=Pt(7))

    format_cell(t1.cell(1, 3), EVENT_COUNT, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 4), CONFIRMED_COUNT, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 5), MORPHOLOGY, bold=True, italic=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in t1.rows:
        for j, w in enumerate(t1_widths):
            row.cells[j].width = w

    # Spacing between tables
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(0)

    # ---------------- Table 2 ----------------
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run(f"Table 2: Environmental Monitoring from Testing Performed on {TEST_DATE}")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(8.5)
    r_t2.underline = True

    t2 = doc.add_table(rows=0, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    t2_widths = [
        Inches(1.05), # Sampling site
        Inches(0.65), # Frequency
        Inches(0.70), # Date (DDMMMYY)
        Inches(0.55), # Analyst
        Inches(1.15), # Day / Week(s)
        Inches(1.15), # Observation
        Inches(0.85), # EM Plate ETX Number
        Inches(0.80), # Microbial ID
        Inches(0.60)  # Notes
    ]

    hdr_row = t2.add_row()
    t2_headers = [
        "Environmental\nMonitoring (EM)\nSampling Site",
        "Frequency",
        "Date\n(DDMMMYY)",
        "Analyst",
        "Day /Week(s)",
        "Observation",
        "EM Plate ETX\nNumber",
        "Microbial ID",
        "Notes"
    ]
    for j, h in enumerate(t2_headers):
        format_cell(hdr_row.cells[j], h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    def add_section_divider(table, title_text):
        row = table.add_row()
        cell = row.cells[0]
        for c in row.cells[1:]:
            cell.merge(c)
        format_cell(cell, title_text, bold=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex="E8E8E8")
        return row

    def add_data_row(table, site, freq, date_str, analyst, day_week, obs, etx_num, micro_id, notes, etx_url=None):
        row = table.add_row()
        format_cell(row.cells[0], site, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], freq, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[2], date_str, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[3], analyst, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[4], day_week, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[5], obs, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        
        if etx_url:
            c = row.cells[6]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(c)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_hyperlink(p, etx_url, etx_num, font_size=Pt(7))
        else:
            format_cell(row.cells[6], etx_num, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
            
        format_cell(row.cells[7], micro_id, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[8], notes, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        return row

    # --- Section 1: Personnel EM ---
    add_section_divider(t2, "Personnel EM")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", TEST_DATE, PROCESSING_INITIAL, "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", TEST_DATE, CHANGEOVER_INITIAL, "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t2, f"Biological Safety Cabinet EM for E00{PROCESSING_BSC} and E00{CHANGEOVER_BSC}")
    add_data_row(t2, f"Surface Sampling of\nISO 5 BSC E00{PROCESSING_BSC}\n(4 Locations)", "Daily", TEST_DATE, PROCESSING_INITIAL, "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, f"Surface Sampling of\nISO 5 BSC E00{CHANGEOVER_BSC}\n(4 Locations)", "Daily", TEST_DATE, CHANGEOVER_INITIAL, "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, f"Settling Sampling of\nISO 5 BSC E00{PROCESSING_BSC}", "Daily", TEST_DATE, PROCESSING_INITIAL, "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, f"Settling Sampling of\nISO 5 BSC E00{CHANGEOVER_BSC}", "Daily", TEST_DATE, CHANGEOVER_INITIAL, "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 3: Weekly Active Air L-Suite (RULE 0B: streamlined to (L-Suite)) ---
    add_section_divider(t2, "Weekly Active Air Sampling of Cleanroom 145, 144, 143, 142 (L-Suite) with Processing BSCs")

    # Create 2 rows for the 2 active air plates, merging the left columns
    r_air1 = t2.add_row()
    r_air2 = t2.add_row()

    for col_idx in [0, 1, 2, 3, 4, 8]:
        r_air1.cells[col_idx].merge(r_air2.cells[col_idx])

    format_cell(r_air1.cells[0], "Active Air Sampling\nof Cleanrooms", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT)
    format_cell(r_air1.cells[1], "Weekly", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(r_air1.cells[2], AIR_DATE, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(r_air1.cells[3], AIR_ANALYST, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(r_air1.cells[4], "Week of Testing", bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(r_air1.cells[8], "None", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 1 (144 Sec II): 1 CFU
    format_cell(r_air1.cells[5], "1 CFU\n(ISO 7 144 Sec II)", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
    c_etx1 = r_air1.cells[6]
    c_etx1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_etx1)
    p_e1 = c_etx1.paragraphs[0]
    p_e1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_e1.paragraph_format.space_before = Pt(0)
    p_e1.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_e1, URL_144, ETX_144, font_size=Pt(6.5))

    c_id1 = r_air1.cells[7]
    c_id1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_id1)
    p_i1 = c_id1.paragraphs[0]
    p_i1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_i1.paragraph_format.space_before = Pt(0)
    p_i1.paragraph_format.space_after = Pt(0)
    p_i1.paragraph_format.line_spacing = 1.0
    r_i1_a = p_i1.add_run("S. epidermidis\n")
    r_i1_a.font.name = 'Times New Roman'
    r_i1_a.font.size = Pt(6.2)
    r_i1_a.italic = True
    r_i1_b = p_i1.add_run("(Gram (+) cocci)")
    r_i1_b.font.name = 'Times New Roman'
    r_i1_b.font.size = Pt(5.8)

    # Row 2 (142): 5 CFUs
    format_cell(r_air2.cells[5], "5 CFUs\n(ISO 8 142)", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
    c_etx2 = r_air2.cells[6]
    c_etx2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_etx2)
    p_e2 = c_etx2.paragraphs[0]
    p_e2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_e2.paragraph_format.space_before = Pt(0)
    p_e2.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_e2, URL_142, ETX_142, font_size=Pt(6.5))

    format_cell(r_air2.cells[7], "Gram (+) rods,\nGram (-) rods,\nGram (+) cocci", bold=False, font_size=Pt(6.0), align=WD_ALIGN_PARAGRAPH.CENTER)

    # --- Section 4: Weekly Surface Sampling L-Suite (RULE 0B: streamlined to (L-Suite)) ---
    add_section_divider(t2, "Weekly Surface Sampling of Cleanroom 145, 144, 143, 142 (L-Suite) with Processing BSCs")
    add_data_row(t2, "Surface Sampling of\nCleanrooms", "Weekly", AIR_DATE, AIR_ANALYST, "Week of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # Set column widths
    for row in t2.rows:
        if len(row.cells) == 9:
            for j, w in enumerate(t2_widths):
                row.cells[j].width = w

    return doc

if __name__ == "__main__":
    print(f"=== Generating Standalone Tables for OOS-{OOS_ID} ({CLIENT_NAME}) ===")
    doc = build_standalone_tables_doc()

    docx_filename = f"Tables OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.docx"
    pdf_filename = f"Tables OOS-{OOS_ID} {CLIENT_NAME} - ScanRDI.pdf"

    scratch_docx = os.path.join(SCRIPT_DIR, docx_filename)
    scratch_pdf = os.path.join(SCRIPT_DIR, pdf_filename)
    doc.save(scratch_docx)
    print(f"Saved DOCX to: {scratch_docx}")

    print("Converting to PDF via Microsoft Word...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_w = word.Documents.Open(os.path.abspath(scratch_docx))
    doc_w.SaveAs(os.path.abspath(scratch_pdf), FileFormat=17)
    doc_w.Close()
    word.Quit()
    print(f"Exported PDF to: {scratch_pdf}")

    # Verify Page Count
    doc_fitz = fitz.open(scratch_pdf)
    print(f"Verified PDF Page Count: {len(doc_fitz)} (Strict requirement: 1)")
    assert len(doc_fitz) == 1, f"ERROR: Standalone tables must be 1 page, got {len(doc_fitz)}"

    # Render high-res verification preview
    pix = doc_fitz[0].get_pixmap(dpi=150)
    preview_img = os.path.join(SCRIPT_DIR, f"preview_tables_{OOS_ID}_p1.png")
    pix.save(preview_img)
    print(f"Saved preview image to: {preview_img}")
    doc_fitz.close()

    # Sync to Desktop and Documents\OOS
    for dest_dir in [DESKTOP_DIR, ROOT_DIR]:
        dst_docx = os.path.join(dest_dir, docx_filename)
        dst_pdf = os.path.join(dest_dir, pdf_filename)
        try:
            shutil.copy2(scratch_docx, dst_docx)
            shutil.copy2(scratch_pdf, dst_pdf)
            print(f"Synced deliverables to: {dest_dir}")
        except Exception as e:
            print(f"Sync error for {dest_dir}: {e}")

    print("\n=== Table generation and synchronization complete! ===")
