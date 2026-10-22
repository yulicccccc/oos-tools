import docx

path = r'C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis.docx'
doc = docx.Document(path)

# Let's inspect all cells of Table 1, Table 2, Table 3 and list every cell's text
print("--- TABLE 1 ---")
for r_i, r in enumerate(doc.tables[0].rows):
    row_vals = [c.text.strip() for c in r.cells]
    print(f"R{r_i}: {row_vals}")

print("\n--- TABLE 2 ---")
for r_i, r in enumerate(doc.tables[1].rows):
    row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
    # dedup merged
    dedup = []
    for x in row_vals:
        if not dedup or x != dedup[-1]:
            dedup.append(x)
    print(f"R{r_i:02d}: {dedup}")

print("\n--- TABLE 3 ---")
for r_i, r in enumerate(doc.tables[2].rows):
    row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
    dedup = []
    for x in row_vals:
        if not dedup or x != dedup[-1]:
            dedup.append(x)
    print(f"R{r_i:02d}: {dedup}")
