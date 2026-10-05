import openpyxl
import pandas as pd

excel_path = r'scratch\Sterile Lab - OOS Tracking Log.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("Sheet names in workbook:")
for name in wb.sheetnames:
    print(f" - {name}")

print("\n" + "="*60)
print("Searching for 'Optimal Balance' or 'E19193' across ALL sheets:")
print("="*60)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    found = []
    # Read rows
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    
    # Try finding header row
    header_idx = None
    for idx, row in enumerate(rows[:10]):
        row_str = " ".join([str(c) for c in row if c is not None])
        if "ETX" in row_str or "OOS" in row_str or "Client" in row_str:
            header_idx = idx
            break
            
    header = rows[header_idx] if header_idx is not None else [f"Col{i}" for i in range(len(rows[0]))]
    
    for r_idx, row in enumerate(rows):
        row_str = " ".join([str(c) for c in row if c is not None])
        if 'optimal balance' in row_str.lower() or 'e19193' in row_str.lower():
            row_dict = {}
            for c_idx, val in enumerate(row):
                if val is not None and str(val).strip() != "":
                    col_name = str(header[c_idx]).strip() if c_idx < len(header) and header[c_idx] is not None else f"Col{c_idx}"
                    row_dict[col_name] = str(val)
            found.append((r_idx + 1, row_dict))
            
    if found:
        print(f"\n>>> Found in Sheet '{sheet_name}' ({len(found)} rows):")
        for r_num, row_data in found:
            print(f"  Row {r_num}:")
            for k, v in row_data.items():
                print(f"    {k}: {v}")
    else:
        # print(f"Sheet '{sheet_name}': No records found")
        pass

print("\n" + "="*60)
print("Inspecting Sheet 'Celsis Sterility OOS' Details:")
print("="*60)
ws_celsis = wb['Celsis Sterility OOS']
c_rows = list(ws_celsis.iter_rows(values_only=True))
for idx in range(min(15, len(c_rows))):
    print(f"Row {idx+1}: {[str(x)[:20] if x is not None else '' for x in c_rows[idx][:10]]}")

print("\nInspecting Sheet 'OOS info for Priority Clients' Details:")
print("="*60)
ws_prio = wb['OOS info for Priority Clients']
p_rows = list(ws_prio.iter_rows(values_only=True))
for idx in range(min(15, len(p_rows))):
    print(f"Row {idx+1}: {[str(x)[:20] if x is not None else '' for x in p_rows[idx][:10]]}")
