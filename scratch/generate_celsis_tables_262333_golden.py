import os
import sys
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
LOCAL_DESKTOP_DIR = r"C:\Users\qchen\Desktop"
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"

OOS_ID = "262333"
CLIENT_NAME = "Optimal Balance Pharmacy (E19193)"
TEST_NAME = "Celsis"
SAMPLE_ID = "ETX-261002-0288"
SAMPLE_URL = "https://etrax.eagleanalytical.com/Submission/Details/fZlFl0vWDpqR-VWBsez7IA__"
SUBCULTURE_ID = "ETX-261012-0654"
SUBCULTURE_URL = "https://etrax.eagleanalytical.com/Submission/Details/-2W3%24lpbkR6si5jTUhP7pg__"
ANALYTE = "TESAMORELIN 8 MG/ML (3 ML) Injection"
LOT_NUMBER = "LG342011076"
MEDIA = "1 x 300mL FTM"
MICROBIAL_ID = "Bacillus paralicheniformis"

PRO_DATE = "05Oct26"
PRO_ANALYST_FULL = "Alex Saravia"
PRO_INITIAL = "ES"

ALQ_DATE = "12Oct26"
ALQ_ANALYST_FULL = "Cuong Du"
ALQ_INITIAL = "CCD"

AIR_ANALYST = "SMO"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=35, bottom=35, left=45, right=45):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_hyperlink(paragraph, url, text, font_size=Pt(7), color_rgb="0000FF"):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color_rgb}"/>'))
    rPr.append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
    new_run.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def format_cell(cell, text, bold=False, italic=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None, v_align=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = v_align
    set_cell_margins(cell)
    if fill_hex:
        set_cell_background(cell, fill_hex)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = font_size
    r.bold = bold
    r.italic = italic
    return r

def build_celsis_tables():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # =========================================================================
    # Table 1: Information for Sample under investigation
    # =========================================================================
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"Table 1: Information for {SAMPLE_ID} under investigation")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(8.5)
    r_t1.underline = True

    t1 = doc.add_table(rows=2, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = [
        "Processing Analyst",
        "Aliquoting Analyst",
        "Sample ID",
        "Related Microbial ID",
        "Media with microbial growth",
        "Microbial ID"
    ]
    t1_widths = [Inches(1.20), Inches(1.20), Inches(1.10), Inches(1.10), Inches(1.20), Inches(1.70)]

    for j, h in enumerate(t1_headers):
        format_cell(t1.cell(0, j), h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    format_cell(t1.cell(1, 0), PRO_ANALYST_FULL, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 1), ALQ_ANALYST_FULL, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    # Sample ID Hyperlink
    c_sid = t1.cell(1, 2)
    c_sid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sid)
    p_sid = c_sid.paragraphs[0]
    p_sid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sid.paragraph_format.space_before = Pt(0)
    p_sid.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sid, SAMPLE_URL, SAMPLE_ID, font_size=Pt(7))

    # Subculture ID Hyperlink
    c_sub = t1.cell(1, 3)
    c_sub.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sub)
    p_sub = c_sub.paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sub, SUBCULTURE_URL, SUBCULTURE_ID, font_size=Pt(7))

    format_cell(t1.cell(1, 4), MEDIA, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
    # Rule 0P.II.1 & Rule 61.II.1: "有Microbial ID就不用staining了", binomial italicized
    format_cell(t1.cell(1, 5), MICROBIAL_ID, bold=False, italic=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in t1.rows:
        for j, w in enumerate(t1_widths):
            row.cells[j].width = w

    # Spacing between tables
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(0)

    # =========================================================================
    # Helper to build 9-column EM Table (Table 2 & Table 3)
    # =========================================================================
    t2_widths = [
        Inches(1.20), # Sampling site
        Inches(0.55), # Frequency
        Inches(0.60), # Date (DDMMMYY)
        Inches(0.50), # Analyst
        Inches(0.95), # Day / Week(s)
        Inches(0.95), # Observation
        Inches(0.95), # EM Plate ETX Number
        Inches(1.20), # Microbial ID
        Inches(0.60)  # Notes
    ]

    t2_headers = [
        "Environmental\nMonitoring (EM)\nSampling Site",
        "Frequency",
        "Date\n(DDMMMYY)",
        "Analyst",
        "Day /Week(s)",
        "Observation",
        "EM Plate ETX Number",
        "Microbial ID",
        "Notes"
    ]

    def add_section_divider(table, title_text):
        row = table.add_row()
        cell = row.cells[0]
        for c in row.cells[1:]:
            cell.merge(c)
        format_cell(cell, title_text, bold=True, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex="E8E8E8")
        return row

    def add_data_row(table, site, freq, date_str, analyst, day_week, obs, etx_num, micro_id, notes, micro_italic=False):
        row = table.add_row()
        format_cell(row.cells[0], site, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], freq, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[2], date_str, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[3], analyst, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[4], day_week, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[5], obs, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[6], etx_num, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[7], micro_id, bold=False, italic=micro_italic, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[8], notes, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)
        return row

    # =========================================================================
    # Table 2: Environmental Monitoring from Processing Performed on 05Oct26
    # =========================================================================
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run(f"Table 2: Environmental Monitoring from Processing Performed on {PRO_DATE}")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(8.5)
    r_t2.underline = True

    t2 = doc.add_table(rows=0, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    hdr2 = t2.add_row()
    for j, h in enumerate(t2_headers):
        format_cell(hdr2.cells[j], h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    # --- Section 1: Personnel EM ---
    add_section_divider(t2, "Personnel EM")
    add_data_row(t2, "Personnel (Left Touch and Right Touch)", "Daily", "02Oct26", PRO_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personnel (Left Touch and Right Touch)", "Daily", "05Oct26", PRO_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personnel (Left Touch and Right Touch)", "Daily", "06Oct26", PRO_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t2, "Biological Safety Cabinet EM for E001316 and E001798")
    add_data_row(t2, "Surface Sampling of ISO 5 (4 locations)", "Daily", "02Oct26", PRO_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Surface Sampling of ISO 5 (4 locations)", "Daily", "05Oct26", PRO_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Surface Sampling of ISO 5 (4 locations)", "Daily", "06Oct26", PRO_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of ISO 5 (2 locations)", "Daily", "02Oct26", PRO_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of ISO 5 (2 locations)", "Daily", "05Oct26", PRO_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Settling Sampling of ISO 5 (2 locations)", "Daily", "06Oct26", PRO_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 3: Cleanroom Active Air ---
    add_section_divider(t2, "Weekly Active Air Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs")
    add_data_row(t2, "Active Air Sampling of Cleanrooms", "Weekly", "05Oct26", AIR_ANALYST, "Week of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 4: Cleanroom Surface ---
    add_section_divider(t2, "Weekly Surface Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs")
    add_data_row(t2, "Surface Sampling of Cleanrooms", "Weekly", "05Oct26", AIR_ANALYST, "Week of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    for row in t2.rows:
        for j, w in enumerate(t2_widths):
            row.cells[j].width = w

    # =========================================================================
    # Page Break Before Table 3 (PRD Rule 12.d & Golden Standard)
    # =========================================================================
    doc.add_page_break()

    # =========================================================================
    # Table 3: Environmental Monitoring from Aliquoting Performed on 12Oct26
    # =========================================================================
    p_t3 = doc.add_paragraph()
    p_t3.paragraph_format.space_before = Pt(0)
    p_t3.paragraph_format.space_after = Pt(2)
    r_t3 = p_t3.add_run(f"Table 3: Environmental Monitoring from Aliquoting Performed on {ALQ_DATE}")
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(8.5)
    r_t3.underline = True

    t3 = doc.add_table(rows=0, cols=9)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)

    hdr3 = t3.add_row()
    for j, h in enumerate(t2_headers):
        format_cell(hdr3.cells[j], h, bold=False, font_size=Pt(7), align=WD_ALIGN_PARAGRAPH.CENTER)

    # --- Section 1: Personnel EM ---
    add_section_divider(t3, "Personnel EM")
    add_data_row(t3, "Personnel (Left Touch and Right Touch)", "Daily", "09Oct26", ALQ_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Personnel (Left Touch and Right Touch)", "Daily", "12Oct26", ALQ_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Personnel (Left Touch and Right Touch)", "Daily", "13Oct26", ALQ_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t3, "Biological Safety Cabinet EM for E001316 and E001798")
    add_data_row(t3, "Surface Sampling of ISO 5 (4 locations)", "Daily", "09Oct26", ALQ_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Surface Sampling of ISO 5 (4 locations)", "Daily", "12Oct26", ALQ_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Surface Sampling of ISO 5 (4 locations)", "Daily", "13Oct26", ALQ_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Settling Sampling of ISO 5 (2 locations)", "Daily", "09Oct26", ALQ_INITIAL, "Date Before Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Settling Sampling of ISO 5 (2 locations)", "Daily", "12Oct26", ALQ_INITIAL, "Date of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t3, "Settling Sampling of ISO 5 (2 locations)", "Daily", "13Oct26", ALQ_INITIAL, "Date After Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 3: Cleanroom Active Air ---
    add_section_divider(t3, "Weekly Active Air Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs")
    add_data_row(t3, "Active Air Sampling of Cleanrooms", "Weekly", "12Oct26", AIR_ANALYST, "Week of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 4: Cleanroom Surface ---
    add_section_divider(t3, "Weekly Surface Sampling of Cleanroom 114, 114A, 114B (Suite 114) with Processing BSCs")
    add_data_row(t3, "Surface Sampling of Cleanrooms", "Weekly", "12Oct26", AIR_ANALYST, "Week of Testing", "No Growth", "Not Applicable", "Not Applicable", "None")

    for row in t3.rows:
        for j, w in enumerate(t2_widths):
            row.cells[j].width = w

    # =========================================================================
    # Rule 0N / Rule 0P / Rule 61: Universal Strict 7.0pt (14 half-points)
    # =========================================================================
    for elem in doc._element.iter():
        if elem.tag.endswith('sz') or elem.tag.endswith('szCs'):
            elem.set(qn('w:val'), '14')

    # Save to scratch
    scratch_dest = f"scratch/Tables_OOS_{OOS_ID}_Golden.docx"
    doc.save(scratch_dest)
    print(f"Generated clean Golden Standard tables: {scratch_dest}")

    # Deploy to Desktop locations
    filename = f"Tables OOS-{OOS_ID} {CLIENT_NAME} - {TEST_NAME}.docx"
    desktop_dest = os.path.join(DESKTOP_DIR, filename)
    local_dest = os.path.join(LOCAL_DESKTOP_DIR, filename)
    doc_dest = os.path.join(DOCUMENTS_DIR, filename)

    doc.save(desktop_dest)
    print(f"Deployed to OneDrive Desktop: {desktop_dest}")
    try:
        doc.save(local_dest)
        print(f"Deployed to Local Desktop: {local_dest}")
    except Exception as e:
        print(f"Local Desktop save error: {e}")
    try:
        doc.save(doc_dest)
        print(f"Deployed to Documents: {doc_dest}")
    except Exception as e:
        print(f"Documents save error: {e}")

    # Also keep compatibility with Draft name if user clicks recent
    draft_filename = f"Tables OOS-{OOS_ID} {CLIENT_NAME} - {TEST_NAME} (Draft).docx"
    doc.save(os.path.join(DESKTOP_DIR, draft_filename))
    try:
        doc.save(os.path.join(LOCAL_DESKTOP_DIR, draft_filename))
    except Exception as e:
        pass

    return desktop_dest

if __name__ == "__main__":
    build_celsis_tables()
