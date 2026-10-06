import docx

def update_week_text(filepath):
    doc = docx.Document(filepath)
    updated = 0
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    if 'Week on Testing Date' in p.text:
                        for run in p.runs:
                            if 'Week on Testing Date' in run.text:
                                run.text = run.text.replace('Week on Testing Date', 'Week of Testing')
                                updated += 1
                        if 'Week on Testing Date' in p.text:
                            p.text = p.text.replace('Week on Testing Date', 'Week of Testing')
                            updated += 1
    doc.save(filepath)
    print(f"Updated {filepath}, replaced {updated} instances.")

update_week_text("tables for celsis.docx")
update_week_text("tables for 71.docx")
