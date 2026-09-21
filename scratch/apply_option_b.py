import fitz, os, shutil, docx, win32clipboard
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
pdf_path = os.path.join(desktop_dir, "OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf")
docx_path = os.path.join(desktop_dir, "OOS-261186 Review Email.docx")
html_path = os.path.join(desktop_dir, "OOS-261186_Review_Email_Robin_Format.html")

# ==========================================
# 1. UPDATE PDF (Restore Legacy Root Cause)
# ==========================================
doc = fitz.open(pdf_path)

legacy_rc = (
    "Based on the findings outlined in the preceding sections, the Out-Of-Specification (OOS) "
    "result observed for the Environmental Monitoring (EM) Active Air Sampling plate may be attributed "
    "to a potential analyst error. It is to be noted that the growth observed on the weekly active air "
    "and weekly surface sampling plates for the day of testing do not follow a trend, indicating that "
    "the contamination was transient in nature and that routine daily disinfection procedures were "
    "effective in eliminating the contamination. Furthermore, no trend was observed in the clean rooms "
    "previous weekly EM data, therefore, no preventive and corrective actions are deemed necessary at this time."
)

for page in doc:
    for w in page.widgets():
        if w.field_name == "Text Field50":
            t50 = w.field_value
            # Replace defensive narrative back to legacy analyst error sentence
            defensive_substr = (
                "is considered an isolated, transient environmental event. "
                "The analyst adhered to all aseptic protocols and standard operating procedures with no deviations noted during interview. "
                "Laboratory error is not considered the assignable cause."
            )
            if defensive_substr in t50:
                t50 = t50.replace(
                    "is considered an isolated, transient environmental event. "
                    "The analyst adhered to all aseptic protocols and standard operating procedures with no deviations noted during interview. "
                    "Laboratory error is not considered the assignable cause.",
                    "may be attributed to a potential analyst error."
                )
            elif "potential analyst error" not in t50:
                # If neither matched, inject legacy_rc at the end of Suite 116 sentence
                s116_marker = "failed testing."
                idx = t50.find(s116_marker)
                if idx != -1:
                    t50 = t50[:idx + len(s116_marker)] + "\r\r" + legacy_rc
            
            # Ensure Suite 116 and formatting fixes remain
            t50 = t50.replace("suite 115", "Suite 116")
            t50 = t50.replace("Suite 115", "Suite 116")
            t50 = t50.replace("MICRO-SOP--9", "MICRO-SOP-9")
            t50 = t50.replace("on 31-May-2026 It was documented", "on 31-May-2026. It was documented")
            
            w.field_value = t50
            w.text_fontsize = 8.0
            w.update()
            print("Updated PDF Text Field50 to Option B (Legacy Root Cause retained).")

temp_out = os.path.join(desktop_dir, "temp_b.pdf")
doc.save(temp_out)
doc.close()

shutil.move(temp_out, pdf_path)
shutil.copyfile(pdf_path, os.path.join(desktop_dir, "OOS-261186.pdf"))
print("Saved updated PDF to Desktop (both standard and short name).")

# ==========================================
# 2. UPDATE EMAIL (Remove Root Cause Row)
# ==========================================
intro_text = (
    "I have reviewed this OOS. All necessary corrections to both the investigation form and the EM table "
    "have already been completed on your behalf (including attaching the finalized EM table as Page 8 and "
    "updating the Gram stain results). Please find the summary of edits made in the table below."
)
closing_text = (
    "@Simin Mohammad all corrections have been completed and incorporated into the attached package. "
    "Please review and sign the updated OOS investigation form: <b>OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf</b>."
)

table_data = [
    (
        "EM Table Attachment & Formatting",
        "The EM Table was omitted from the PDF submission. The draft Word document contained multiple formatting and typographical inconsistencies, including extra spacing in the Table 1 date (07MAY   2026), missing Gram stain identification for Kocuria palustris, a non-standard Table 2 title, and multiple microbial identifications running together without line breaks.",
        "Amended and attached as Page 8 (following Page 7 Version History) to create the complete 8-page package. Standardized Table 1 date to \"07MAY 2026\", updated microbial identification to \"Kocuria palustris, Gram (+) cocci\", revised title to \"Table 2: Environmental Monitoring Plates for Analyst and Cleanroom Bracketing\", and formatted multi-organism entries with clean line breaks.",
        "Document Assembly & Technical Uniformity Correction.\nCompleted on analyst's behalf."
    ),
    (
        "Name of Analyst who Performed the Test",
        "Listed without spacing or line breaks as \"Simin Mohammad(Weekly Active Air Sampling Plate Setup)Sophia Santamaria (Weekly Active Air Sampling Plate Readers)\".",
        "Revised with distinct line breaks, added appropriate role spacing, and corrected plural \"Readers\" to singular \"Reader\":\n\nSimin Mohammad\n(Weekly Active Air Sampling Analyst)\n\nSophia Santamaria\n(Weekly Active Air Sampling Plate Reader)",
        "Formatting & Grammatical Correction."
    ),
    (
        "SOP Reference & Effective Date",
        "Listed as \"MICRO-SOP-2\" Rev 15 with an Effective Date of 05-AUG-2026.",
        "Revised to \"MICRO-SOP-2\" Rev 16 with an Effective Date of 23-Jul-2026.",
        "Compliance Correction."
    ),
    (
        "Limits / Specification",
        "Listed as \"Action level: >10\".",
        "Revised to \"Action level: >= 10 CFU/Plate\".",
        "Technical Specification Correction."
    ),
    (
        "Analyst interviewed? (Comments)",
        "\"Yes, analyst Simin Mohammad, and Sophia Santamaria were comprehensively interviewed.\"",
        "Revised to \"Yes, analysts Simin Mohammad and Sophia Santamaria were interviewed comprehensively.\"",
        "Grammatical Improvement."
    ),
    (
        "Equipment & Calibration Details",
        "Incubator IDs, sensor numbers, and calibration dates were running together without spacing: \"Incubator E001034 (Sensor E001501)Incubator E001031 (Sensor E001505)\" and \"Aug-2026Feb-2027Aug-2026Feb-2027\".",
        "Delineated each incubator and associated sensor with paragraph line breaks, and clearly separated calibration dates:\n\nIncubator E001034 (Sensor E001501)\nIncubator E001031 (Sensor E001505)\n\nAug 2026 / Feb 2027",
        "Critical Readability & Formatting Correction."
    ),
    (
        "Phase I Summary – Plate Setup & Incubation Narrative",
        "Contained multiple missing spaces after periods (\"document.The\", \"system.Active\", \"Facility.The\"), a missing period after Gram stain (\"cocci)To observe\"), an unnecessary comma (\"clean room, were\"), a double hyphen in \"MICRO-SOP--9\", and a typographical error in the reader's surname (\"Sophia Sanatamaria\").",
        "Corrected all punctuation and spacing, updated the spelling to \"Sophia Santamaria\", removed the double hyphen from the SOP reference, and refined sentence structure.",
        "Minor Grammatical & Typographical Correction."
    ),
    (
        "Phase I Summary – Cleanroom Suite Reference",
        "\"It is also important to note that no samples processed in suite 115 for the week of testing (03-May-2026 to 09-May-2026) failed testing.\"",
        "Revised cleanroom reference from Suite 115 to Suite 116: \"It is also important to note that no samples processed in Suite 116 for the week of testing (03-May-2026 to 09-May-2026) failed testing.\"",
        "Critical Assessment Correction."
    )
]

# Update DOCX
doc_email = docx.Document()
section = doc_email.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

p1 = doc_email.add_paragraph()
r1 = p1.add_run("Good morning ")
r1.font.name = "Calibri"; r1.font.size = Pt(11); r1.bold = True
r2 = p1.add_run("@Simin Mohammad")
r2.font.name = "Calibri"; r2.font.size = Pt(11); r2.bold = True; r2.font.color.rgb = RGBColor(0, 102, 204)
r3 = p1.add_run(",")
r3.font.name = "Calibri"; r3.font.size = Pt(11); r3.bold = True

p2 = doc_email.add_paragraph()
r_intro = p2.add_run(intro_text)
r_intro.font.name = "Calibri"; r_intro.font.size = Pt(11)

table = doc_email.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT

tblPr = table._tbl.tblPr
tblInd = OxmlElement('w:tblInd')
tblInd.set(qn('w:w'), '0')
tblInd.set(qn('w:type'), 'dxa')
tblPr.append(tblInd)

headers = ["Section", "Unreviewed", "Reviewed", "Impact"]
hdr_cells = table.rows[0].cells
widths = [Inches(1.5), Inches(2.2), Inches(2.3), Inches(1.5)]

for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].width = widths[i]
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFF00"/>')
    hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
    p = hdr_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = "Calibri"; r.font.size = Pt(10); r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

for row_data in table_data:
    row_cells = table.add_row().cells
    for i in range(4):
        row_cells[i].width = widths[i]
        row_cells[i].text = row_data[i]
        p = row_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Calibri"; r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True

p3 = doc_email.add_paragraph()
r_close1 = p3.add_run("@Simin Mohammad all corrections have been completed and incorporated into the attached package. Please review and sign the updated OOS investigation form: ")
r_close1.font.name = "Calibri"; r_close1.font.size = Pt(11)
r_close2 = p3.add_run("OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf")
r_close2.font.name = "Calibri"; r_close2.font.size = Pt(11); r_close2.bold = True
r_close3 = p3.add_run(".")
r_close3.font.name = "Calibri"; r_close3.font.size = Pt(11)

p4 = doc_email.add_paragraph()
r_thanks = p4.add_run("Thanks,")
r_thanks.font.name = "Calibri"; r_thanks.font.size = Pt(11)

doc_email.save(docx_path)
print("Saved updated DOCX email.")

# Generate HTML
html_rows = ""
for row in table_data:
    u_txt = row[1].replace("\n", "<br>")
    r_txt = row[2].replace("\n", "<br>")
    i_txt = row[3].replace("\n", "<br>")
    html_rows += f"""
    <tr>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; font-weight: bold; vertical-align: top;">{row[0]}</td>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top;">{u_txt}</td>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top;">{r_txt}</td>
      <td style="border: 1px solid #B0B0B0; padding: 6px 8px; vertical-align: top;">{i_txt}</td>
    </tr>"""

full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #000; line-height: 1.5; }}
  .email-container {{ max-width: 950px; margin: 10px 0; }}
  .tag {{ color: #0066cc; font-weight: bold; }}
  table {{ border-collapse: collapse; width: 100%; font-family: Calibri, Arial, sans-serif; font-size: 10pt; margin: 10px 0 15px 0; }}
  th {{ background-color: #FFFF00 !important; color: #000; font-weight: bold; text-align: center; border: 1px solid #B0B0B0; padding: 6px 8px; }}
</style>
</head>
<body>
<div class="email-container">
  <p>Good morning <span class="tag">@Simin Mohammad</span>,</p>
  <p>{intro_text}</p>
  <table>
    <thead>
      <tr>
        <th style="width: 20%;">Section</th>
        <th style="width: 30%;">Unreviewed</th>
        <th style="width: 32%;">Reviewed</th>
        <th style="width: 18%;">Impact</th>
      </tr>
    </thead>
    <tbody>
      {html_rows}
    </tbody>
  </table>
  <p>{closing_text}</p>
  <p>Thanks,</p>
</div>
</body>
</html>"""

with open(html_path, "w", encoding="utf-8") as f:
    f.write(full_html)
print("Saved updated HTML email.")

# Update Windows Clipboard
def set_html_clipboard(html_fragment):
    MARKER_BLOCK = (
        "Version:0.9\r\n"
        "StartHTML:{start_html:08d}\r\n"
        "EndHTML:{end_html:08d}\r\n"
        "StartFragment:{start_frag:08d}\r\n"
        "EndFragment:{end_frag:08d}\r\n"
    )
    prefix = "<html><body><!--StartFragment-->"
    suffix = "<!--EndFragment--></body></html>"
    dummy = MARKER_BLOCK.format(start_html=0, end_html=0, start_frag=0, end_frag=0)
    start_html = len(dummy)
    start_frag = start_html + len(prefix)
    end_frag = start_frag + len(html_fragment.encode("utf-8"))
    end_html = end_frag + len(suffix.encode("utf-8"))
    header = MARKER_BLOCK.format(start_html=start_html, end_html=end_html, start_frag=start_frag, end_frag=end_frag)
    full_data = (header + prefix + html_fragment + suffix).encode("utf-8")

    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        cf_html = win32clipboard.RegisterClipboardFormat("HTML Format")
        win32clipboard.SetClipboardData(cf_html, full_data)
        plain = (
            f"Good morning @Simin Mohammad,\n\n"
            f"{intro_text}\n\n"
            f"@Simin Mohammad all corrections have been completed and incorporated into the attached package. "
            f"Please review and sign the updated OOS investigation form: OOS-261186 EM SMO 116A Air 07MAY2026 - EM.pdf.\n\n"
            f"Thanks,"
        )
        win32clipboard.SetClipboardText(plain)
    finally:
        win32clipboard.CloseClipboard()

body_start = full_html.find('<div class="email-container">')
body_end = full_html.find('</div>', body_start) + 6
set_html_clipboard(full_html[body_start:body_end])
print("Clipboard refreshed with Option B email.")
