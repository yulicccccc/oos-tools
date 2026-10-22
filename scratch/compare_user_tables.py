import docx

path = r'C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\Tables OOS-262333 Optimal Balance Pharmacy (E19193) - Celsis.docx'
doc = docx.Document(path)

for t_idx, t in enumerate(doc.tables):
    print(f'\n=== TABLE {t_idx+1} ===')
    for r_idx, row in enumerate(t.rows):
        row_cells = []
        for c_idx, cell in enumerate(row.cells):
            cell_runs = []
            for p in cell.paragraphs:
                for r in p.runs:
                    cell_runs.append(f'{r.text}')
            # also check hyperlinks
            for p in cell.paragraphs:
                for child in p._p:
                    if child.tag.endswith('hyperlink'):
                        link_text = ''.join(child.itertext())
                        cell_runs.append(f'[HL: {link_text}]')
            row_cells.append(' '.join(cell_runs))
        # dedup merged cells
        dedup = []
        for x in row_cells:
            if not dedup or x != dedup[-1]:
                dedup.append(x)
        print(f'R{r_idx:02d}: {dedup}')
