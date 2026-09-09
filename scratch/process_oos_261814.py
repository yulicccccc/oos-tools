import os, sys, json, re, io, copy, shutil
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
from docxtpl import DocxTemplate
from pypdf import PdfWriter, PdfReader
from datetime import datetime
import win32com.client
from utils import get_room_logic, get_cleanroom_narrative, ordinal, num_to_words, get_full_name

SAVE_FILE = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\SAVE_OOS-261814 GoGoMeds Select (E10747) - ScanRDI.txt"
EM_DOCX_FILE = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\EM Table OOS-261814 07AUG2026.docx"
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"
OUTPUT_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\scratch"
QYC_PDF_PATH = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop\OOS-261814 GoGoMeds Select (E10747) - ScanRDI - QYC.pdf"
CORP_FORM_DOWNLOADED = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\CORP-FORM-21 Laboratory OOS Investigation Form (v11.1) (1).pdf"

# 1. Load base save
with open(SAVE_FILE, "r", encoding="utf-8") as f:
    data = json.load(f)

# 2. Update Personnel and Equipment for GA in L-Suite BSC 1937
data['sample_url'] = data.get('sample_url') or "https://etrax.eagleanalytical.com/Submission/Details/RjtF7Gpyb7QIgSCtaB%24Udw__"
data['prepper_initial'] = 'EN'
data['prepper_name'] = 'Elysse Nioupin'
data['analyst_initial'] = 'SU'
data['analyst_name'] = 'Sonal Uprety'
data['changeover_initial'] = 'GA'
data['changeover_name'] = 'Gerald Anyangwe'
data['reader_initial'] = 'VV'
data['reader_name'] = 'Varsha Subramanian'
data['bsc_id'] = '1313'
data['chgbsc_id'] = '1937'
data['diff_changeover_analyst'] = 'Yes'
data['diff_changeover_bsc'] = 'Yes'
data['organism_morphology'] = 'Rod'
data['control_positive'] = data.get('control_pos', 'A. brasiliensis')
data['control_data'] = data.get('control_exp', '22May28')

# Room logic
t_room, t_suite, t_suffix, t_loc = get_room_logic(data['bsc_id'])
c_room, c_suite, c_suffix, c_loc = get_room_logic(data['chgbsc_id'])
t_suite_phrase = f"Suite {t_suite}{t_suffix}" if t_suite != "L-Suite" else "L-Suite"
c_suite_phrase = f"Suite {c_suite}{c_suffix}" if c_suite != "L-Suite" else "L-Suite"

data['cr_id'] = t_room
data['cr_suit'] = t_suite
data['suit'] = t_suffix
data['bsc_location'] = t_loc
data['smart_bsc_bracketing_header'] = f"Biological Safety Cabinet EM Bracketing Biological Safety Cabinet (BSC) E00{data['bsc_id']} and E00{data['chgbsc_id']}"

d_obj = datetime.strptime(data['test_date'], "%d%b%y")
tr_id = f"{d_obj.strftime('%m%d%y')}-{data['scan_id']}-{data['shift_number']}"
pdf_date_str = d_obj.strftime("%d-%b-%Y")
data['test_record'] = tr_id

# 3. Cleanroom and Equipment Narrative (Dual-Suite / Dual-BSC with MICRO-SOP-9)
p1a = get_cleanroom_narrative(t_suite, t_room=t_room, action_text="testing", verb="consists of")
p1b = get_cleanroom_narrative(c_suite, t_room=c_room, action_text="changeover", verb="consists of")
intro = (
    f"The ISO 5 BSC E00{data['bsc_id']}, located in the {t_loc}, ({t_suite_phrase}), and "
    f"ISO 5 BSC E00{data['chgbsc_id']}, located in the {c_loc}, ({c_suite_phrase}), were thoroughly cleaned and "
    "disinfected prior to their respective procedures in accordance with MICRO-SOP-9, Cleaning and Disinfecting Procedure "
    f"for Microbiology. Furthermore, the BSCs used throughout testing, E00{data['bsc_id']} for sample processing and "
    f"E00{data['chgbsc_id']} for the changeover step, were certified and approved by both the Engineering and Quality Assurance teams."
)
usage_sent = (
    f"Sample processing was conducted within the ISO 5 BSC in the {t_loc} ({t_suite_phrase}, BSC E00{data['bsc_id']}) "
    f"by {data['analyst_name']} and the changeover step was conducted within the ISO 5 BSC in the {c_loc} "
    f"({c_suite_phrase}, BSC E00{data['chgbsc_id']}) by {data['changeover_name']} on {data['test_date']}."
)
data['equipment_summary'] = f"{p1a}\n\n{p1b}\n\n{intro} {usage_sent}"

# 4. Personnel block & interview
names_only_phrase = "Elysse Nioupin, Sonal Uprety, Gerald Anyangwe, and Varsha Subramanian"
analysts_with_prefix_phrase = f"analysts {names_only_phrase}"
smart_comment_interview = f"Yes, {analysts_with_prefix_phrase} were interviewed comprehensively."
data['smart_comment_interview'] = smart_comment_interview

analyst_sig_text = f"{data['analyst_name']} (Written by: Qiyue Chen)"
smart_personnel_block = (
    f"Prepper: \n{data['prepper_name']} ({data['prepper_initial']})\n\n"
    f"Processor:\n{data['analyst_name']} ({data['analyst_initial']})\n\n"
    f"Changeover Processor:\n{data['changeover_name']} ({data['changeover_initial']})\n\n"
    f"Reader:\n{data['reader_name']} ({data['reader_initial']})"
)
smart_incident_opening = f"On {data['test_date']}, sample {data['sample_id']} was found positive for viable microorganisms after ScanRDI testing."
smart_comment_samples = f"Yes, {data['sample_id']}"
smart_comment_records = f"Yes, See {tr_id} for more information."
smart_comment_storage = f"Yes, Information is available in Eagle Trax Sample Location History under {data['sample_id']}"

p1 = f"All analysts involved in the prepping, processing, changeover, and reading of the samples – {names_only_phrase} – were interviewed and their answers are recorded throughout this document."
p2 = f"The sample was stored upon arrival according to the Client’s instructions. Analysts {data['prepper_name']} and {data['analyst_name']} confirmed the integrity of the samples throughout both the preparation and processing stages. No leaks or turbidity were observed at any point, verifying the integrity of the sample."
p3 = "All reagents and supplies mentioned in the material section above were stored according to the suppliers’ recommendations, and their integrity was visually verified before utilization. Moreover, each reagent and supply had valid expiration dates."
p4 = (
    f"During the preparation phase, {data['prepper_name']} disinfected the samples using acidified bleach and placed them into a pre-disinfected storage bin. "
    f"On {data['test_date']}, prior to sample processing, {data['analyst_name']} performed a second disinfection with acidified bleach, allowing a minimum contact time of 10 minutes "
    "before transferring the samples into the cleanroom suites. A final disinfection step was completed immediately before the samples were introduced into the "
    f"ISO 5 Biological Safety Cabinet (BSC), E00{data['bsc_id']}, located within the {t_loc}, ({t_suite_phrase}). All activities were performed in accordance with "
    "MICRO-SOP-12, Rapid Scan RDI® Test using FIFU Method."
)
p5 = data['equipment_summary']
p6 = (
    f"The analyst, {data['reader_name']}, confirmed that the equipment was set up as per ENG-SOP-4 "
    "(Scan RDI® System – Operations (Standard C3 Quality Check and Microscope Setup) and Maintenance), and the negative control "
    f"and the positive control for the analyst, {data['reader_name']}, yielded expected results."
)
smart_phase1_part1 = "\n\n".join([p1, p2, p3, p4, p5, p6])
data['smart_phase1_part1'] = smart_phase1_part1
data['smart_phase1_summary'] = smart_phase1_part1

# 5. History & Cross-Contamination
data['sample_history_paragraph'] = f"Analyzing a 6-month sample history for {data['client_name']}, this specific analyte \"{data['sample_name']}\" has had no prior failures using the Scan RDI method during this period."
data['cross_contamination_summary'] = "To evaluate the potential for sample-to-sample contamination, all samples processed on the same day were reviewed. All other samples processed by the same analyst and by other analysts on that day yielded negative results, indicating that cross-contamination is unlikely."

# 6. FDA-Aligned cGMP Defense Engine & EM Details (Dual-BSC & Dual-Suite with Fungal ID Defense)
p_transposition_1 = (
    f"During the OOS investigation and subsequent review of the testing records, a result-transposition event was identified "
    f"involving {data['sample_id']} and ETX-260804-0101. Processing Analyst SU completed ScanRDI sterility testing for sample "
    f"ETX-260804-0101, associated with DCA Pharmacy (E12860), and sample {data['sample_id']}, associated with {data['client_name']}, "
    "in accordance with MICRO-SOP-12, ScanRDI Sterility Testing Process. During sample reading and result verification using "
    "the ScanRDI instrument, Second Shift Reading Analyst VV determined that ETX-260804-0101 met the established acceptance "
    f"criteria and generated a passing result, whereas {data['sample_id']} generated a failing result. During subsequent result "
    "entry, Analyst VV inadvertently transposed the results between the two samples. As a result, the passing result generated "
    f"for ETX-260804-0101 was incorrectly assigned to {data['sample_id']}, while the failing result generated for {data['sample_id']} "
    "was incorrectly assigned to ETX-260804-0101. During the subsequent data review, reviewer OA did not identify the result "
    "transposition and approved the incorrectly assigned failing result for ETX-260804-0101. Consequently, ETX-260804-0101, "
    "which had actually met the acceptance criteria, was placed on hold pending an Out-of-Specification (OOS) investigation."
)

p_transposition_2 = (
    "Review of the original ScanRDI testing and instrument-generated data confirmed that the testing process itself was "
    "performed as intended and that the instrument results were correctly generated for each sample. The discrepancy occurred "
    "after testing, during manual result transcription and subsequent review, and did not reflect an analytical failure or an "
    "actual failing result for ETX-260804-0101. The event was therefore identified as an unintentional result-transposition "
    "error during data entry that was not detected during the initial data review. The original instrument-generated result "
    f"for ETX-260804-0101 was confirmed as passing, while {data['sample_id']} was confirmed as the sample associated with "
    "the failing ScanRDI result. Following confirmation of the correct result attribution, the investigation proceeded to "
    f"evaluate potential laboratory sources of the microbial recovery for {data['sample_id']}."
)

p_em_intro = (
    f"Upon review of the environmental monitoring data associated with the sterility test, no microbial growth was recovered from "
    f"any of the four ISO 5 work-surface monitoring locations within BSC E00{data['bsc_id']} (Suite 115A) or "
    f"BSC E00{data['chgbsc_id']} (L-Suite Room 144) on the date of testing. In addition, no microbial growth was observed on the "
    f"changeover personnel monitoring (analyst {data['changeover_name']}) or settling plates within BSC E00{data['chgbsc_id']}. "
    f"Low-level microbial recoveries were identified from processing personnel and settling-plate monitoring within BSC E00{data['bsc_id']}, "
    "as well as from routine weekly monitoring of the surrounding cleanroom areas."
)

p_personnel = (
    "One CFU was recovered from Processing Analyst SU's right-hand touch plate and submitted for microbial identification "
    "under sample ID ETX-260817-0447. The isolate was identified as Micrococcus luteus. This recovery was not microbiologically "
    "consistent with the organism observed in the test sample: Micrococcus luteus is coccal in morphology, whereas the microorganism "
    "recovered from the test sample exhibited rod-shaped morphology. Accordingly, the personnel monitoring result does not support "
    "direct transfer of the recovered right-glove organism to the test sample. In contrast, personnel monitoring for Changeover Analyst GA "
    "yielded no microbial growth."
)

p_settling = (
    f"Two CFUs were also recovered from settling plate Sett 1 within BSC E00{data['bsc_id']} and submitted for microbial identification "
    "under sample ID ETX-260817-0507. The isolates were identified as Sporisorium graminicola and Ustilago maydis. These microorganisms "
    "are fungal (smut) species, which are taxonomically and morphologically distinct from the bacterial rod-shaped microorganisms "
    "recovered from the test sample. Additionally, the settling plate was documented as desiccated at the 5-day read. Therefore, "
    "no microbiological match exists between the settling-plate isolates and the test-sample contaminant. Settling plates within "
    f"changeover BSC E00{data['chgbsc_id']} yielded no growth. This finding was evaluated in conjunction with the remaining contemporaneous "
    f"environmental monitoring data, including the complete absence of microbial recovery from all ISO 5 work surfaces within both "
    f"BSC E00{data['bsc_id']} and BSC E00{data['chgbsc_id']}."
)

p_weekly = (
    "Routine weekly facility monitoring during the week of testing recovered 1 CFU of Gram-positive short rods from active air "
    "monitoring in ISO 8 Cleanroom 115 (ETX-260817-0370) and 2 CFUs identified as Bacillus megaterium from the floor of ISO 7 "
    "Suite 115A (ETX-260817-0366). In the L-Suite, routine weekly surface sampling of cleanroom Room 144 (CR1978) yielded no microbial "
    "growth, while active air monitoring recovered typical human-associated flora (Staphylococcus, Micrococcus, Corynebacterium spp.) "
    "strictly confined to lower-classified background areas (ISO 8 anteroom 143 and outer room 142). All of these recoveries occurred "
    "in lower-classified background areas physically separated from the critical ISO 5 testing zones. Furthermore, the positive pressure cascade "
    "(flowing outwards from ISO 7 rooms 145 and 144 toward ISO 8 rooms 143 and 142) and the transfer of samples in disinfected, closed containers "
    "provide robust barriers preventing ingress of airborne contaminants into the ISO 5 work areas. The complete absence of microbial recovery "
    "on all ISO 5 critical surfaces within both BSCs further demonstrates that environmental controls effectively prevented transfer into the "
    "critical testing zone."
)

p_monthly = (
    f"Monthly cleaning and disinfection, using H2O2, of the cleanrooms (ISO 7) and their containing Biosafety Cabinets (BSCs, ISO 5) "
    f"were performed on {data['monthly_cleaning_date']}, as per MICRO-SOP-9, Cleaning and Disinfecting Procedure for Microbiology. "
    "It was documented that all H2O2 indicators passed."
)

p_history = data['sample_history_paragraph']
p_cross = data['cross_contamination_summary']
p_conclusion = (
    "Based on the cumulative evidence, it is highly unlikely that the failing results were due to reagents, supplies, "
    "the cleanroom environment, the process, or analyst involvement. Consequently, the possibility of laboratory error "
    "contributing to this failure is minimal. Therefore, the original test result is deemed valid."
)

p7 = f"On {data['test_date']}, a rapid sterility test was conducted on the sample using the ScanRDI method. The sample was initially prepared by Analyst {data['prepper_name']}, processed by {data['analyst_name']}, and subsequently read by {data['reader_name']}. The test revealed {data['confirm_number']} rod-shaped viable microorganisms, see Table 1."
p8 = f"Table 2 (see attached tables) presents the environmental monitoring results for {data['sample_id']}. The environmental monitoring (EM) plates were incubated for no less than 48 hours at 30–35°C and for no less than an additional five days at 20–25°C, as per SOP 2.600.002, Environmental Monitoring of the Clean-room Facility."

smart_phase1_part2_page5 = "\r \r".join([
    p7,
    p_transposition_1,
    p_transposition_2,
    p8,
    p_em_intro
])

smart_phase1_part2_page6 = "\r \r".join([
    p_personnel,
    p_settling,
    p_weekly,
    p_monthly,
    p_history,
    p_cross,
    p_conclusion
])

smart_phase1_part2 = "\n\n".join([
    p7,
    p_transposition_1,
    p_transposition_2,
    p8,
    p_em_intro,
    p_personnel,
    p_settling,
    p_weekly,
    p_monthly,
    p_history,
    p_cross,
    p_conclusion
])

data['narrative_summary'] = smart_phase1_part2
data['smart_justification'] = p_conclusion
data['smart_phase1_part2'] = smart_phase1_part2
data['smart_phase1_continued'] = smart_phase1_part2
data['smart_cr_id'] = f"CR{t_suite} (E00{t_room}) & CR{c_room} (L-Suite)"
data['smart_scan_id'] = f"E00{data['scan_id']}"
data['analyst_signature'] = analyst_sig_text
data['report_header'] = f"{data['sample_id']}\n\n{data['client_name']}"
data['smart_personnel_block'] = smart_personnel_block
data['smart_incident_opening'] = smart_incident_opening
data['smart_comment_samples'] = smart_comment_samples
data['smart_comment_records'] = smart_comment_records
data['smart_comment_storage'] = smart_comment_storage

# Save updated JSON state
out_json_path = os.path.join(OUTPUT_DIR, f"UPDATED_SAVE_OOS-{data['oos_id']} {data['client_name']} - ScanRDI.txt")
with open(out_json_path, "w", encoding="utf-8") as f:
    json.dump(data, f, indent=2)
print("Saved updated JSON state to:", out_json_path)

# 7. Helper functions for docx formatting
def set_cell_hyperlink(cell, url, text):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    part = p.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink_xml = (
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rStyle w:val="Hyperlink"/>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="467886"/>'
        f'<w:sz w:val="14"/>'
        f'<w:szCs w:val="14"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    p._p.append(parse_xml(hyperlink_xml))

def update_cell_text(cell, text, bold=False, italic=False, font_size=Pt(7)):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic

def build_corrected_em_table_element():
    doc_raw_em = docx.Document(EM_DOCX_FILE)
    t = doc_raw_em.tables[0]
    
    update_cell_text(t.rows[0].cells[2], 'Date (DDMMM YY)', bold=True, font_size=Pt(6))
    
    # SU Personal
    update_cell_text(t.rows[2].cells[2], '07Aug26')
    update_cell_text(t.rows[2].cells[8], 'Plate was desiccated on 5 day read')
    
    # GA Personal
    update_cell_text(t.rows[3].cells[2], '07Aug26')
    
    # BSC Header
    update_cell_text(t.rows[4].cells[0], 'Biological Safety Cabinet EM Bracketing Biological Safety Cabinet (BSC) E001313 and E001937', bold=True, font_size=Pt(7))
    
    # SU Surface BSC 1313
    update_cell_text(t.rows[5].cells[0], 'Surface Sampling of ISO 5 E001313 (4 locations)')
    update_cell_text(t.rows[5].cells[2], '07Aug26')
    
    # GA Surface BSC 1937
    update_cell_text(t.rows[6].cells[0], 'Surface Sampling of ISO 5 E001937 (4 locations)')
    update_cell_text(t.rows[6].cells[2], '07Aug26')
    
    # SU Settling BSC 1313 - with newly identified fungal species
    update_cell_text(t.rows[7].cells[0], 'Settling Sampling of ISO 5 E001313 (2 locations)')
    update_cell_text(t.rows[7].cells[2], '07Aug26')
    update_cell_text(t.rows[7].cells[7], 'Sporisorium graminicola\nUstilago maydis')
    update_cell_text(t.rows[7].cells[8], 'Plate was desiccated on 5 day read')
    
    # GA Settling BSC 1937
    update_cell_text(t.rows[8].cells[0], 'Settling Sampling of ISO 5 E001937 (2 locations)')
    update_cell_text(t.rows[8].cells[2], '07Aug26')
    
    # Header Suite 115 Air
    update_cell_text(t.rows[9].cells[0], 'Weekly Active Air Sampling Bracketing - Cleanroom 115 - CR1737', bold=True, font_size=Pt(7))
    
    # Suite 115 Air
    update_cell_text(t.rows[10].cells[2], '06Aug26')
    
    # Header L-Suite Air
    update_cell_text(t.rows[11].cells[0], 'Weekly Active Air Sampling Bracketing - Cleanroom 144 - CR1978 (L-Suite)', bold=True, font_size=Pt(7))
    
    # L-Suite Air
    update_cell_text(t.rows[12].cells[2], '07Aug26')
    
    # Header Suite 115 Surface
    update_cell_text(t.rows[13].cells[0], 'Surface Sampling of Anteroom and Cleanroom Bracketing - Cleanroom 115 - CR1737', bold=True, font_size=Pt(7))
    
    # Suite 115 Surface
    update_cell_text(t.rows[14].cells[2], '06Aug26')
    
    # Header L-Suite Surface (Corrected title)
    update_cell_text(t.rows[15].cells[0], 'Surface Sampling of Anteroom and Cleanroom Bracketing - Cleanroom 144 - CR1978 (L-Suite)', bold=True, font_size=Pt(7))
    
    # L-Suite Surface
    update_cell_text(t.rows[16].cells[2], '07Aug26')
    
    return copy.deepcopy(t._element)

# 8. Render Standalone Tables Document
out_doc_tables = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - ScanRDI.docx")
doc_tables = docx.Document("tables for scan.docx")

# Remove Table 2 (trend table)
if len(doc_tables.tables) >= 3:
    t_trend = doc_tables.tables[2]._element
    t_trend.getparent().remove(t_trend)
for p in list(doc_tables.paragraphs):
    if 'In Trend of Past OOS Results' in p.text:
        p._element.getparent().remove(p._element)

# Fill Table 0 (Sample Information)
t0 = doc_tables.tables[0]
update_cell_text(t0.rows[1].cells[0], data['analyst_name'], font_size=Pt(8))
update_cell_text(t0.rows[1].cells[1], data['reader_name'], font_size=Pt(8))
set_cell_hyperlink(t0.rows[1].cells[2], data['sample_url'], data['sample_id'])
update_cell_text(t0.rows[1].cells[3], str(data.get('event_number', '187')), font_size=Pt(8))
update_cell_text(t0.rows[1].cells[4], str(data.get('confirm_number', '4')), font_size=Pt(8))
update_cell_text(t0.rows[1].cells[5], f"{data['organism_morphology']}-shaped Morphology", font_size=Pt(8))

# Replace Table 1 with the 17-row EM table
t1_old = doc_tables.tables[1]._element
t1_parent = t1_old.getparent()
idx1 = t1_parent.index(t1_old)
t1_parent.remove(t1_old)
t_new_em = build_corrected_em_table_element()
t1_parent.insert(idx1, t_new_em)

doc_tables.save(out_doc_tables)
print("Saved Standalone Tables Docx to:", out_doc_tables)

# 9. Export Standalone Tables Docx to PDF via Word COM
out_pdf_tables = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - ScanRDI.pdf")
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
try:
    d = word.Documents.Open(os.path.abspath(out_doc_tables), ReadOnly=True)
    d.SaveAs(os.path.abspath(out_pdf_tables), FileFormat=17)
    d.Close(SaveChanges=False)
    print("Saved Standalone Tables PDF via Word to:", out_pdf_tables)
finally:
    word.Quit()

r_check = PdfReader(out_pdf_tables)
print(f"Verified Tables PDF Page Count: {len(r_check.pages)}")

# 10. Render Master Word Report (ScanRDI OOS P1 template.docx)
primary_tpl = "ScanRDI OOS P1 template.docx" if os.path.exists("ScanRDI OOS P1 template.docx") else "ScanRDI OOS template 0.docx"
tpl_report = DocxTemplate(primary_tpl)
tpl_report.render(data)

# Remove Table 3 (trend table)
if len(tpl_report.docx.tables) >= 4:
    t3 = tpl_report.docx.tables[3]._element
    t3.getparent().remove(t3)
for p in list(tpl_report.docx.paragraphs):
    if 'In Trend of Past OOS Results' in p.text:
        p._element.getparent().remove(p._element)

# In Table 1: set hyperlink on Sample ID
if len(tpl_report.docx.tables) >= 2:
    set_cell_hyperlink(tpl_report.docx.tables[1].rows[1].cells[2], data['sample_url'], data['sample_id'])

# In Table 2: replace with the 17-row EM table
if len(tpl_report.docx.tables) >= 3:
    t2_rep = tpl_report.docx.tables[2]._element
    t2_parent = t2_rep.getparent()
    idx2 = t2_parent.index(t2_rep)
    t2_parent.remove(t2_rep)
    t2_parent.insert(idx2, build_corrected_em_table_element())

out_doc_report = os.path.join(OUTPUT_DIR, f"OOS-{data['oos_id']} {data['client_name']} - ScanRDI.docx")
tpl_report.save(out_doc_report)
print(f"Saved Master Word Report to: {out_doc_report}")

# 11. Populate CORP-FORM-21 PDF and Append Page 7 Tables PDF
pdf_map = {
    'Text Field57': data['oos_id'],
    'Date Field0': pdf_date_str,
    'Date Field1': pdf_date_str,
    'Date Field2': pdf_date_str,
    'Date Field3': pdf_date_str,
    'Text Field2': data['sample_id'],
    'Text Field6': data['lot_number'],
    'Text Field4': data['sample_name'] + "\n\n\n\n",
    'Text Field5': data['dosage_form'],
    'Text Field8': "MICRO-SOP-12 (16)\rENG-SOP-4 (05)",
    'Text Field9': "24Jul26\r24Jul26",
    'Text Field10': "Rev: 16\rRev: 05",
    'Text Field15': "Yes, as per MICRO-SOP-12, ENG-SOP-4",
    'Text Field16': "Yes, as per MICRO-SOP-12, ENG-SOP-4",
    'Text Field0': analyst_sig_text,
    'Text Field3': smart_personnel_block,
    'Text Field7': smart_incident_opening + "\n\n",
    'Text Field13': smart_comment_interview,
    'Text Field14': smart_comment_samples,
    'Text Field17': smart_comment_records,
    'Text Field21': smart_comment_storage,
    'Text Field30': f"E00{data['scan_id']}",
    'Text Field32': f"E001737 (CR115) & E001978 (L-Suite)",
    'Text Field34': f"E00{data['scan_id']}",
    'Text Field24': data['control_pos'],
    'Text Field25': data['control_lot'] + "\n\n\n",
    'Text Field26': data['control_exp'],
    'Text Field48': f"N/A QYC {datetime.now().strftime('%d%b%y')}",
    'Text Field49': smart_phase1_part1,
    'Text Field50': smart_phase1_part2_page5,
    'Text Field51': smart_phase1_part2_page6
}

# Pull prefilled boilerplate checkboxes and fields from QYC.pdf / template
prefilled_boilerplate = {}
source_prefill = QYC_PDF_PATH if os.path.exists(QYC_PDF_PATH) else "ScanRDI OOS template.pdf"
if os.path.exists(source_prefill):
    reader_pre = PdfReader(source_prefill)
    fields_pre = reader_pre.get_fields()
    if fields_pre:
        prefilled_boilerplate = {k: v.get('/V') for k, v in fields_pre.items() if v.get('/V') is not None}

combined_pdf_map = {**prefilled_boilerplate, **pdf_map}

base_pdf_source = CORP_FORM_DOWNLOADED if os.path.exists(CORP_FORM_DOWNLOADED) else source_prefill
writer = PdfWriter(clone_from=base_pdf_source)
for p in writer.pages:
    writer.update_page_form_field_values(p, combined_pdf_map)

# Append Tables PDF as Page 7
if os.path.exists(out_pdf_tables):
    table_reader = PdfReader(out_pdf_tables)
    for table_page in table_reader.pages:
        writer.add_page(table_page)
    print(f"Appended Tables PDF to form. Total pages: {len(writer.pages)}")

out_pdf_report = os.path.join(OUTPUT_DIR, f"OOS-{data['oos_id']} {data['client_name']} - ScanRDI.pdf")
with open(out_pdf_report, "wb") as f:
    writer.write(f)
print("Saved complete 7-Page PDF Report to:", out_pdf_report)

# 12. Deploy / Sync all outputs to Documents directory
files_to_sync = [
    out_doc_report,
    out_doc_tables,
    out_pdf_report,
    out_pdf_tables,
    out_json_path
]
for fpath in files_to_sync:
    if fpath and os.path.exists(fpath):
        dest = os.path.join(DOCUMENTS_DIR, os.path.basename(fpath))
        try:
            shutil.copy2(fpath, dest)
            print(f"Synced to Documents: {dest}")
        except PermissionError:
            print(f"Notice: {dest} is currently open in another application, skipped overwrite.")
        except Exception as e:
            print(f"Warning: Could not sync {fpath} to {dest}: {e}")

print("\n--- ALL GENERATION AND DEPLOYMENT COMPLETED SUCCESSFULLY! ---")
