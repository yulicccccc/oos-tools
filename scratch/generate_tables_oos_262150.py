"""
Script: generate_tables_oos_262150.py
Purpose: Generate Standalone Tables (Word .docx and PDF) for Scan RDI OOS-262150
Client: BelieveRX Pharmacy (E74120)
Sample ID: ETX-260914-0306
Sample Name: Tirzepatide 5mg, Glycine 0.5mg/mL
Lot: 091126-66A
Test Date: 16Sep26
Processing Analyst: Elizabeth Bennett (ELB) (47th sample processed)
Processing BSC: BSC E001938 (Cleanroom 145 / L-Suite)
Changeover Analyst: Elizabeth Bennett (ELB)
Changeover BSC: BSC E001937 (Buffer Room 144 / L-Suite)
Reading Analyst: Sonal Uprety (SU)
"""

import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import fitz

sys.stdout.reconfigure(encoding='utf-8')

DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
SCRATCH_DIR = r"c:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\scratch"

OOS_ID = "262150"
CLIENT_NAME = "BelieveRX Pharmacy (E74120)"
SAMPLE_ID = "ETX-260914-0306"
SAMPLE_URL = "https://etrax.eagleanalytical.com/Submission/Details/j$B8MDDTkRSGXhr3At81EQ__"
SAMPLE_NAME = "Tirzepatide 5mg, Glycine 0.5mg/mL"
LOT_NUMBER = "091126-66A"
TEST_DATE = "16Sep26"
TEST_DATE_FULL = "16Sep2026"

PROCESSING_ANALYST = "Elizabeth Bennett"
PROCESSING_INITIAL = "ELB"
PROCESSING_BSC = "1938"

CHANGEOVER_ANALYST = "Elizabeth Bennett"
CHANGEOVER_INITIAL = "ELB"
CHANGEOVER_BSC = "1937"

READING_ANALYST = "Sonal Uprety"
EVENT_COUNT = "56"
CONFIRMED_COUNT = "5"
MORPHOLOGY = "Cocci shaped morphology"

# XML Helper Functions
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=25, bottom=25, left=35, right=35):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_hyperlink(paragraph, url, text, font_size=Pt(6.5), color_rgb="2980B9"):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size.pt * 2)}"/>'))
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{color_rgb}"/>'))
    rPr.append(parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>'))
    new_run.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def format_cell(cell, text, bold=False, italic=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None, v_align=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = v_align
    set_cell_margins(cell)
    if fill_hex:
        set_cell_background(cell, fill_hex)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = font_size
    r.bold = bold
    r.italic = italic
    return r

def build_tables_doc(output_docx_path):
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    # ---------------- Table 1 ----------------
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"Table 1: Information for {SAMPLE_ID} under investigation")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(8.5)
    r_t1.underline = True

    t1 = doc.add_table(rows=2, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = [
        "Processing Analyst",
        "Reading Analyst",
        "Sample ID",
        "Events",
        "Confirmed Microbial\nEvents",
        "Morphology Description"
    ]
    t1_widths = [Inches(1.3), Inches(1.3), Inches(1.1), Inches(0.7), Inches(1.2), Inches(2.1)]

    for j, h in enumerate(t1_headers):
        cell = t1.cell(0, j)
        format_cell(cell, h, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

    format_cell(t1.cell(1, 0), PROCESSING_ANALYST, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 1), READING_ANALYST, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

    c_sid = t1.cell(1, 2)
    c_sid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_sid)
    p_sid = c_sid.paragraphs[0]
    p_sid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sid.paragraph_format.space_before = Pt(0)
    p_sid.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_sid, SAMPLE_URL, SAMPLE_ID, font_size=Pt(6.5))

    format_cell(t1.cell(1, 3), str(EVENT_COUNT), bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 4), str(CONFIRMED_COUNT), bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t1.cell(1, 5), MORPHOLOGY, bold=True, italic=True, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in t1.rows:
        for j, w in enumerate(t1_widths):
            row.cells[j].width = w

    # Space between tables
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(0)

    # ---------------- Table 2 ----------------
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run(f"Table 2: Environmental Monitoring from Testing Performed on {TEST_DATE}")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(8.5)
    r_t2.underline = True

    t2 = doc.add_table(rows=0, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    t2_widths = [
        Inches(1.05), # Sampling site
        Inches(0.60), # Frequency
        Inches(0.70), # Date (DDMMMYY)
        Inches(0.55), # Analyst
        Inches(1.05), # Day / Week(s)
        Inches(1.10), # Observation
        Inches(0.85), # EM Plate ETX Number
        Inches(1.20), # Microbial ID
        Inches(0.60)  # Notes
    ]

    hdr_row = t2.add_row()
    t2_headers = [
        "Environmental\nMonitoring (EM)\nSampling Site",
        "Frequency",
        "Date\n(DDMMMYY)",
        "Analyst",
        "Day /Week(s)",
        "Observation",
        "EM Plate ETX\nNumber",
        "Microbial ID",
        "Notes"
    ]
    for j, h in enumerate(t2_headers):
        format_cell(hdr_row.cells[j], h, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

    def add_section_divider(table, title_text):
        row = table.add_row()
        cell = row.cells[0]
        for c in row.cells[1:]:
            cell.merge(c)
        format_cell(cell, title_text, bold=True, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex="E8E8E8")
        return row

    def add_data_row(table, site, freq, date_str, analyst, day_week, obs, etx_num, micro_id, notes, etx_url=None):
        row = table.add_row()
        format_cell(row.cells[0], site, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], freq, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[2], date_str, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[3], analyst, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[4], day_week, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[5], obs, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

        if etx_url:
            c = row.cells[6]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(c)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_hyperlink(p, etx_url, etx_num, font_size=Pt(6.5))
        else:
            format_cell(row.cells[6], etx_num, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)

        format_cell(row.cells[7], micro_id, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[8], notes, bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER)
        return row

    # --- Section 1: Personnel EM ---
    add_section_divider(t2, f"Personnel EM for {TEST_DATE}")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", TEST_DATE, PROCESSING_INITIAL, "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, "Personal (Left Touch\nand Right Touch)", "Daily", TEST_DATE, CHANGEOVER_INITIAL, "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 2: BSC EM ---
    add_section_divider(t2, f"Biological Safety Cabinet (BSC) EM for BSC E00{PROCESSING_BSC} and BSC E00{CHANGEOVER_BSC} for {TEST_DATE}")
    add_data_row(t2, f"Surface Sampling of\nISO 5 BSC E00{PROCESSING_BSC}\n(4 Locations)", "Daily", TEST_DATE, PROCESSING_INITIAL, "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, f"Surface Sampling of\nISO 5 BSC E00{CHANGEOVER_BSC}\n(4 Locations)", "Daily", TEST_DATE, CHANGEOVER_INITIAL, "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, f"Settling Sampling of\nISO 5 BSC E00{PROCESSING_BSC}", "Daily", TEST_DATE, PROCESSING_INITIAL, "Date of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")
    add_data_row(t2, f"Settling Sampling of\nISO 5 BSC E00{CHANGEOVER_BSC}", "Daily", TEST_DATE, CHANGEOVER_INITIAL, "Date of Testing\n(Scan C/O)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # --- Section 3: Weekly Active Air Sampling L-Suite ---
    add_section_divider(t2, f"Weekly Active Air Sampling of Cleanroom 145 (E001979, L-Suite) with Processing and Changeover BSCs for {TEST_DATE}")

    row_l = t2.add_row()
    format_cell(row_l.cells[0], "Active Air Sampling\nof Cleanrooms", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.LEFT, v_align=WD_ALIGN_VERTICAL.CENTER)
    format_cell(row_l.cells[1], "Weekly", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)
    format_cell(row_l.cells[2], "15Sep26", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)
    format_cell(row_l.cells[3], "SMO", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)
    format_cell(row_l.cells[4], "Week of Testing\n(Scan)", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)

    # We set top alignment for columns 5, 6, 7 so that each observation, ETX, and organism block align horizontally
    c_obs = row_l.cells[5]
    c_obs.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(c_obs)

    c_etx = row_l.cells[6]
    c_etx.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(c_etx)

    c_id = row_l.cells[7]
    c_id.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(c_id)

    id_lines_0347 = [
        "Corynebacterium ureicelerivorans",
        "Brevibacterium sp",
        "Gram (+) short rods",
        "Micrococcus luteus",
        "Gram (+) cocci",
        "Moraxella osloensis",
        "Gram (-) coccobacilli"
    ]
    id_lines_0349 = [
        "Psychrobacter sp",
        "Moraxella osloensis",
        "Gram (-) coccobacilli",
        "Corynebacterium segmentosum",
        "Corynebacterium sp",
        "Gram (+) short rods",
        "Kocuria indica",
        "Kocuria rhizophila/ sp. BT304",
        "Staphylococcus lugdunensis",
        "Staphylococcus saprophyticus",
        "Staphylococcus saprophyticus",
        "Staphylococcus pasteuri",
        "Staphylococcus capitis",
        "Micrococcus luteus",
        "Gram (+) cocci",
        "Curvularia sorghina",
        "Hyphae"
    ]
    id_lines_0356 = [
        "Micrococcus luteus",
        "Staphylococcus hominis",
        "Staphylococcus saprophyticus",
        "Gram (+) cocci",
        "Corynebacterium ureicelerivorans",
        "Gram (+) short rods",
        "Moraxella osloensis",
        "Gram (-) coccobacilli"
    ]

    # --- Block 1: 0347 ---
    p_obs0 = c_obs.paragraphs[0]
    p_obs0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_obs0.paragraph_format.space_before = Pt(2)
    p_obs0.paragraph_format.space_after = Pt(0)
    p_obs0.paragraph_format.line_spacing = 1.0
    r_o0 = p_obs0.add_run("6 CFUs\n(ISO 8 143- Section I)")
    r_o0.font.name = 'Times New Roman'
    r_o0.font.size = Pt(6.5)

    p_e0 = c_etx.paragraphs[0]
    p_e0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_e0.paragraph_format.space_before = Pt(2)
    p_e0.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_e0, "https://etrax.eagleanalytical.com/Submission/Details/GVgiOQzTswrcZrcY7Rs3sg__", "ETX-260922-0347", font_size=Pt(6.5))

    p_id0 = c_id.paragraphs[0]
    p_id0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_id0.paragraph_format.space_before = Pt(2)
    p_id0.paragraph_format.space_after = Pt(0)
    p_id0.paragraph_format.line_spacing = 1.0
    r_id0 = p_id0.add_run('\n'.join(id_lines_0347))
    r_id0.font.name = 'Times New Roman'
    r_id0.font.size = Pt(5.5)

    # --- Block 2: 0349 ---
    p_obs2 = c_obs.add_paragraph()
    p_obs2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_obs2.paragraph_format.space_before = Pt(36)
    p_obs2.paragraph_format.space_after = Pt(0)
    p_obs2.paragraph_format.line_spacing = 1.0
    r_o2 = p_obs2.add_run("21 CFUs\n(ISO8 143- Section II)")
    r_o2.font.name = 'Times New Roman'
    r_o2.font.size = Pt(6.5)

    p_e2 = c_etx.add_paragraph()
    p_e2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_e2.paragraph_format.space_before = Pt(42)
    p_e2.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_e2, "https://etrax.eagleanalytical.com/Submission/Details/Vex6NX4Yn4WJMGsntGoMAg__", "ETX-260922-0349", font_size=Pt(6.5))

    p_id2 = c_id.add_paragraph()
    p_id2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_id2.paragraph_format.space_before = Pt(8)
    p_id2.paragraph_format.space_after = Pt(0)
    p_id2.paragraph_format.line_spacing = 1.0
    r_id2 = p_id2.add_run('\n'.join(id_lines_0349))
    r_id2.font.name = 'Times New Roman'
    r_id2.font.size = Pt(5.5)

    # --- Block 3: 0356 ---
    p_obs4 = c_obs.add_paragraph()
    p_obs4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_obs4.paragraph_format.space_before = Pt(106)
    p_obs4.paragraph_format.space_after = Pt(0)
    p_obs4.paragraph_format.line_spacing = 1.0
    r_o4 = p_obs4.add_run("9 CFUs\n(ISO8 142)")
    r_o4.font.name = 'Times New Roman'
    r_o4.font.size = Pt(6.5)

    p_e4 = c_etx.add_paragraph()
    p_e4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_e4.paragraph_format.space_before = Pt(112)
    p_e4.paragraph_format.space_after = Pt(0)
    add_hyperlink(p_e4, "https://etrax.eagleanalytical.com/Submission/Details/21hYsvDl7c1cAl93FSN-sw__", "ETX-260922-0356", font_size=Pt(6.5))

    p_id4 = c_id.add_paragraph()
    p_id4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_id4.paragraph_format.space_before = Pt(8)
    p_id4.paragraph_format.space_after = Pt(0)
    p_id4.paragraph_format.line_spacing = 1.0
    r_id4 = p_id4.add_run('\n'.join(id_lines_0356))
    r_id4.font.name = 'Times New Roman'
    r_id4.font.size = Pt(5.5)

    # Notes cell
    format_cell(row_l.cells[8], "None", bold=False, font_size=Pt(6.5), align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)

    # --- Section 4: Weekly Surface Sampling L-Suite ---
    add_section_divider(t2, f"Weekly Surface Sampling of Cleanroom 145 (E001979, L-Suite) with Processing and Changeover BSCs for {TEST_DATE}")
    add_data_row(t2, "Surface Sampling of\nCleanrooms", "Weekly", "15Sep26", "SMO", "Week of Testing\n(Scan)", "No Growth", "Not Applicable", "Not Applicable", "None")

    # Set column widths
    for row in t2.rows:
        if len(row.cells) == 9:
            for j, w in enumerate(t2_widths):
                row.cells[j].width = w

    doc.save(output_docx_path)
    print(f"Saved DOCX to: {output_docx_path}")

def convert_docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
        print(f"Saved PDF to: {pdf_path}")
    finally:
        word.Quit()

if __name__ == "__main__":
    out_docx = os.path.join(DESKTOP_DIR, "Tables OOS-262150 BelieveRX Pharmacy (E74120) - ScanRDI.docx")
    out_pdf = os.path.join(DESKTOP_DIR, "Tables OOS-262150 BelieveRX Pharmacy (E74120) - ScanRDI.pdf")
    build_tables_doc(out_docx)
    convert_docx_to_pdf(out_docx, out_pdf)

    doc_fitz = fitz.open(out_pdf)
    print(f"Generated PDF Page Count: {len(doc_fitz)}")
    for i, p in enumerate(doc_fitz):
        pix = p.get_pixmap(dpi=150)
        img_out = os.path.join(SCRATCH_DIR, f"tables_262150_p{i+1}.png")
        pix.save(img_out)
        print(f"Rendered page {i+1} to {img_out}")
