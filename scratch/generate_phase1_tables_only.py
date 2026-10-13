import os, sys, shutil, copy
import docx
from docx.shared import Pt, Inches
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE
from docxtpl import DocxTemplate
import win32com.client
from pypdf import PdfReader
import fitz

OOS_ROOT = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS"
sys.path.insert(0, OOS_ROOT)
os.chdir(OOS_ROOT)

OUTPUT_DIR = os.path.join(OOS_ROOT, "scratch")
DESKTOP_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
DOCUMENTS_DIR = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents"

TPL_BASE = os.path.join(OUTPUT_DIR, "tables_71_pure_celsis.docx")
if not os.path.exists(TPL_BASE):
    TPL_BASE = os.path.join(OOS_ROOT, "tables for celsis.docx")

data = {
    "oos_id": "262098",
    "client_name": "Solyn LLC (E75000)",
    "sample_id": "ETX-260902-0505",
    "sample_id_1": "ETX-260902-0505",
    "sample_url": "https://etrax.eagleanalytical.com/Submission/Details/%247Ydn%24JOiYgmsnaKjbxn-g__",
    "sample_name": "GLP3R/Cagrilinitide",
    "lot_number": "2608-216",
    "analyst_name": "Alex Saravia",
    "reading_name": "Elysse Nioupin",
    "positive_id": "ETX-260910-0290",
    "positive_url": "https://etrax.eagleanalytical.com/Submission/Details?id=oZa2TD1UI-0eAJOFysBJ3A__",
    "positive_media": "1 x 100mL TSB bottle",
    "positive_org": "Microbacterium sp. PM5\n(Gram (+) rods)",

    # Table 2 Dates & Analysts (04Sep26)
    "process_date": "04Sep26",
    "pro_before_test": "03Sep26",
    "pro_test_date": "04Sep26",
    "pro_after_test": "08Sep26",
    "pro_analyst_initial": "AS",
    "pro_date_of_weekly": "04Sep26",

    # Personnel EM
    "pro_be_obs_pers_dur_pro": "No growth", "pro_be_etx_pers_dur_pro": "N/A", "pro_be_id_pers_dur_pro": "N/A",
    "pro_obs_pers_dur_pro": "No growth",    "pro_etx_pers_dur_pro": "N/A", "pro_id_pers_dur_pro": "N/A",
    "pro_af_obs_pers_dur_pro": "No growth", "pro_af_etx_pers_dur_pro": "N/A", "pro_id_pers_dur_pro": "N/A",

    # Surface ISO 5
    "pro_be_obs_surf_dur_pro": "No growth", "pro_be_etx_surf_dur_pro": "N/A", "pro_be_id_surf_dur_pro": "N/A",
    "pro_obs_surf_dur_pro": "No growth",    "pro_etx_surf_dur_pro": "N/A", "pro_id_surf_dur_pro": "N/A",
    "pro_af_obs_surf_dur_pro": "No growth", "pro_af_etx_surf_dur_pro": "N/A", "pro_id_surf_dur_pro": "N/A",

    # Settling ISO 5
    "pro_be_obs_sett_dur_pro": "No growth", "pro_be_etx_sett_dur_pro": "N/A", "pro_be_id_sett_dur_pro": "N/A",
    "pro_obs_sett_dur_pro": "No growth",    "pro_etx_sett_dur_pro": "N/A", "pro_id_sett_dur_pro": "N/A",
    "pro_af_obs_sett_dur_pro": "No growth", "pro_af_etx_sett_dur_pro": "N/A", "pro_id_sett_dur_pro": "N/A",

    # Weekly Active Air & Surface
    "pro_obs_air_wk_of": "2 CFUs (ISO 8 114)",
    "pro_etx_air_wk_of": "ETX-260914-0487",
    "pro_id_air_wk_of": "Corynebacterium ureicelerivorans\nMycobacterium grossiae",
    "pro_obs_room_wk_of": "No growth",
    "pro_etx_room_wk_of": "N/A",
    "pro_id_room_wk_of": "N/A"
}

OUT_DOCX = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - Phase I.docx")
OUT_PDF = os.path.join(OUTPUT_DIR, f"Tables OOS-{data['oos_id']} {data['client_name']} - Phase I.pdf")

tpl = DocxTemplate(TPL_BASE)
tpl.render(data)
tpl.save(OUT_DOCX)

doc_rendered = docx.Document(OUT_DOCX)

URL_0487 = "https://etrax.eagleanalytical.com/SubmissionTest/Details/fd3G2StZClcy1TP2ES6BLw__"

def set_cell_clean_text(cell, text, font_size_pt=7, bold=False):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    b_tag = '<w:b/><w:bCs/>' if bold else ''
    lines = text.split('\n')
    for line in lines:
        p_xml = (
            f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:pPr>'
            f'<w:jc w:val="center"/>'
            f'<w:spacing w:before="0" w:after="0" w:line="220" w:lineRule="auto"/>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'{b_tag}'
            f'<w:sz w:val="{int(font_size_pt*2)}"/>'
            f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
            f'</w:rPr>'
            f'</w:pPr>'
            f'<w:r>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'{b_tag}'
            f'<w:sz w:val="{int(font_size_pt*2)}"/>'
            f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
            f'</w:rPr>'
            f'<w:t>{line.strip()}</w:t>'
            f'</w:r>'
            f'</w:p>'
        )
        tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_hyperlink_to_cell(cell, url, text, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    part = cell.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:hyperlink r:id="{r_id}" w:history="1">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:color w:val="0000FF"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'<w:u w:val="single"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_dual_microbial_id_cell(cell, org1, org2, font_size_pt=6.5):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{org1}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    p2_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{org2}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    p_empty_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="120" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_empty_xml))
    tc.append(parse_xml(p2_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_table1_microbial_id_cell(cell, species, gram_stain, font_size_pt=7):
    tc = cell._tc
    for p in tc.xpath('w:p'):
        tc.remove(p)
    p_xml = (
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:i/>'
        f'<w:iCs/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:t>{species}</w:t>'
        f'</w:r>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:sz w:val="{int(font_size_pt*2)}"/>'
        f'<w:szCs w:val="{int(font_size_pt*2)}"/>'
        f'</w:rPr>'
        f'<w:br/>'
        f'<w:t>{gram_stain}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    tc.append(parse_xml(p_xml))
    cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER

# 1. Finalize Table 1 for Phase I (ETX-260902-0505 only, exactly 1 row)
# Update Heading of Table 1
for p in doc_rendered.paragraphs:
    if "Table 1:" in p.text:
        p.text = "Table 1: Information for ETX-260902-0505 under investigation"
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(9)
        p.runs[0].bold = True
        p.runs[0].underline = True

t1_rendered = doc_rendered.tables[0]
add_hyperlink_to_cell(t1_rendered.rows[1].cells[2], data["sample_url"], data["sample_id_1"], font_size_pt=7)
add_hyperlink_to_cell(t1_rendered.rows[1].cells[3], data["positive_url"], data["positive_id"], font_size_pt=7)
set_table1_microbial_id_cell(t1_rendered.rows[1].cells[5], "Microbacterium sp. PM5", "(Gram (+) rods)", font_size_pt=7)

# 2. Finalize Table 2 (04Sep26 Processing)
t2_rendered = doc_rendered.tables[1]
set_cell_clean_text(t2_rendered.rows[1].cells[0], "Personnel EM", font_size_pt=7, bold=True)
set_cell_clean_text(t2_rendered.rows[5].cells[0], "Biological Safety Cabinet EM for E001316", font_size_pt=7, bold=True)
set_cell_clean_text(t2_rendered.rows[12].cells[0], "Weekly Active Air Sampling of Cleanroom Suite 114 (E001736) with Processing BSC", font_size_pt=7, bold=True)
set_cell_clean_text(t2_rendered.rows[14].cells[0], "Weekly Surface Sampling of Cleanroom Suite 114 (E001736) with Processing BSC", font_size_pt=7, bold=True)

# Row 13: Weekly Active Air
set_cell_clean_text(t2_rendered.rows[13].cells[3], "04Sep26", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[13].cells[4], "ISS", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[13].cells[6], "Week of Testing", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[13].cells[8], "2 CFUs\n(ISO 8 114)", font_size_pt=7)
add_hyperlink_to_cell(t2_rendered.rows[13].cells[9], URL_0487, "ETX-260914-0487", font_size_pt=7)
set_dual_microbial_id_cell(t2_rendered.rows[13].cells[10], "Corynebacterium ureicelerivorans", "Mycobacterium grossiae", font_size_pt=6.5)

# Row 15: Weekly Surface
set_cell_clean_text(t2_rendered.rows[15].cells[3], "04Sep26", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[4], "ISS", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[6], "Week of Testing", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[8], "No growth", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[9], "N/A", font_size_pt=7)
set_cell_clean_text(t2_rendered.rows[15].cells[10], "N/A", font_size_pt=7)

doc_rendered.save(OUT_DOCX)
print("Saved 1-page Phase I tables docx to:", OUT_DOCX)

# Convert to PDF via Word COM
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    doc_com = word.Documents.Open(os.path.abspath(OUT_DOCX))
    doc_com.SaveAs(os.path.abspath(OUT_PDF), FileFormat=17) # 17 = wdFormatPDF
    doc_com.Close()
    print("Converted Phase I tables to PDF:", OUT_PDF)
except Exception as e:
    print("Word COM conversion error:", e)
finally:
    word.Quit()

# Check page count
doc_pdf = fitz.open(OUT_PDF)
print(f"Phase I Table PDF total pages: {len(doc_pdf)}")
doc_pdf.close()
