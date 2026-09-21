import docx, os, sys, shutil
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

desktop_dir = r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Desktop"
src_docx = os.path.join(desktop_dir, "EM table OOS-261185 08MAY2026.docx")
backup_docx = os.path.join(r"C:\Users\qchen\OneDrive - Professional Compounding Centers of America, Inc\Documents\OOS\.history", "EM_table_261185_backup.docx")

# Backup original
os.makedirs(os.path.dirname(backup_docx), exist_ok=True)
if not os.path.exists(backup_docx):
    shutil.copyfile(src_docx, backup_docx)
    print(f"Backed up original table to: {backup_docx}")

doc = docx.Document(src_docx)

# Set tight margins so it fits strictly on 1 page
for section in doc.sections:
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

# Update Table 2 title if needed
for p in doc.paragraphs:
    if "Table 2:" in p.text:
        p.text = ""
        r = p.add_run("Table 2: Environmental Monitoring for Analyst & Cleanroom Bracketing")
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.bold = True
    elif "Table 1:" in p.text:
        p.text = ""
        r = p.add_run("Table 1: Read Dates & Incubation Observation")
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.bold = True

# --- TABLE 1 AMENDMENT ---
t1 = doc.tables[0]
# Cell (1, 0): ETX Submission ID -> ETX-260518-0250
t1.rows[1].cells[0].text = "ETX-260518-0250"
# Cell (1, 1): SMO 08MAY 2026
t1.rows[1].cells[1].text = "SMO\n08MAY 2026"
# Cell (1, 7): Microbial ID - fix Talaromyces purpurogenus to Penicillium decumbens...
t1.rows[1].cells[7].text = (
    "Penicillium decumbens,\n"
    "Cladosporium tenuissimum,\n"
    "Cladosporium langeronii,\n"
    "Cladosporium halotolerans"
)

# --- TABLE 2 AMENDMENT ---
t2 = doc.tables[1]

# Row 2 (29APR 2026 Active Air):
t2.rows[2].cells[4].text = "Week Before Testing Date"

# Row 3 (08MAY 2026 Active Air - Week of testing):
t2.rows[3].cells[4].text = "Week of Testing Date"
t2.rows[3].cells[6].text = "ETX-260518-0249\nETX-260518-0250"
t2.rows[3].cells[7].text = (
    "Staphylococcus aureus\n"
    "Penicillium decumbens\n"
    "Cladosporium tenuissimum\n"
    "Cladosporium langeronii\n"
    "Cladosporium halotolerans"
)

# Row 4 (14MAY 2026 Active Air - Week after testing):
t2.rows[4].cells[4].text = "Week After Testing Date"
t2.rows[4].cells[6].text = "ETX-260526-0391\nETX-260526-0395"
t2.rows[4].cells[7].text = "Gram (+) cocci\nCandida orthopsilosis\nBudding yeast"

# Row 6 (29APR 2026 Surface):
t2.rows[6].cells[0].text = "Surface Sampling of Cleanrooms"
t2.rows[6].cells[4].text = "Week Before Testing Date"

# Row 7 (08MAY 2026 Surface):
t2.rows[7].cells[4].text = "Week of Testing Date"

# Row 8 (14MAY 2026 Surface):
t2.rows[8].cells[4].text = "Week After Testing Date"
t2.rows[8].cells[6].text = "ETX-260526-0424\nETX-260526-0425"
t2.rows[8].cells[7].text = "Gram (+) rods\nStaphylococcus capitis"

# Apply formatting (font size, cell margins, compact padding)
def set_cell_margins(cell, top=30, bottom=30, left=50, right=50):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

for t in [t1, t2]:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(t.rows):
        for cell in row.cells:
            set_cell_margins(cell, top=25, bottom=25, left=40, right=40)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(7.5)
                    if r_idx == 0:
                        run.font.bold = True

doc.save(src_docx)
print(f"Amended table saved to: {src_docx}")

# Also save to standardized name
target_tbl_docx = os.path.join(desktop_dir, "Tables OOS-261185 EM SMO 115B Air 08MAY2026 - EM.docx")
shutil.copyfile(src_docx, target_tbl_docx)
print(f"Saved standardized table DOCX: {target_tbl_docx}")
